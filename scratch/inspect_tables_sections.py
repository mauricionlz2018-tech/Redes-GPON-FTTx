import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
print("Num sections:", len(doc.sections))
for s_idx, s in enumerate(doc.sections):
    print(f"Section {s_idx}: header pars={len(s.header.paragraphs)}, footer pars={len(s.footer.paragraphs)}")

print("\nTables count:", len(doc.tables))
for t_idx, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns)
    first_cell = tbl.cell(0, 0).text.strip().replace('\n', ' ')[:40]
    print(f"Table {t_idx}: {rows}x{cols} | first cell='{first_cell}'")

