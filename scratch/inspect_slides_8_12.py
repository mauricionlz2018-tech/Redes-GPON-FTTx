import pptx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(os.path.join('docs', 'PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx'))

for idx in [7, 8, 9, 10, 11]:
    s = prs.slides[idx]
    print(f"\n================ SLIDE {idx+1} ================")
    for shp in s.shapes:
        txt = shp.text.strip().replace('\n', ' // ') if shp.has_text_frame else ''
        print(f"  Shape: '{shp.name}' [{shp.shape_type}], pos=({shp.left/914400:.2f}, {shp.top/914400:.2f}), size=({shp.width/914400:.2f} x {shp.height/914400:.2f}) | Text: '{txt[:60]}'")
