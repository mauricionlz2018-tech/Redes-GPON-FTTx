import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
print("Total paragraphs:", len(doc.paragraphs))

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if p.style.name.startswith("Heading 1") or p.style.name.startswith("Heading 2"):
        print(f"P[{i}] ({p.style.name}): {t}")
    elif any(k in t.upper() for k in ["CAPÍTULO", "CAPITULO", "CONCLUSIONES", "ANEXOS"]):
        print(f"P[{i}] (SPECIAL {p.style.name}): {t}")

