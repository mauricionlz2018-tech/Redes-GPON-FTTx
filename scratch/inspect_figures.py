import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")

for i, p in enumerate(doc.paragraphs):
    has_img = any('drawing' in r._element.xml for r in p.runs)
    if 'Figura' in p.text or has_img:
        print(f"P[{i}]: imgs={1 if has_img else 0} | text={p.text[:80]}")
