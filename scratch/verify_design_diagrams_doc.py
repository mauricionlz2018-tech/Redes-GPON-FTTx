import docx
import re

def verify_doc(doc_path):
    print(f"\n==========================================")
    print(f"VERIFYING: {doc_path}")
    print(f"==========================================")
    
    doc = docx.Document(doc_path)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total sections: {len(doc.sections)}")
    
    # 1. Margins check on content sections
    for i, s in enumerate(doc.sections):
        print(f"Section {i}: {s.page_width.cm:.2f} x {s.page_height.cm:.2f} cm | Margins: L={s.left_margin.cm:.2f}, R={s.right_margin.cm:.2f}, T={s.top_margin.cm:.2f}, B={s.bottom_margin.cm:.2f}")
        
    # 2. Check figure index
    idx_figs = []
    in_fig_idx = False
    for i, p in enumerate(doc.paragraphs[:120]):
        txt = p.text.strip()
        if 'ndice de figuras' in txt:
            in_fig_idx = True
            continue
        if 'ndice de tablas' in txt:
            in_fig_idx = False
            break
        if in_fig_idx and txt.startswith("Figura "):
            idx_figs.append((i, txt))
            
    print(f"\nFigure index entries found: {len(idx_figs)}")
    assert len(idx_figs) == 73, f"Expected 73 index entries, found {len(idx_figs)}"
    for k, (p_idx, txt) in enumerate(idx_figs):
        expected_num = k + 1
        m = re.match(r'^Figura\s+(\d+)\.', txt)
        assert m is not None, f"TOC entry {p_idx} does not match format: {txt}"
        num = int(m.group(1))
        assert num == expected_num, f"TOC entry {k} has num {num} but expected {expected_num}: {txt}"
        
    print("ALL 73 TOC ENTRIES STRICTLY SEQUENTIAL (1 to 73)!")
    
    # 3. Check body figures
    body_figs = []
    for i, p in enumerate(doc.paragraphs[105:], start=105):
        txt = p.text.strip()
        if txt.startswith("Figura "):
            # Check drawing in previous 2 paragraphs (to allow for blank line)
            prev_p1 = doc.paragraphs[i-1]
            prev_p2 = doc.paragraphs[i-2] if i >= 2 else None
            dw_count = len(prev_p1._p.xpath('.//w:drawing')) + len(prev_p1._p.xpath('.//w:pict'))
            if prev_p2 is not None:
                dw_count += len(prev_p2._p.xpath('.//w:drawing')) + len(prev_p2._p.xpath('.//w:pict'))
                
            next_p = doc.paragraphs[i+1]
            next_txt = next_p.text.strip()
            
            body_figs.append({
                'p_idx': i,
                'text': txt,
                'has_drawing_before': dw_count > 0,
                'next_text': next_txt,
                'style': p.style.name
            })
            
    print(f"\nBody figures found: {len(body_figs)}")
    assert len(body_figs) == 73, f"Expected 73 body figures, found {len(body_figs)}"
    
    for k, fig in enumerate(body_figs):
        expected_num = k + 1
        m = re.match(r'^Figura\s+(\d+)\.', fig['text'])
        assert m is not None, f"Figure at P[{fig['p_idx']}] format invalid: {fig['text']}"
        num = int(m.group(1))
        assert num == expected_num, f"Figure {k} has num {num} but expected {expected_num}: {fig['text']}"
        assert fig['has_drawing_before'], f"Figure {num} at P[{fig['p_idx']}] is missing drawing in previous paragraph!"
        
    print("ALL 73 BODY FIGURES STRICTLY SEQUENTIAL (1 to 73) WITH VALID DRAWINGS!")

    # 4. Print specifically the 7 newly added figures in the document
    print("\n--- NEWLY INSERTED DESIGN DIAGRAMS VERIFICATION ---")
    new_indices = [27, 30, 33, 34, 36, 37, 38] # 0-indexed for Fig 28, 31, 34, 35, 37, 38, 39
    for idx in new_indices:
        f = body_figs[idx]
        print(f"P[{f['p_idx']}]: {f['text']}")
        print(f"   Image preceding: OK | Note following: '{f['next_text']}' | Style: {f['style']}")
        assert 'laboraci' in f['next_text'], f"New figure missing Elaboración propia: {f['next_text']}"
        assert f['style'] == 'Figuras', f"New figure invalid style: {f['style']}"
        
    print("\nVERIFICATION PASSED 100%!")

verify_doc('docs/Documentacion_Residencias_avance.docx')
verify_doc('docs/Documentacion_Residencias (4).docx')

