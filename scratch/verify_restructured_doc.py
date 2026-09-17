import docx
import os
import re

def verify():
    doc_path = 'docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx'
    print(f"=== INICIANDO AUDITORÍA AUTOMATIZADA: {doc_path} ===")
    
    if not os.path.exists(doc_path):
        print(f"ERROR: No existe el archivo {doc_path}")
        return False
        
    doc = docx.Document(doc_path)
    file_size_mb = os.path.getsize(doc_path) / (1024 * 1024)
    print(f"Tamaño de archivo: {file_size_mb:.2f} MB")
    print(f"Total de párrafos: {len(doc.paragraphs)}")
    print(f"Total de tablas: {len(doc.tables)}")

    # 1. Capítulos
    chapters_found = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip().upper()
        if t.startswith("CAPÍTULO") or t.startswith("CAPITULO") or t == "ANEXOS":
            chapters_found.append((i, p.text.strip()))

    print("\n--- Capítulos Detectados ---")
    for idx, chap in chapters_found:
        print(f"  P[{idx}]: {chap}")

    # 2. Figuras en el cuerpo (con imagen y fuente)
    body_figures = []
    index_figures = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("Figura ") or t.startswith("figura "):
            if i < 120:
                index_figures.append((i, t))
            else:
                prev_p = doc.paragraphs[i-1] if i > 0 else None
                has_drawing = ('<w:drawing' in prev_p._p.xml or '<w:pict' in prev_p._p.xml) if prev_p else False
                next_p = doc.paragraphs[i+1] if i+1 < len(doc.paragraphs) else None
                has_source = (next_p and (next_p.text.strip().startswith("Fuente:") or "Fuente" in next_p.text)) if next_p else False
                body_figures.append({
                    "index": i,
                    "caption": t,
                    "has_drawing_above": has_drawing,
                    "has_source_below": has_source
                })

    print(f"\n--- Figuras en Índice: {len(index_figures)} | Figuras en Cuerpo: {len(body_figures)} ---")
    for f in body_figures:
        num_match = re.search(r'Figura (\d+)', f['caption'])
        num = num_match.group(1) if num_match else '?'
        status = "OK" if (f['has_drawing_above'] and f['has_source_below']) else "CHECK"
        print(f"  Fig #{num:2s} (P[{f['index']:3d}]): {status} | {f['caption'][:70]}")

    # 3. Tablas
    tables_found = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith("Tabla ") or t.startswith("tabla "):
            tables_found.append((i, t))

    print(f"\n--- Total de Leyendas de Tablas Detectadas: {len(tables_found)} ---")
    for idx, tbl in tables_found:
        print(f"  P[{idx}]: {tbl[:70]}")

    # 4. Prohibiciones
    em_dashes = 0
    semicolons = 0
    emojis = 0
    casio_count = 0
    leonardo_count = 0

    emoji_pattern = re.compile(r'[\U00010000-\U0010ffff]|[\u2600-\u27bf]|[\u2300-\u23ff]|[\u2b50-\u2b55]')

    for p in doc.paragraphs:
        t = p.text
        em_dashes += t.count('—') + t.count('–')
        semicolons += t.count(';')
        if emoji_pattern.search(t):
            emojis += 1
        if 'casio' in t.lower():
            casio_count += 1
        if 'leonardo' in t.lower():
            leonardo_count += 1

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                txt = cell.text
                em_dashes += txt.count('—') + txt.count('–')
                semicolons += txt.count(';')
                if emoji_pattern.search(txt):
                    emojis += 1
                if 'casio' in txt.lower():
                    casio_count += 1
                if 'leonardo' in txt.lower():
                    leonardo_count += 1

    print("\n--- Auditoría de Restricciones y Calidad ---")
    print(f"  Guiones largos ('—' / '–'): {em_dashes} (Debe ser 0)")
    print(f"  Puntos y coma (';'):         {semicolons} (Debe ser 0)")
    print(f"  Emojis o glifos informales: {emojis} (Debe ser 0)")
    print(f"  Menciones a 'Casio':        {casio_count} (Debe ser 0)")
    print(f"  Menciones a 'Leonardo':     {leonardo_count} (Asesor oficial)")

    all_ok = (
        len(chapters_found) >= 7 and
        len(index_figures) == 29 and
        len(body_figures) == 29 and
        em_dashes == 0 and
        semicolons == 0 and
        emojis == 0 and
        casio_count == 0 and
        leonardo_count > 0
    )
    print(f"\nRESULTADO FINAL DE LA AUDITORÍA: {'APROBADO AL 100%' if all_ok else 'REVISIÓN REQUERIDA'}")
    return all_ok

if __name__ == '__main__':
    verify()
