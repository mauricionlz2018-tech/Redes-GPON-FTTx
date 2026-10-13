import docx
from docx.shared import Pt, Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import shutil
import os

def apply_updates():
    doc_path = 'docs/Documentacion_Residencias_avance.docx'
    backup_path = 'docs/Documentacion_Residencias_avance_BACKUP_08OCT_PAGINACION.docx'
    
    # 1. Crear respaldo preventivo
    print(f"Creando respaldo en {backup_path}...")
    shutil.copyfile(doc_path, backup_path)
    print("Respaldo creado con éxito.")

    # 2. Cargar documento
    print(f"Abriendo {doc_path}...")
    doc = docx.Document(doc_path)

    # ----------------------------------------------------
    # PARTE A: REORDENAMIENTO Y DIVISIÓN DE REFERENCIAS
    # ----------------------------------------------------
    print("Configurando sección de Referencias...")

    # P1184: Título principal en mayúsculas y negritas (Regla 13)
    p_ref_title = None
    for p in doc.paragraphs:
        if p.text.strip().lower() == 'referencias':
            p_ref_title = p
            break
            
    if p_ref_title:
        p_ref_title.text = 'REFERENCIAS'
        p_ref_title.runs[0].font.name = 'Arial'
        p_ref_title.runs[0].font.size = Pt(14)
        p_ref_title.runs[0].font.bold = True
        p_ref_title.paragraph_format.space_before = Pt(18)
        p_ref_title.paragraph_format.space_after = Pt(12)
        # Asegurar outlineLvl 0
        p_el = p_ref_title._element.get_or_add_pPr()
        for lvl in p_el.xpath('./w:outlineLvl'):
            p_el.remove(lvl)
        p_el.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
        print("  -> Título REFERENCIAS actualizado a mayúsculas y negrita.")

    # Subtítulo 1: Referencias bibliográficas
    p_bib_sub = None
    for p in doc.paragraphs:
        if p.text.strip().lower() == 'referencias bibliográficas' or p.text.strip().lower() == 'referencias bibliograficas':
            p_bib_sub = p
            break

    if p_bib_sub:
        p_bib_sub.text = 'Referencias bibliográficas'
        p_bib_sub.runs[0].font.name = 'Arial'
        p_bib_sub.runs[0].font.size = Pt(12)
        p_bib_sub.runs[0].font.bold = True
        p_bib_sub.paragraph_format.space_before = Pt(12)
        p_bib_sub.paragraph_format.space_after = Pt(6)
        p_el = p_bib_sub._element.get_or_add_pPr()
        for lvl in p_el.xpath('./w:outlineLvl'):
            p_el.remove(lvl)
        p_el.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="1"/>'))
        print("  -> Subtítulo 'Referencias bibliográficas' verificado (Nivel 2).")

    # Localizar y remover GPON Telecom de la sección bibliográfica si existe ahí
    gpon_text = 'GPON Telecom S.A. de C.V. (2026). Infraestructura para la vida: Nosotros. https://gpontelecom.com.mx/#div_nosotros'
    gpon_p_old = None
    for p_idx, p in enumerate(doc.paragraphs):
        if 'gpon telecom s.a. de c.v. (2026)' in p.text.lower() and 'infraestructura para la vida' in p.text.lower():
            print(f"  -> Encontrada referencia previa de GPON Telecom en párrafo {p_idx}.")
            gpon_p_old = p
            break

    # Localizar punto de inicio de referencias electrónicas (Agafonkin)
    p_agafonkin = None
    agafonkin_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith('Agafonkin, V. (2023)'):
            p_agafonkin = p
            agafonkin_idx = idx
            break

    if p_agafonkin is not None:
        print(f"  -> Punto de división para referencias electrónicas localizado en párrafo {agafonkin_idx} (Agafonkin).")
        
        # Insertar subtítulo 'Referencias electrónicas' antes de Agafonkin
        new_sub_p = parse_xml(f'<w:p {nsdecls("w")}/>')
        p_agafonkin._element.addprevious(new_sub_p)
        p_elec_sub = docx.text.paragraph.Paragraph(new_sub_p, doc)
        p_elec_sub.text = 'Referencias electrónicas'
        p_elec_sub.runs[0].font.name = 'Arial'
        p_elec_sub.runs[0].font.size = Pt(12)
        p_elec_sub.runs[0].font.bold = True
        p_elec_sub.paragraph_format.space_before = Pt(18)
        p_elec_sub.paragraph_format.space_after = Pt(6)
        
        # outlineLvl = 1
        p_el = p_elec_sub._element.get_or_add_pPr()
        p_el.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="1"/>'))
        print("  -> Insertado subtítulo 'Referencias electrónicas' (Nivel 2).")

        # Si gpon_p_old existía en la sección bibliográfica, removerlo de ahí
        if gpon_p_old is not None:
            # Eliminar su elemento XML del cuerpo
            gpon_p_old._element.getparent().remove(gpon_p_old._element)
            print("  -> Removida referencia GPON Telecom de la sección bibliográfica.")

        # Ahora insertar GPON Telecom en la sección electrónica en orden alfabético
        # Buscar Google Chrome Developers para insertar GPON después de él (y antes de Grafana Labs)
        p_target_google = None
        for p in doc.paragraphs:
            if 'Google Chrome Developers' in p.text:
                p_target_google = p
                break

        if p_target_google is not None:
            new_gpon_xml = parse_xml(f'<w:p {nsdecls("w")}/>')
            p_target_google._element.addnext(new_gpon_xml)
            p_gpon = docx.text.paragraph.Paragraph(new_gpon_xml, doc)
            
            # Formato APA 7 con sangría francesa y cursiva en el título del recurso
            p_gpon.paragraph_format.line_spacing = 1.5
            p_gpon.paragraph_format.space_after = Pt(6)
            p_gpon.paragraph_format.left_indent = Pt(36)
            p_gpon.paragraph_format.first_line_indent = Pt(-36)
            
            r1 = p_gpon.add_run('GPON Telecom S.A. de C.V. (2026). ')
            r2 = p_gpon.add_run('Infraestructura para la vida: Nosotros. ')
            r2.font.italic = True
            r3 = p_gpon.add_run('https://gpontelecom.com.mx/#div_nosotros')
            
            for r in [r1, r2, r3]:
                r.font.name = 'Arial'
                r.font.size = Pt(11)
            print("  -> Insertada referencia formal GPON Telecom en 'Referencias electrónicas' en orden alfabético estricto.")

    # ----------------------------------------------------
    # PARTE B: CONFIGURACIÓN EXACTA DE PAGINACIÓN INSTITUCIONAL
    # ----------------------------------------------------
    print("\nConfigurando paginación institucional por secciones...")

    # Sección 0 (Carátula y Oficios):
    # Regla 7: No se numeran, pero cuentan como 1 y 2
    s0 = doc.sections[0]
    s0.footer.is_linked_to_previous = False
    for p in s0.footer.paragraphs:
        p.text = ''
    # Limpiar pgNumType o establecer start=1 decimal
    s0_el = s0._sectPr
    for el in s0_el.xpath('./w:pgNumType'):
        s0_el.remove(el)
    s0_el.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="decimal"/>'))
    print("  -> Sección 0 (Carátula): Sin paginación visible. Contabiliza páginas 1 y 2.")

    # Sección 1 (Preliminares: Índices de contenido, figuras y tablas):
    # Regla 8: Números romanos mayúsculas (upperRoman), parte inferior derecha, consecutivo a carátula (start=3)
    s1 = doc.sections[1]
    s1.footer.is_linked_to_previous = False
    s1_el = s1._sectPr
    for el in s1_el.xpath('./w:pgNumType'):
        s1_el.remove(el)
    s1_el.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="3" w:fmt="upperRoman"/>'))
    print("  -> Sección 1 (Preliminares): Formato upperRoman iniciando en III (página 3).")

    # Sección 2 (Cuerpo Principal / Capítulos I al VI):
    # Regla 9: Números arábigos (decimal), consecutivos a la numeración anterior (inicia en 14), parte inferior derecha
    s2 = doc.sections[2]
    s2.footer.is_linked_to_previous = False
    s2_el = s2._sectPr
    for el in s2_el.xpath('./w:pgNumType'):
        s2_el.remove(el)
    # Sin w:start para que sea automáticamente consecutivo a la sección anterior (14 en adelante)
    s2_el.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal"/>'))
    print("  -> Sección 2 (Capítulos y cuerpo principal): Formato arábigo consecutivo desde la sección anterior (página 14 en adelante).")

    # Sección 3 (Conclusiones y Glosario):
    # Regla 10: Números romanos en mayúscula nuevamente (upperRoman) y consecutivos a la paginación, parte inferior derecha
    s3 = doc.sections[3]
    s3.footer.is_linked_to_previous = False
    s3_el = s3._sectPr
    for el in s3_el.xpath('./w:pgNumType'):
        s3_el.remove(el)
    # Sin w:start para que sea consecutivo a la paginación del documento en números romanos mayúsculos
    s3_el.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="upperRoman"/>'))
    print("  -> Sección 3 (Conclusiones y Glosario): Formato upperRoman consecutivo a la paginación general.")

    # Sección 4 (Referencias):
    # Regla 11: Sin paginación, pero mencionadas en el índice
    s4 = doc.sections[4]
    s4.footer.is_linked_to_previous = False
    for p in s4.footer.paragraphs:
        p.text = ''
    s4_el = s4._sectPr
    for el in s4_el.xpath('./w:pgNumType'):
        s4_el.remove(el)
    print("  -> Sección 4 (Referencias): Sin paginación visible.")

    # Sección 5 (Anexos):
    # Regla 11: Sin paginación, pero mencionados en el índice
    s5 = doc.sections[5]
    s5.footer.is_linked_to_previous = False
    for p in s5.footer.paragraphs:
        p.text = ''
    s5_el = s5._sectPr
    for el in s5_el.xpath('./w:pgNumType'):
        s5_el.remove(el)
    print("  -> Sección 5 (Anexos): Sin paginación visible.")

    # 4. Guardar archivo principal
    print(f"\nGuardando cambios en {doc_path}...")
    doc.save(doc_path)
    print("Documentacion_Residencias_avance.docx guardado exitosamente.")

if __name__ == '__main__':
    apply_updates()
