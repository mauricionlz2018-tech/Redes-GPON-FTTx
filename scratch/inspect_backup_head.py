import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')
print("Total paragraphs:", len(doc.paragraphs))
print("Total tables:", len(doc.tables))

for i in range(0, 20):
    p = doc.paragraphs[i]
    t = p.text.strip()
    xml = p._p.xml
    has_break = 'w:br' in xml or 'w:lastRenderedPageBreak' in xml
    print(f"P[{i}] (style={p.style.name}, break={has_break}): '{t}'")

