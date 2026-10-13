import docx

doc_dev = docx.Document('docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx')
for i in range(len(doc_dev.paragraphs) - 15, len(doc_dev.paragraphs)):
    p = doc_dev.paragraphs[i]
    print(f"P{i} [{p.style.name}]: '{p.text.strip()[:70]}'")

