import pptx

prs = pptx.Presentation(r'docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx')
print(f'Total slides: {len(prs.slides)}')
for idx, s in enumerate(prs.slides):
    texts = []
    page_num = ''
    for shp in s.shapes:
        if shp.has_text_frame:
            t = shp.text_frame.text.strip().replace('\n', ' | ')
            if '/' in t and len(t) <= 6:
                page_num = t
            elif t:
                texts.append(t[:70])
    print(f'Slide {idx+1:2d} [Pag: {page_num:<5}] Shapes: {len(s.shapes):2d} | Content: {" /// ".join(texts[:2])}')

