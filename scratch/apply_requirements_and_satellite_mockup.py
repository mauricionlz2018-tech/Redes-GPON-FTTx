import os
import re
import copy
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_ADVANCE = 'docs/Documentacion_Residencias_avance.docx'
DOC_MIRROR = 'docs/Documentacion_Residencias (4).docx'
ARTIFACT_DIR = r'C:/Users/karen/.gemini/antigravity/brain/a33b5b22-c89b-4ff3-a6cc-733a0c69ceac'

print("=== STARTING REQUIREMENTS EXPANSION AND SATELLITE MOCKUP INTEGRATION ===")

# Copy new satellite mockup to artifact directory
shutil.copyfile('scratch/mockup_vista_mapa_satelital_rutas.png', os.path.join(ARTIFACT_DIR, 'mockup_vista_mapa_satelital_rutas.png'))
print("Copied satellite mockup to artifact dir.")

doc = docx.Document(DOC_ADVANCE)

# 1. Define complete 83 Figures metadata
FIGURES_METADATA = [
    # Cap I & II (Figures 1 to 26)
    (1, "Ubicación de la empresa GPON TELECOM S.A de C.V", "8"),
    (2, "Organigrama de la empresa GPON TELECOM S.A de C.V.", "9"),
    (3, "Modelo de referencia OSI de 7 capas frente a la arquitectura TCP/IP.", "13"),
    (4, "Principio físico de propagación: Refracción, ángulo crítico y reflexión interna total (TIR).", "15"),
    (5, "Protocolo UDP.", "17"),
    (6, "Paradigma de FTT existente actualmente.", "19"),
    (7, "Espectro y plan de longitudes de onda en GPON (ITU-T G.984.2 WDM).", "21"),
    (8, "Arquitectura de superposición de capas en un Sistema de Información Geográfica (GIS).", "26"),
    (9, "Librería Leaflet para mapas en react.", "28"),
    (10, "Base de datos PostgreSQL.", "30"),
    (11, "¿Qué es un ORM?", "32"),
    (12, "Logo de Sequelize.js", "32"),
    (13, "Arquitectura Cliente-Servidor.", "37"),
    (14, "Integraciones y entregas continuas.", "55"),
    (15, "Metodología de cascada.", "57"),
    (16, "Ejemplo de diagrama de casos de uso.", "59"),
    (17, "Arquitectura Modelo-Vista-Controlador.", "61"),
    (18, "Ejemplo de Responsive Web Design y Mobile-First Web Design.", "64"),
    (19, "Interfaz de Adobe Color.", "66"),
    (20, "Tipos de modelos de caja (Caja negra y Caja blanca).", "68"),
    (21, "Pruebas de API REST.", "69"),
    (22, "Pasos para realizar un test de usabilidad (ejemplo).", "70"),
    (23, "Diferencia en BSS y OSS.", "71"),
    (24, "Ejemplo de sistema de inventario de redes de telecomunicaciones.", "72"),
    (25, "Cuestionario en Google Forms sobre redes GPON/FTTx para el levantamiento de requerimientos.", "76"),
    (26, "Modelado de actores y sus funciones.", "78"),

    # Cap III - Requerimientos & Diagramas de Casos de Uso (27 to 30)
    (27, "Diagrama de casos de uso - Módulo 1: Seguridad, autenticación y control de acceso (RBAC).", "79"),
    (28, "Diagrama de casos de uso - Módulo 2: Cartografía GIS, trazado troncal y cajas terminales NAP.", "80"),
    (29, "Diagrama de casos de uso - Módulo 3: Operación de puertos físicos, concurrencia ACID y abonados.", "81"),
    (30, "Diagrama de casos de uso - Módulo 4: Operación móvil Offline-First y generación de reportes ejecutivos PDF.", "82"),

    # Cap III - Robustez, Flujo, Actividades (31 to 33)
    (31, "Diagrama de robustez (V-O-C de Jacobson) para la asignación concurrente de puertos y control transaccional.", "84"),
    (32, "Diagrama de flujo general de operación y procesos del sistema GPON.", "89"),
    (33, "Diagrama de actividades UML (Swimlanes) para el proceso de asignación y provisión concurrente de puertos.", "91"),

    # Cap III - Topología, Clases, DER (34 to 37)
    (34, "Cadena de distribución de la red GPON desde la cabecera central (NOC) hasta la acometida domiciliaria (ONT).", "93"),
    (35, "Topología jerárquica de la red óptica GPON/FTTx y matriz de presupuesto óptico de potencia (ITU-T G.984.2).", "95"),
    (36, "Diagrama de clases del dominio y entidades del sistema bajo estándar UML 2.5.", "97"),
    (37, "Diagrama Entidad-Relación (DER) conceptual con notación Chen del sistema de inventario GPON.", "99"),

    # Cap III - Secuencia concurrencia, FSM estados, Secuencia offline (38 to 40)
    (38, "Diagrama de secuencia transaccional de concurrencia con SELECT ... FOR UPDATE.", "112"),
    (39, "Diagrama de máquina de estados finitos (FSM) del ciclo de vida del puerto óptico SC-APC.", "113"),
    (40, "Diagrama de secuencia UML para sincronización diferida y arquitectura móvil Offline-First.", "114"),

    # Cap III - SECTION 3.6.1 Maquetado Wireframes Figma (41)
    (41, "Maquetado de interfaces (Wireframes) y arquitectura visual del sistema en Figma.", "115"),

    # Cap III - SECTION 3.6.1 Vistas de Maquetado (42 to 46)
    (42, "Maquetado de interfaz — Visor cartográfico GIS interactivo y modal transaccional de inspección de puertos NAP.", "116"),
    (43, "Maquetado de interfaz — Visor cartográfico GIS en capa satelital con telemetría de ruta troncal y semaforización de cajas NAP.", "117"),
    (44, "Maquetado de interfaz — Padrón consolidado de abonados FTTx y gestión de expedientes técnicos.", "118"),
    (45, "Maquetado de interfaz — Consola analítica de reportes ejecutivos e indicadores de saturación de red GPON.", "119"),
    (46, "Maquetado de interfaz — Módulo de gestión centralizada de personal operativo y control de acceso (RBAC).", "120"),

    # Cap III - SECTION 3.6.1 Diagramas de Navegación, IA, Paquetes (47, 48, 49)
    (47, "Diagrama de navegación del sistema y flujo heurístico de interfaces de usuario.", "121"),
    (48, "Diagrama jerárquico de arquitectura de información (IA) del sistema web y móvil FTTx.", "122"),
    (49, "Diagrama de paquetes y componentes de software bajo estándar UML.", "123"),

    # Cap III - SECTION 3.6.2 Pruebas de Contraste Adobe Color (50 to 58)
    (50, "Prueba de colores #FFFFFF y #4F46ES", "124"),
    (51, "Prueba de colores para Botones Soporte / Técnico", "124"),
    (52, "Pruebas de color para Badge \"Datos de Prueba\"", "125"),
    (53, "Contraste para Pestaña \"Mapa de Red\" (Activa)", "125"),
    (54, "Contraste de pestañas \"Abonados\" / \"Reportes\"", "125"),
    (55, "Contraste de Botón \"APK\"", "126"),
    (56, "Contraste de color Tag Rol \"ADMIN\"", "126"),
    (57, "Contraste de Botón \"+ Troncal / Ramal\"", "127"),
    (58, "Contraste de Botón \"+ Mufa\" (Empalme)", "127"),

    # Cap IV - Implementación y Desarrollo de Código (59 to 66)
    (59, "Configuración de la conexión a PostgreSQL con Sequelize (database.ts).", "129"),
    (60, "Modelo Client.ts y estructura de la capa de modelos (backend/src/models/).", "130"),
    (61, "Middleware de autenticación con verificación de token JWT (auth.ts).", "131"),
    (62, "Controlador de asignación concurrente de puertos con transacción pesimista ACID (portController.ts).", "132"),
    (63, "Definición de rutas y endpoints de la API REST con middlewares de seguridad y RBAC (api.ts).", "133"),
    (64, "Interceptor de Axios para la inyección del token JWT en el frontend (client.ts).", "135"),
    (65, "Inicialización del servidor Express, configuración CORS y conexión Sequelize (index.ts).", "137"),
    (66, "Cliente HTTP Axios con interceptor de autorización Bearer JWT y manejo de sesión (client.ts).", "139"),

    # Cap IV - Cartografía Leaflet y Estados de Puertos (67 to 76)
    (67, "Marcador de caja NAP con ocupación inferior al 80% (NAP-SJR-26, 0/8).", "141"),
    (68, "Marcador de caja NAP en umbral preventivo, con ocupación entre 80% y 99% (NAP-SJR-01, 14/16).", "141"),
    (69, "Marcador de caja NAP saturada, con ocupación del 100% (NAP-SJR-02, 16/16).", "141"),
    (70, "Ventana emergente con el detalle de un cable troncal trazado como polilínea en el mapa.", "142"),
    (71, "Modal de captura de coordenadas GPS en campo (GpsCaptureModal.tsx).", "143"),
    (72, "Puertos de la matriz NAP en estado Libre.", "143"),
    (73, "Puertos de la matriz NAP en estado Ocupado.", "144"),
    (74, "Puerto de la matriz NAP en estado Dañado.", "144"),
    (75, "Puerto de la matriz NAP en estado Reservado.", "144"),
    (76, "Puertos libres y ocupados.", "145"),

    # Cap IV - Frontend, Offline PWA, Reportes PDF, Docker (77 to 83)
    (77, "Arquitectura de componentes frontend React y visor cartográfico.", "145"),
    (78, "Flujo de decisión y sincronización diferida de la arquitectura móvil Offline-First.", "148"),
    (79, "Esquema de persistencia local IndexedDB con Dexie.js para trabajo sin conexión (offlineDb.ts).", "149"),
    (80, "Encabezado institucional y resumen del reporte ejecutivo de auditoría de red en PDF.", "150"),
    (81, "Inventario y nivel de saturación por caja NAP en el reporte ejecutivo en PDF.", "151"),
    (82, "Flujo de generación de reportes técnicos ejecutivos en PDF mediante streaming en memoria.", "152"),
    (83, "Arquitectura de contenerización multicontenedor con Docker Compose y red aislada.", "154"),
]

assert len(FIGURES_METADATA) == 83, f"Expected 83 figures, got {len(FIGURES_METADATA)}"

# ==============================================================================
# PART 1: EXPAND SECTION 3.2 WITH COMPLETE REQUIREMENTS TAXONOMY
# ==============================================================================
print("\n=== EXPANDING SECTION 3.2 WITH COMPLETE REQUIREMENTS TAXONOMY ===")

def create_heading3(doc, text):
    p = doc.add_paragraph()
    p.style = 'Heading 3'
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(12)
    r.bold = True
    return p

def create_body_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Arial'
        r_b.font.size = Pt(11)
        r_b.bold = True
    r_t = p.add_run(text)
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(11)
    return p

# Locate insertion point in Section 3.2: after 3.2.4 RNF and before Matriz de Trazabilidad
p_target_matriz = None
p_target_flujo = None
p_target_activ = None

for i, p in enumerate(doc.paragraphs[117:], start=117):
    if '3.2.5 Matriz de Trazabilidad' in p.text:
        p_target_matriz = p
    elif '3.2.6 Diagrama de Flujo General' in p.text:
        p_target_flujo = p
    elif '3.2.7 Diagrama de Actividades' in p.text:
        p_target_activ = p
        break

print(f"Target Matriz: {p_target_matriz.text[:50]}")

# Update subsequent subsection headings
if p_target_matriz:
    p_target_matriz.text = "3.2.11 Matriz de Trazabilidad de Permisos por Rol y Códigos HTTP"
    p_target_matriz.style = 'Heading 3'
    for r in p_target_matriz.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(12)
        r.bold = True

if p_target_flujo:
    p_target_flujo.text = "3.2.12 Diagrama de Flujo General de Operación y Procesos del Sistema"
    p_target_flujo.style = 'Heading 3'
    for r in p_target_flujo.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(12)
        r.bold = True

if p_target_activ:
    p_target_activ.text = "3.2.13 Diagrama de Actividades UML para Asignación y Provisión Concurrente de Puertos"
    p_target_activ.style = 'Heading 3'
    for r in p_target_activ.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(12)
        r.bold = True

# Build new requirement paragraphs
new_req_paras = []

# 3.2.5 UI/UX Design Requirements
h_325 = create_heading3(doc, "3.2.5 Requerimientos de Diseño e Interfaz de Usuario (UI/UX)")
new_req_paras.append(h_325)
p_intro_325 = create_body_para(
    doc,
    "El diseño de la experiencia de usuario (UX) y de las interfaces gráficas (UI) se fundamenta en el estándar internacional "
    "ISO 9241-210 (Human-centred design for interactive systems) y en las pautas de accesibilidad para el contenido web WCAG 2.1 nivel AA, "
    "adaptadas a las condiciones operativas de las cuadrillas en campo en el municipio de San José del Rincón:"
)
new_req_paras.append(p_intro_325)
new_req_paras.append(create_body_para(doc, "La interfaz debe implementar una relación de contraste cromático mínima de 4.5:1 para texto normal y de 3.0:1 para elementos de control gráfico contra el fondo, garantizando total legibilidad bajo condiciones de radiación solar directa en campo abierto.", "1. RDI-01: Ergonomía Visual y Accesibilidad Bajo Luz Solar: "))
new_req_paras.append(create_body_para(doc, "Todos los elementos interactivos, botones de acción, marcadores cartográficos y conectores de puerto deben presentar un área táctil mínima de 48×48 píxeles con separación perimetral de al menos 8 px, facultando la pulsación precisa con una sola mano por técnicos que utilizan guantes de protección industrial.", "2. RDI-02: Dimensiones de Controles Táctiles (Touch Targets): "))
new_req_paras.append(create_body_para(doc, "El estado operativo de las terminales ópticas NAP debe expresarse inequívocamente mediante tres estados visuales: Verde pulsante ('#10B981') para holgura operativa (<80%), Ámbar preventivo ('#F59E0B') para umbral de alerta (80-99%) y Rojo de bloqueo ('#EF4444') para saturación total (100%), complementado con etiquetas textuales y numéricas de ocupación (e.g. '14/16').", "3. RDI-03: Semaforización Cromática Tri-Estado Intuitiva: "))
new_req_paras.append(create_body_para(doc, "La vista de inspección de caja NAP debe proyectar una réplica isomórfica 2×8 fiel a la distribución física del chasis interno de 16 puertos SC-APC, mostrando diodos LED emulados con grabado numérico secuencial del 1 al 16 en cada conector.", "4. RDI-04: Matriz Isomórfica de Chasis de Puertos: "))
new_req_paras.append(create_body_para(doc, "El diseño visual debe estructurarse bajo el principio Mobile-First con Tailwind CSS, adaptando formularios, cuadrículas de abonados y herramientas cartográficas a pantallas táctiles verticales de teléfonos inteligentes (5 a 6.7 pulgadas) con barras de navegación ergonómicas al alcance del pulgar.", "5. RDI-05: Arquitectura Responsiva Mobile-First: "))
new_req_paras.append(create_body_para(doc, "El visor cartográfico debe prevenir la sobrecarga cognitiva del operador aplicando renderizado progresivo y agrupamiento de capas, desplegando los datos técnicos detallados únicamente a través de modales contextuales activados bajo demanda.", "6. RDI-06: Densidad de Información y Jerarquía Visual Geoespacial: "))

# 3.2.6 Hardware Requirements
h_326 = create_heading3(doc, "3.2.6 Requerimientos de Hardware y Entorno Físico de Despliegue")
new_req_paras.append(h_326)
p_intro_326 = create_body_para(
    doc,
    "Para garantizar la compatibilidad operativa de la plataforma tanto en la planta externa como en el centro de datos corporativo, se establecen los siguientes requerimientos de hardware:"
)
new_req_paras.append(p_intro_326)
new_req_paras.append(create_body_para(doc, "Los teléfonos inteligentes operados por las cuadrillas de campo deben contar con procesador de arquitectura ARM de 64 bits (octa-core a 2.0 GHz mínimo), memoria RAM mínima de 3 GB, almacenamiento libre de 4 GB y pantalla táctil capacitiva con brillo de al menos 450 nits.", "1. RHW-01: Dispositivos Móviles de Cuadrilla (Clientes de Campo): "))
new_req_paras.append(create_body_para(doc, "El hardware móvil debe integrar receptor satelital multi-constelación (A-GPS, GLONASS y Galileo) con precisión geodésica inferior a 5 metros bajo cielo abierto, permitiendo la captura fehaciente de coordenadas geográficas (WGS84) al pie de poste.", "2. RHW-02: Sensor de Posicionamiento Global Satelital (GPS): "))
new_req_paras.append(create_body_para(doc, "La infraestructura centralizada que aloja el backend REST y la base de datos PostgreSQL debe contar con arquitectura x86_64, mínimo 4 núcleos virtuales de CPU (vCPU), 8 GB de memoria RAM DDR4 ECC y 100 GB de almacenamiento en disco de estado sólido NVMe en arreglo redundante.", "3. RHW-03: Servidor de Producción Central (Host de Servicios y Base de Datos): "))

# 3.2.7 Software Platform Requirements
h_327 = create_heading3(doc, "3.2.7 Requerimientos de Software y Plataforma Tecnológica")
new_req_paras.append(h_327)
p_intro_327 = create_body_para(
    doc,
    "La compatibilidad del software se define para asegurar interoperabilidad fluida, mantenibilidad y modularidad de desarrollo:"
)
new_req_paras.append(p_intro_327)
new_req_paras.append(create_body_para(doc, "La plataforma debe ejecutarse sin requerir complementos propietarios en navegadores compatibles con ECMAScript 2022+: Google Chrome (v100+), Mozilla Firefox (v100+), Apple Safari (v15+) y Microsoft Edge (v100+).", "1. RSW-01: Motores de Navegación Web Homologados: "))
new_req_paras.append(create_body_para(doc, "El entorno cliente debe registrar un Service Worker que gestione la Cache Storage API y un Web App Manifest para permitir la instalación de la plataforma como aplicación web progresiva (PWA) en el dispositivo móvil.", "2. RSW-02: Soporte para Progressive Web App (PWA): "))
new_req_paras.append(create_body_para(doc, "El servidor debe ejecutarse sobre el entorno Node.js en su versión 18.x o 20.x LTS, estructurado bajo el framework Express 4.x y el compilador TypeScript 5.x configurado con directivas de tipado estricto.", "3. RSW-03: Entorno de Ejecución Backend: "))
new_req_paras.append(create_body_para(doc, "El gestor de base de datos relacional debe ser PostgreSQL versión 16+, con soporte nativo para transacciones concurrentes ACID, operadores JSONB y extensiones criptográficas integradas.", "4. RSW-04: Motor de Base de Datos Relacional: "))

# 3.2.8 Network and Communications Requirements
h_328 = create_heading3(doc, "3.2.8 Requerimientos de Red, Comunicaciones y Conectividad")
new_req_paras.append(h_328)
p_intro_328 = create_body_para(
    doc,
    "Los requerimientos de red regulan la transferencia de información y la resiliencia en zonas rurales con conectividad intermitente:"
)
new_req_paras.append(p_intro_328)
new_req_paras.append(create_body_para(doc, "La totalidad de las comunicaciones cliente-servidor debe cifrarse obligatoriamente mediante HTTPS sobre TLS 1.3 con certificados X.509 validados, bloqueando cualquier intento de conexión bajo texto plano HTTP por el puerto 80.", "1. RCOM-01: Cifrado Forzoso de Canal de Transporte: "))
new_req_paras.append(create_body_para(doc, "Los endpoints de la API REST deben transferir exclusivamente datos en formato JSON compacto con compresión HTTP activa (GZIP o Brotli), limitando el tamaño promedio de los payloads a menos de 50 KB para optimizar el consumo de datos móviles en campo.", "2. RCOM-02: Optimización de Ancho de Banda y Cargas Útiles: "))
new_req_paras.append(create_body_para(doc, "La aplicación debe monitorizar dinámicamente el estado de la conexión mediante eventos de red, alternando de manera transparente a persistencia local en Dexie.js (IndexedDB) ante desconexión y encolando transacciones para su sincronización diferida al reanudar enlace.", "3. RCOM-03: Resiliencia ante Desconexión Celular: "))

# 3.2.9 Data Integrity and Database Requirements
h_329 = create_heading3(doc, "3.2.9 Requerimientos de Datos e Integridad Transaccional")
new_req_paras.append(h_329)
p_intro_329 = create_body_para(
    doc,
    "La gobernanza y consistencia de los datos del inventario óptico se rige por las siguientes restricciones técnicas:"
)
new_req_paras.append(p_intro_329)
new_req_paras.append(create_body_para(doc, "Toda mutación del estado de un puerto óptico debe procesarse dentro de una transacción gestionada bajo nivel de aislamiento Read Committed, aplicando bloqueo pesimista a nivel de fila mediante 'SELECT ... FOR UPDATE' para eliminar colisiones concurrentes.", "1. RDAT-01: Integridad Transaccional ACID y Bloqueo Pesimista: "))
new_req_paras.append(create_body_para(doc, "El esquema relacional debe cumplir estrictamente con la Tercera Forma Normal (3FN), eliminando dependencias transitivas y anomalías de actualización mediante claves foráneas y normalización de catálogos.", "2. RDAT-02: Normalización en Tercera Forma Normal (3FN): "))
new_req_paras.append(create_body_para(doc, "Se deben programar restricciones de clave foránea con cláusula 'ON DELETE RESTRICT' en elementos padre (cajas NAP y splitters), prohibiendo la eliminación física de cualquier nodo que mantenga puertos ocupados por clientes con servicio activo.", "3. RDAT-03: Restricciones de Integridad Referencial: "))
new_req_paras.append(create_body_para(doc, "Cada asignación, liberación o cambio de estado de puerto debe registrar de forma automática e inmutable el identificador del usuario, dirección IP, fecha/hora UTC y valores previos en la tabla de bitácora transaccional.", "4. RDAT-04: Registro Inmutable de Auditoría: "))

# 3.2.10 Organizational and Regulatory Requirements
h_3210 = create_heading3(doc, "3.2.10 Requerimientos Organizacionales, de Operación y Normativos")
new_req_paras.append(h_3210)
p_intro_3210 = create_body_para(
    doc,
    "Requisitos vinculados con los procesos humanos de la empresa y la regulación técnica del sector de telecomunicaciones:"
)
new_req_paras.append(p_intro_3210)
new_req_paras.append(create_body_para(doc, "La interfaz debe ser lo suficientemente intuitiva para que un técnico instalador de cuadrilla aprenda a consultar una caja NAP y registrar una nueva acometida en un periodo de capacitación menor a dos horas.", "1. ROPR-01: Curva de Aprendizaje y Usabilidad Operativa: "))
new_req_paras.append(create_body_para(doc, "Los cálculos de atenuación y el presupuesto de potencia óptica deben regirse bajo los parámetros técnicos de las recomendaciones ITU-T G.984.1 y G.984.2 para redes GPON Clase B+ (margen admisible de pérdida entre 13 dB y 28 dB).", "2. ROPR-02: Apego a Estándares de Redes GPON (ITU-T): "))
new_req_paras.append(create_body_para(doc, "El sistema debe desplegar la nomenclatura de los 12 hilos ópticos del cable de distribución conforme al código cromático normalizado TIA/EIA-598-A (Azul, Naranja, Verde, Marrón, Gris, Blanco, Rojo, Negro, Amarillo, Violeta, Rosa y Aqua), eliminando confusiones en campo.", "3. ROPR-03: Código de Colores de Fibra Óptica (TIA/EIA-598-A): "))
new_req_paras.append(create_body_para(doc, "El sistema debe proveer la generación inmediata de reportes de auditoría técnica en PDF con encabezados institucionales y firmas de responsabilidad técnica para soportar inspecciones de calidad y control de inventario patrimonial.", "4. ROPR-04: Trazabilidad y Dictámenes Formales en PDF: "))

# Insert all new requirement paragraphs before p_target_matriz
target_elem = p_target_matriz._p
for para in new_req_paras:
    target_elem.addprevious(para._p)

print("Inserted all 6 new Requirement Subsections (3.2.5 to 3.2.10) successfully!")

# ==============================================================================
# PART 2: INSERT SATELLITE MOCKUP (FIGURA 43) INTO SECTION 3.6.1
# ==============================================================================
print("\n=== INSERTING SATELLITE MOCKUP INTO SECTION 3.6.1 ===")

# Find Figura 42 note paragraph in Section 3.6.1
p_fig42_note = None
for i, p in enumerate(doc.paragraphs[117:], start=117):
    if 'Figura 42.' in p.text:
        # Next paragraph is the Elaboración propia note
        p_fig42_note = doc.paragraphs[i+1]
        break

print(f"Found Figura 42 note: {p_fig42_note.text}")

# Detailed explanation text explaining every rectangle, badge, button, and marker
desc_satellite = (
    "Para complementar la visualización general de planta externa y dotar a las cuadrillas de una herramienta de alta "
    "precisión sobre el terreno, se maquetó la vista del Visor Cartográfico GIS en modo satelital de alta resolución. "
    "La interfaz integra múltiples componentes visuales e interactivos diseñados para facilitar la orientación espacial y la "
    "supervisión del despliegue físico de la red de fibra óptica en San José del Rincón, estructurados formalmente de la siguiente manera:\n"
    "1. Controles de Zoom y Navegación Escalar ('[+]' y '[-]'): Ubicados en el rectángulo superior izquierdo de la pantalla, "
    "permiten al operador aproximar milimétricamente la escena visual para inspeccionar postes individuales o alejar la perspectiva "
    "para abarcar corredores viales completos.\n"
    "2. Insignia / Badge de Cobertura Total ('Capas de Red 352.51 km'): Rectángulo flotante con bordes redondeados, fondo azul claro "
    "e icono de visibilidad que cuantifica en tiempo real la longitud lineal acumulada de todos los cables de fibra óptica (troncales y "
    "ramales secundarios) registrados activamente en el inventario del municipio.\n"
    "3. Píldora Flotante de Telemetría de Trazo Activo ('ruta manzan... 17.62 km (17,620.9 ML)'): Rectángulo interactivo centrado en la "
    "cabecera que reporta la distancia geodésica del segmento de cable troncal seleccionado en la Manzana actual (17.62 kilómetros equivalentes "
    "a 17,620.9 Metros Lineales - ML). Incorpora el botón rojo de acción rápida 'Eliminar' (con icono de cesto de basura) para suprimir o "
    "recalibrar la polilínea del tendido, así como un control de cierre rápido '[X]'.\n"
    "4. Selector Desplegable de Capas Base ('Satélite v'): Rectángulo de control ubicado en el margen superior derecho que habilita la "
    "conmutación instantánea entre la ortofoto fotográfica satelital aérea y la cartografía vectorial de calles de OpenStreetMap.\n"
    "5. Traza Vectorial de Cable Troncal (Polilínea continua de color púrpura): Representación gráfica de la trayectoria geodésica del cable "
    "de fibra óptica dieléctrico autosoportado (ADSS) tendido sobre la postería municipal.\n"
    "6. Marcadores Cromáticos Triangulares de Cajas Terminales NAP:\n"
    "   - Marcador Triangular Verde (Vértice inferior izquierdo): Identifica una caja NAP en estado operativo óptimo con ocupación inferior al "
    "80%, indicando disponibilidad inmediata de puertos SC-APC libres para nuevas contrataciones domiciliarias.\n"
    "   - Marcador Triangular Rojo (Vértice central): Identifica una caja NAP en condición de saturación crítica al 100% de capacidad (16 de 16 "
    "puertos ocupados), notificando de forma visual restrictiva la prohibición de conectar nuevos clientes en dicho punto sin previo tendido de ampliación.\n"
    "   - Marcador Triangular Amarillo / Ámbar (Vértice superior derecho): Identifica una caja NAP en umbral preventivo (80% a 99% de ocupación), "
    "alertando a la jefatura técnica sobre la necesidad próxima de instalar un splitter secundario adicional o derivar clientes hacia un cierre adyacente."
)

p_sat_desc = doc.add_paragraph()
p_sat_desc.style = 'Normal'
p_sat_desc.paragraph_format.line_spacing = 1.5
p_sat_desc.paragraph_format.space_before = Pt(6)
p_sat_desc.paragraph_format.space_after = Pt(6)
r_sat_d = p_sat_desc.add_run(desc_satellite)
r_sat_d.font.name = 'Arial'
r_sat_d.font.size = Pt(11)

p_sat_img = doc.add_paragraph()
p_sat_img.style = 'Normal'
p_sat_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sat_img.paragraph_format.space_before = Pt(6)
p_sat_img.paragraph_format.space_after = Pt(4)
r_sat_i = p_sat_img.add_run()
r_sat_i.add_picture('scratch/mockup_vista_mapa_satelital_rutas.png', width=Inches(6.05))

p_sat_cap = doc.add_paragraph()
p_sat_cap.style = 'Figuras'
p_sat_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sat_cap.paragraph_format.space_before = Pt(4)
p_sat_cap.paragraph_format.space_after = Pt(2)
r_sat_c = p_sat_cap.add_run("FIGURA_DUMMY_43")
r_sat_c.font.name = 'Arial'
r_sat_c.font.size = Pt(9)
r_sat_c.bold = True

p_sat_note = doc.add_paragraph()
p_sat_note.style = 'Normal'
p_sat_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sat_note.paragraph_format.space_before = Pt(0)
p_sat_note.paragraph_format.space_after = Pt(12)
r_sat_n = p_sat_note.add_run("Elaboración propia.")
r_sat_n.font.name = 'Arial'
r_sat_n.font.size = Pt(9)
r_sat_n.bold = True

# Insert sequentially right after p_fig42_note
curr_el = p_fig42_note._p
for el in [p_sat_desc, p_sat_img, p_sat_cap, p_sat_note]:
    curr_el.addnext(el._p)
    curr_el = el._p

print("Inserted Satellite Mockup and explanation successfully right after Figura 42!")

# ==============================================================================
# PART 3: RENUMBER ALL 83 FIGURES IN THE BODY
# ==============================================================================
print("\n=== IDENTIFYING AND RENUMBERING ALL 83 FIGURES IN BODY ===")
body_after_intro = doc.paragraphs[117:]
figure_captions = []
for p in body_after_intro:
    txt = p.text.strip()
    if p.style.name == 'Figuras' and '\t' not in txt:
        figure_captions.append(p)
    elif txt.startswith("FIGURA_DUMMY_"):
        figure_captions.append(p)

print(f"Total figure captions detected in body: {len(figure_captions)}")
if len(figure_captions) != 83:
    print(f"ERROR: Expected 83 captions, found {len(figure_captions)}")
    for i, p in enumerate(figure_captions):
        print(f"  [{i}]: {p.text[:60]}")
    raise RuntimeError("Figure caption count mismatch!")

for k, p_cap in enumerate(figure_captions):
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_id = str(700 + k)
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    # Clean paragraph XML
    p_elem = p_cap._p
    pPr = p_elem.find(qn('w:pPr'))
    for child in list(p_elem):
        if child != pPr:
            p_elem.remove(child)
            
    p_cap.style = 'Figuras'
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(2)
    
    bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{bookmark_name}"/>')
    p_elem.append(bm_start)
    
    r_elem = parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>'
        f'<w:b/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>'
        f'<w:t xml:space="preserve">{new_caption_text}</w:t></w:r>'
    )
    p_elem.append(r_elem)
    
    bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>')
    p_elem.append(bm_end)

print("All 83 figure captions successfully renumbered and bookmarked in body!")

# ==============================================================================
# PART 4: SYNCHRONIZE ÍNDICE DE FIGURAS (83 ENTRIES)
# ==============================================================================
print("\n=== SYNCHRONIZING ÍNDICE DE FIGURAS ===")
idx_title_p = None
idx_end_p = None
for i, p in enumerate(doc.paragraphs[:130]):
    if 'ndice de figuras' in p.text:
        idx_title_p = i
    if 'ndice de tablas' in p.text:
        idx_end_p = i
        break

existing_toc_p = []
for i in range(idx_title_p + 1, idx_end_p):
    p = doc.paragraphs[i]
    if p.text.strip().startswith("Figura "):
        existing_toc_p.append(p)

print(f"Found {len(existing_toc_p)} existing TOC entries. Updating to 83...")

# Update existing 82
for k in range(min(len(existing_toc_p), 83)):
    p = existing_toc_p[k]
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    hl = p._p.xpath('.//w:hyperlink')
    if hl:
        hl[0].attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor'] = bookmark_name
        
    instr = p._p.xpath('.//w:instrText')
    for inst in instr:
        if 'PAGEREF' in inst.text:
            inst.text = f" PAGEREF {bookmark_name} \\h "
            
    t_nodes = p._p.xpath('.//w:hyperlink//w:t')
    if len(t_nodes) >= 2:
        t_nodes[0].text = new_caption_text
        t_nodes[-1].text = fig_page
    elif len(t_nodes) == 1:
        t_nodes[0].text = new_caption_text

# Append 83rd entry if needed
if len(existing_toc_p) < 83:
    last_toc_p = existing_toc_p[-1]
    for k in range(len(existing_toc_p), 83):
        fig_num, fig_title, fig_page = FIGURES_METADATA[k]
        new_caption_text = f"Figura {fig_num}. {fig_title}"
        bookmark_name = f"_Toc240964{71 + k:03d}"
        
        new_p_xml = copy.deepcopy(last_toc_p._p)
        hl = new_p_xml.xpath('.//w:hyperlink')
        if hl:
            hl[0].attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor'] = bookmark_name
        instr = new_p_xml.xpath('.//w:instrText')
        for inst in instr:
            if 'PAGEREF' in inst.text:
                inst.text = f" PAGEREF {bookmark_name} \\h "
        t_nodes = new_p_xml.xpath('.//w:hyperlink//w:t')
        if len(t_nodes) >= 2:
            t_nodes[0].text = new_caption_text
            t_nodes[-1].text = fig_page
            
        last_toc_p._p.addnext(new_p_xml)
        last_toc_p = docx.text.paragraph.Paragraph(new_p_xml, doc)

print("Successfully synchronized all 83 entries in Índice de figuras!")

# ==============================================================================
# PART 5: SAVE DOCUMENTS AND VERIFY
# ==============================================================================
print("\n=== SAVING DOCUMENTS ===")
doc.save(DOC_ADVANCE)
print(f"Saved: {DOC_ADVANCE}")

doc.save(DOC_MIRROR)
print(f"Saved: {DOC_MIRROR}")

print("\n=== VERIFICATION AUDIT ===")
for path in [DOC_ADVANCE, DOC_MIRROR]:
    print(f"Auditing {path}...")
    doc_chk = docx.Document(path)
    b_figs = [p.text.strip() for p in doc_chk.paragraphs[117:] if p.style.name == 'Figuras' and '\t' not in p.text]
    print(f"  Body figures: {len(b_figs)}")
    assert len(b_figs) == 83, f"Expected 83 body figures, found {len(b_figs)}"
    for idx, t in enumerate(b_figs):
        assert t.startswith(f"Figura {idx+1}."), f"Mismatch at Fig {idx+1}: {t}"
    
    # Check TOC
    toc_f = [p.text.strip() for p in doc_chk.paragraphs[idx_title_p+1:idx_title_p+85] if p.text.strip().startswith("Figura ")]
    print(f"  TOC figures: {len(toc_f)}")
    assert len(toc_f) == 83, f"Expected 83 TOC figures, found {len(toc_f)}"
    
    # Check tables
    b_tbls = [p.text.strip() for p in doc_chk.paragraphs[117:] if p.style.name == 'Tablas' and '\t' not in p.text]
    print(f"  Tables: {len(b_tbls)}")
    assert len(b_tbls) == 26, f"Expected 26 tables, found {len(b_tbls)}"

print("\n=== ALL TASKS COMPLETED SUCCESSFULLY! ===")

