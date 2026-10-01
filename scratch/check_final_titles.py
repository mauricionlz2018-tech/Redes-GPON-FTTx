import pptx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(os.path.join('docs', 'PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx'))
print(f"Slide count: {len(prs.slides)}")
for i, s in enumerate(prs.slides):
    t_shape = None
    for shp in s.shapes:
        if shp.has_text_frame and any(k in shp.text for k in ['Agenda', 'Empresa', 'Introducción', 'Planteamiento', 'Justificación', 'Objetivos', 'Herramientas', 'Evidencias', 'Cronograma', 'Conclusiones', 'Bibliografía', 'Gracias']):
            t_shape = shp.text.strip().replace('\n', ' ')
            break
    print(f"Slide {i+1}: title=\"{t_shape}\"")
