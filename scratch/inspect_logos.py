import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

print("=== BUSCANDO LOGOS EN ENCABEZADOS Y CUERPO ===")
for i, s in enumerate(doc.sections):
    h = s.header
    for p in h.paragraphs:
        draws = len(p._element.xpath('.//w:drawing'))
        if 'ESTADO DE' in p.text or 'BICENTENARIO' in p.text or draws > 0:
            print(f"Header en Sec {i}: text={repr(p.text)} drawings={draws}")
    for t in h.tables:
        print(f"Tabla en header de Sec {i}:")
        for r in t.rows:
            for c in r.cells:
                print("   cell:", repr(c.text.strip()))

for i in range(25):
    p = doc.paragraphs[i]
    draws = len(p._element.xpath('.//w:drawing'))
    if 'ESTADO DE' in p.text or 'BICENTENARIO' in p.text or draws > 0 or p.text.strip():
        print(f"Body P{i}: text={repr(p.text[:40])} drawings={draws}")

for t_idx, t in enumerate(doc.tables[:4]):
    print(f"Body Tabla {t_idx}:")
    for r in t.rows:
        for c in r.cells:
            if c.text.strip():
                print(f"   cell: {repr(c.text.strip()[:40])}")

