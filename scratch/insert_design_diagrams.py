import os
import re
import copy
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_ADVANCE = 'docs/Documentacion_Residencias_avance.docx'
DOC_MIRROR = 'docs/Documentacion_Residencias (4).docx'

print("=== STARTING DESIGN DIAGRAMS INSERTION AND RENUMBERING ===")

doc = docx.Document(DOC_ADVANCE)

# 1. Define the 73 figures with their exact title, and estimated page number
FIGURES_METADATA = [
    # Cap I & II (Figures 1 to 27)
    (1, "Ubicaci\u00f3n de la empresa GPON TELECOM S.A de C.V", "8"),
    (2, "Organigrama de la empresa GPON TELECOM S.A de C.V.", "9"),
    (3, "Modelo de referencia OSI de 7 capas frente a la arquitectura TCP/IP.", "13"),
    (4, "Principio f\u00edsico de propagaci\u00f3n: Refracci\u00f3n, \u00e1ngulo cr\u00edtico y reflexi\u00f3n interna total (TIR).", "15"),
    (5, "Protocolo UDP.", "17"),
    (6, "Paradigma de FTT existente actualmente.", "19"),
    (7, "Espectro y plan de longitudes de onda en GPON (ITU-T G.984.2 WDM).", "21"),
    (8, "Arquitectura de superposici\u00f3n de capas en un Sistema de Informaci\u00f3n Geogr\u00e1fica (GIS).", "26"),
    (9, "Librer\u00eda Leaflet para mapas en react.", "28"),
    (10, "Base de datos PostgreSQL.", "30"),
    (11, "\u00bfQu\u00e9 es un ORM?", "32"),
    (12, "Logo de Sequelize.js", "32"),
    (13, "Arquitectura Cliente-Servidor.", "37"),
    (14, "Integraciones y entregas continuas.", "55"),
    (15, "Metodolog\u00eda de cascada.", "57"),
    (16, "Ejemplo de diagrama de casos de uso.", "59"),
    (17, "Arquitectura Modelo-Vista-Controlador.", "61"),
    (18, "Ejemplo de Responsive Web Design y Mobile-First Web Design.", "64"),
    (19, "Interfaz de Adobe Color.", "66"),
    (20, "Tipos de modelos de caja (Caja negra y Caja blanca).", "68"),
    (21, "Pruebas de API REST.", "69"),
    (22, "Pasos para realizar un test de usabilidad (ejemplo).", "70"),
    (23, "Diferencia en BSS y OSS.", "71"),
    (24, "Ejemplo de sistema de inventario de redes de telecomunicaciones.", "72"),
    (25, "Cuestionario en Google Forms sobre redes GPON/FTTx para el levantamiento de requerimientos.", "76"),
    (26, "Modelado de actores y sus funciones.", "78"),
    (27, "Diagrama general de casos de uso del sistema bajo est\u00e1ndar UML.", "79"),
    
    # Cap III - SECTION 3.2.2 (NEW 28)
    (28, "Diagrama de robustez (V-O-C de Jacobson) para la asignaci\u00f3n concurrente de puertos y control transaccional.", "82"),
    
    # Cap III - Existing 28 & 29 (now 29 & 30)
    (29, "Diagrama de flujo general de operaci\u00f3n y procesos del sistema GPON.", "89"),
    (30, "Cadena de distribuci\u00f3n de la red GPON desde la cabecera central (NOC) hasta la acometida domiciliaria (ONT).", "91"),
    
    # Cap III - SECTION 3.3.1 (NEW 31)
    (31, "Topolog\u00eda jer\u00e1rquica de la red \u00f3ptica GPON/FTTx y matriz de presupuesto \u00f3ptico de potencia (ITU-T G.984.2).", "93"),
    
    # Cap III - Existing 30 & 31 (now 32 & 33)
    (32, "Diagrama Entidad-Relaci\u00f3n (DER) conceptual con notaci\u00f3n Chen del sistema de inventario GPON.", "95"),
    (33, "Diagrama de secuencia transaccional de concurrencia con SELECT ... FOR UPDATE.", "112"),
    
    # Cap III - SECTION 3.5.2 (NEW 34 & 35)
    (34, "Diagrama de m\u00e1quina de estados finitos (FSM) del ciclo de vida del puerto \u00f3ptico SC-APC.", "113"),
    (35, "Diagrama de secuencia UML para sincronizaci\u00f3n diferida y arquitectura m\u00f3vil Offline-First.", "114"),
    
    # Cap III - Existing 32 (now 36)
    (36, "Maquetado de interfaces (Wireframes) y arquitectura visual del sistema en Figma.", "115"),
    
    # Cap III - SECTION 3.6.1 (NEW 37, 38, 39)
    (37, "Diagrama de navegaci\u00f3n del sistema y flujo heur\u00edstico de interfaces de usuario.", "116"),
    (38, "Diagrama jer\u00e1rquico de arquitectura de informaci\u00f3n (IA) del sistema web y m\u00f3vil FTTx.", "117"),
    (39, "Diagrama de paquetes y componentes de software bajo est\u00e1ndar UML.", "118"),
    
    # Cap III - Adobe Color tests (Old 34 to 42 -> now 40 to 48)
    (40, "Prueba de colores #FFFFFF y #4F46ES", "119"),
    (41, "Prueba de colores para Botones Soporte / T\u00e9cnico", "119"),
    (42, "Pruebas de color para Badge \"Datos de Prueba\"", "120"),
    (43, "Contraste para Pesta\u00f1a \"Mapa de Red\" (Activa)", "120"),
    (44, "Contraste de pesta\u00f1as \"Abonados\" / \"Reportes\"", "120"),
    (45, "Contraste de Bot\u00f3n \"APK\"", "121"),
    (46, "Contraste de color Tag Rol \"ADMIN\"", "121"),
    (47, "Contraste de Bot\u00f3n \"+ Troncal / Ramal\"", "122"),
    (48, "Contraste de Bot\u00f3n \"+ Mufa\" (Empalme)", "122"),
    
    # Cap IV - Coding & Architecture (Old 43 to 50 -> now 49 to 56)
    (49, "Configuraci\u00f3n de la conexi\u00f3n a PostgreSQL con Sequelize (database.ts).", "125"),
    (50, "Modelo Client.ts y estructura de la capa de modelos (backend/src/models/).", "126"),
    (51, "Middleware de autenticaci\u00f3n con verificaci\u00f3n de token JWT (auth.ts).", "127"),
    (52, "Controlador de asignaci\u00f3n concurrente de puertos con transacci\u00f3n pesimista ACID (portController.ts).", "128"),
    (53, "Definici\u00f3n de rutas y endpoints de la API REST con middlewares de seguridad y RBAC (api.ts).", "129"),
    (54, "Interceptor de Axios para la inyecci\u00f3n del token JWT en el frontend (client.ts).", "131"),
    (55, "Inicializaci\u00f3n del servidor Express, configuraci\u00f3n CORS y conexi\u00f3n Sequelize (index.ts).", "133"),
    (56, "Cliente HTTP Axios con interceptor de autorizaci\u00f3n Bearer JWT y manejo de sesi\u00f3n (client.ts).", "135"),
    
    # Cap IV - Map & Port UI (Old 51 to 59 -> now 57 to 65)
    (57, "Marcador de caja NAP con ocupaci\u00f3n inferior al 80% (NAP-SJR-26, 0/8).", "137"),
    (58, "Marcador de caja NAP en umbral preventivo, con ocupaci\u00f3n entre 80% y 99% (NAP-SJR-01, 14/16).", "137"),
    (59, "Marcador de caja NAP saturada, con ocupaci\u00f3n del 100% (NAP-SJR-02, 16/16).", "137"),
    (60, "Ventana emergente con el detalle de un cable troncal trazado como polil\u00ednea en el mapa.", "138"),
    (61, "Modal de captura de coordenadas GPS en campo (GpsCaptureModal.tsx).", "139"),
    (62, "Puertos de la matriz NAP en estado Libre.", "139"),
    (63, "Puertos de la matriz NAP en estado Ocupado.", "140"),
    (64, "Puerto de la matriz NAP en estado Da\u00f1ado.", "140"),
    (65, "Puerto de la matriz NAP en estado Reservado.", "140"),
    
    # Cap IV - UI & Offline & PDF & Docker (Old 61a, 61b, 62 to 67 -> now 66 to 73)
    (66, "Puertos libres y ocupados.", "141"),
    (67, "Arquitectura de componentes frontend React y visor cartogr\u00e1fico.", "141"),
    (68, "Flujo de decisi\u00f3n y sincronizaci\u00f3n diferida de la arquitectura m\u00f3vil Offline-First.", "144"),
    (69, "Esquema de persistencia local IndexedDB con Dexie.js para trabajo sin conexi\u00f3n (offlineDb.ts).", "145"),
    (70, "Encabezado institucional y resumen del reporte ejecutivo de auditor\u00eda de red en PDF.", "146"),
    (71, "Inventario y nivel de saturaci\u00f3n por caja NAP en el reporte ejecutivo en PDF.", "147"),
    (72, "Flujo de generaci\u00f3n de reportes t\u00e9cnicos ejecutivos en PDF mediante streaming en memoria.", "148"),
    (73, "Arquitectura de contenerizaci\u00f3n multicontenedor con Docker Compose y red aislada.", "150"),
]

assert len(FIGURES_METADATA) == 73, f"Expected 73 figures, got {len(FIGURES_METADATA)}"

# Locate targets in the body (paragraphs after index, starting at 105)
body = doc.paragraphs[105:]

target_p1 = None
for p in body:
    if '4.- API Geoespacial GPS / Leaflet' in p.text:
        target_p1 = p
        break

target_p2 = None
for i, p in enumerate(body):
    if 'Figura 29. Cadena de distribuci' in p.text:
        target_p2 = body[i+1] # The Elaboración propia
        break

target_p3 = None
for i, p in enumerate(body):
    if 'Figura 31. Diagrama de secuencia transaccional' in p.text:
        target_p3 = body[i+1] # The Elaboración propia
        break

target_p4 = None
for p in body:
    if '4.- Wireframe 4: Padr' in p.text:
        target_p4 = p
        break

print(f"Target 1: {target_p1.text[:50]}")
print(f"Target 2: {target_p2.text[:50]}")
print(f"Target 3: {target_p3.text[:50]}")
print(f"Target 4: {target_p4.text[:50]}")

def build_figure_elements(intro_text, img_path, caption_dummy, note_text="Elaboraci\u00f3n propia"):
    """Creates [intro_p, img_p, caption_p, note_p] attached to doc."""
    # 1. Intro
    p_intro = doc.add_paragraph()
    p_intro.style = 'Normal'
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(6)
    r_intro = p_intro.add_run(intro_text)
    r_intro.font.name = 'Arial'
    r_intro.font.size = Pt(11)
    
    # 2. Image
    p_img = doc.add_paragraph()
    p_img.style = 'Normal'
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(4)
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=Inches(6.05))
    
    # 3. Caption
    p_cap = doc.add_paragraph()
    p_cap.style = 'Figuras'
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(2)
    r_cap = p_cap.add_run(caption_dummy)
    r_cap.font.name = 'Arial'
    r_cap.font.size = Pt(9)
    r_cap.bold = True
    
    # 4. Note
    p_note = doc.add_paragraph()
    p_note.style = 'Normal'
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(12)
    r_note = p_note.add_run(note_text)
    r_note.font.name = 'Arial'
    r_note.font.size = Pt(9)
    r_note.bold = True
    
    return [p_intro, p_img, p_cap, p_note]

def insert_elements_after(anchor_p, elements):
    curr = anchor_p._p
    for el in elements:
        curr.addnext(el._p)
        curr = el._p
    return el # returns last paragraph

# Group 1: Robustness (after target_p1)
intro_g1 = (
    "Para tender un puente metodol\u00f3gico formal entre los casos de uso descritos y el dise\u00f1o de la arquitectura "
    "de clases y componentes del sistema, se aplic\u00f3 el An\u00e1lisis de Robustez de Jacobson (Boundary-Control-Entity). "
    "Este modelado descompone el caso de uso neur\u00e1lgico CU-08 (Asignaci\u00f3n Concurrente de Puertos en Campo vs. NOC) "
    "en tres niveles de objetos: interfaz perimetral de usuario (Boundary), coordinadores l\u00f3gicos de negocio y control "
    "transaccional (Control), y modelos persistentes de datos en PostgreSQL (Entity), formalizando la segregaci\u00f3n de "
    "responsabilidades y las barreras de contenci\u00f3n concurrente:"
)
elems_g1 = build_figure_elements(intro_g1, 'scratch/diag_robustez_asignacion.png', 'FIGURA_DUMMY_28')
insert_elements_after(target_p1, elems_g1)
print("Inserted Group 1: Robustez")

# Group 2: GPON Topology (after target_p2)
intro_g2 = (
    "Para complementar la cadena de distribuci\u00f3n f\u00edsica y correlacionar la jerarqu\u00eda de planta externa con los niveles "
    "de atenuaci\u00f3n calculados anal\u00edticamente, se dise\u00f1\u00f3 la topolog\u00eda integral de red GPON/FTTx desplegada en el municipio "
    "de San Jos\u00e9 del Rinc\u00f3n. El diagrama formaliza los seis niveles cardinales de la arquitectura (NOC OLT, ODF, Cable Troncal ADSS, "
    "Cierres de Empalme Primarios 1:4, Red de Distribuci\u00f3n con Cajas Terminales NAP 1:16 y Acometida Domiciliaria ONT), integrando la "
    "matriz normalizada del presupuesto \u00f3ptico de potencia bajo la norma ITU-T G.984.2 Clase B+:"
)
elems_g2 = build_figure_elements(intro_g2, 'scratch/diag_topologia_gpon.png', 'FIGURA_DUMMY_31')
insert_elements_after(target_p2, elems_g2)
print("Inserted Group 2: Topología GPON")

# Group 3: FSM and Offline Sequence (after target_p3)
intro_g3_a = (
    "Para complementar el modelado temporal de concurrencia y gobernar formalmente las transiciones del ciclo de vida de la "
    "infraestructura de fibra \u00f3ptica en la base de datos, se elabor\u00f3 el Diagrama de M\u00e1quina de Estados Finitos (FSM - Finite "
    "State Machine) del Puerto \u00d3ptico SC-APC. Este diagrama estipula las seis fases mutuamente excluyentes (Libre, En Proceso, Ocupado, "
    "Da\u00f1ado, En Mantenimiento y Cola Offline), definiendo con exactitud matem\u00e1tica los disparadores l\u00f3gicos, las precondiciones "
    "ACID de base de datos y las restricciones de rol (RBAC) que impiden mutaciones de estado inv\u00e1lidas o inconsistencias en campo:"
)
elems_g3_a = build_figure_elements(intro_g3_a, 'scratch/diag_estados_puerto.png', 'FIGURA_DUMMY_34')

intro_g3_b = (
    "Asimismo, para garantizar la fiabilidad del sistema en cuadrillas de campo que operan en localidades con cobertura celular nula "
    "o intermitente, se model\u00f3 el flujo temporal de persistencia local y sincronizaci\u00f3n diferida mediante un Diagrama de Secuencia "
    "UML para la Arquitectura Offline-First. En este diagrama se desglosan las interacciones as\u00edncronas entre la Interfaz PWA, el Service "
    "Worker, el almac\u00e9n local IndexedDB (Dexie.js), la Cola de Transacciones Pendientes (Offline Queue), la API REST Express y el Motor "
    "PostgreSQL Central, ilustrando las tres fases cr\u00edticas: operaci\u00f3n en desconexi\u00f3n, reconciliaci\u00f3n autom\u00e1tica con "
    "reintentos exponenciales y confirmaci\u00f3n transaccional:"
)
elems_g3_b = build_figure_elements(intro_g3_b, 'scratch/diag_secuencia_offline.png', 'FIGURA_DUMMY_35')

insert_elements_after(target_p3, elems_g3_a + elems_g3_b)
print("Inserted Group 3: FSM & Offline Sequence")

# Group 4: Navigation, IA, Components (after target_p4)
intro_g4_a = (
    "A partir de los esquemas de maquetado aprobados, se formaliz\u00f3 el flujo de interacci\u00f3n del usuario mediante el Diagrama de "
    "Navegaci\u00f3n de la Interfaz de Usuario. Este diagrama especifica la ruta heur\u00edstica de navegaci\u00f3n que siguen los operadores desde "
    "la pantalla de autenticaci\u00f3n y el panel principal (Dashboard), facilitando el acceso contextual al Visor Cartogr\u00e1fico GIS, a la "
    "inspecci\u00f3n y asignaci\u00f3n interactiva en el Chasis de Puertos NAP, al Padr\u00f3n de Abonados, a la Cola de Asignaciones Offline y "
    "al generador de auditor\u00edas ejecutivas en PDF:"
)
elems_g4_a = build_figure_elements(intro_g4_a, 'scratch/diag_navegacion_sistema.png', 'FIGURA_DUMMY_37')

intro_g4_b = (
    "Para asegurar una \u00f3ptima usabilidad y una sobrecarga cognitiva m\u00ednima durante el trabajo a la intemperie, se estructur\u00f3 el "
    "Diagrama Jer\u00e1rquico de Arquitectura de Informaci\u00f3n (IA). En \u00e9l se organizan sistem\u00e1ticamente los niveles de granularidad de la "
    "plataforma: Nivel 0 (Portal de Acceso y Control de Credenciales), Nivel 1 (Vistas Operativas Principales), Nivel 2 (Paneles de Control "
    "y Filtros Din\u00e1micos) y Nivel 3 (Modales Transaccionales, Inspecci\u00f3n de Fibra y Acciones de Campo):"
)
elems_g4_b = build_figure_elements(intro_g4_b, 'scratch/diag_arquitectura_informacion.png', 'FIGURA_DUMMY_38')

intro_g4_c = (
    "Finalmente, para consolidar la estructura t\u00e9cnica del software previo a la fase de codificaci\u00f3n, se dise\u00f1\u00f3 el Diagrama de "
    "Arquitectura de Paquetes y Componentes de Software bajo el est\u00e1ndar UML. Este modelo ilustra la desacoplada relaci\u00f3n entre la capa "
    "de presentaci\u00f3n (Frontend React PWA con Tailwind CSS, Leaflet GIS y Dexie.js), la capa de servicios l\u00f3gicos (Backend Node.js/Express "
    "con Middlewares JWT/RBAC y Controladores REST) y la capa de persistencia h\u00edbrida (PostgreSQL relacional y motor de almacenamiento "
    "local en el navegador):"
)
elems_g4_c = build_figure_elements(intro_g4_c, 'scratch/diag_paquetes_componentes.png', 'FIGURA_DUMMY_39')

insert_elements_after(target_p4, elems_g4_a + elems_g4_b + elems_g4_c)
print("Inserted Group 4: Navigation, IA, Components")

# 2. Identify all figure captions in the document body (from p[105] onwards)
print("\n=== IDENTIFYING ALL FIGURE CAPTIONS IN BODY ===")
body_after_insert = doc.paragraphs[105:]
figure_captions = []
for p in body_after_insert:
    txt = p.text.strip()
    if txt.startswith("Figura ") or txt.startswith("FIGURA_DUMMY_"):
        figure_captions.append(p)

print(f"Total figure captions detected in body: {len(figure_captions)}")
if len(figure_captions) != 73:
    print(f"ERROR: Expected 73 captions, found {len(figure_captions)}")
    for i, p in enumerate(figure_captions):
        print(f"  [{i}]: {p.text[:60]}")
    raise RuntimeError("Figure caption count mismatch!")

# 3. Clean and renumber all 73 figure captions in the body
print("\n=== RENUMBERING ALL FIGURE CAPTIONS IN BODY ===")
for k, p_cap in enumerate(figure_captions):
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_id = str(320 + k)
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    # Clean the paragraph XML
    p_elem = p_cap._p
    # Keep only pPr
    pPr = p_elem.find(qn('w:pPr'))
    for child in list(p_elem):
        if child != pPr:
            p_elem.remove(child)
            
    # Set style to Figuras if not set
    p_cap.style = 'Figuras'
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(2)
    
    # Add bookmarkStart
    bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{bookmark_name}"/>')
    p_elem.append(bm_start)
    
    # Add text run
    r_elem = parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>'
        f'<w:b/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>'
        f'<w:t xml:space="preserve">{new_caption_text}</w:t></w:r>'
    )
    p_elem.append(r_elem)
    
    # Add bookmarkEnd
    bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>')
    p_elem.append(bm_end)

print("All 73 figure captions in body successfully renumbered and bookmarked!")

# 4. Update the Índice de figuras
print("\n=== UPDATING ÍNDICE DE FIGURAS ===")
# Find the exact range of the figure index
# P[8] is 'Índice de figuras'
# Then original had P[9] to P[74] (66 entries)
# Then P[75] was empty with <w:fldChar w:fldCharType="end"/>
# Then P[76] is 'Índice de tablas'

idx_title_p = None
idx_end_p = None
for i, p in enumerate(doc.paragraphs[:100]):
    if 'ndice de figuras' in p.text:
        idx_title_p = i
    if 'ndice de tablas' in p.text:
        idx_end_p = i
        break

print(f"Índice de figuras starts at P[{idx_title_p}], Índice de tablas at P[{idx_end_p}]")

# Let's collect existing figure index paragraphs
existing_toc_p = []
for i in range(idx_title_p + 1, idx_end_p):
    p = doc.paragraphs[i]
    if p.text.strip().startswith("Figura "):
        existing_toc_p.append(p)

print(f"Found {len(existing_toc_p)} existing TOC figure paragraphs.")

# We update existing_toc_p[0..65] for k=0..65
for k in range(min(len(existing_toc_p), 73)):
    p = existing_toc_p[k]
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    # Update hyperlink anchor
    hl = p._p.xpath('.//w:hyperlink')
    if hl:
        hl[0].attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor'] = bookmark_name
        
    # Update PAGEREF instruction
    instr = p._p.xpath('.//w:instrText')
    for inst in instr:
        if 'PAGEREF' in inst.text:
            inst.text = f" PAGEREF {bookmark_name} \\h "
            
    # Update text nodes inside hyperlink: first <w:t> is caption, second <w:t> is page
    t_nodes = p._p.xpath('.//w:hyperlink//w:t')
    if len(t_nodes) >= 2:
        t_nodes[0].text = new_caption_text
        t_nodes[-1].text = fig_page
    elif len(t_nodes) == 1:
        t_nodes[0].text = new_caption_text

# If there are more figures than existing TOC paragraphs (which is true: 73 vs 66, difference 7)
last_toc_p = existing_toc_p[-1]
for k in range(len(existing_toc_p), 73):
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    # Clone the XML of last_toc_p
    new_p_xml = copy.deepcopy(last_toc_p._p)
    
    # Update anchor
    hl = new_p_xml.xpath('.//w:hyperlink')
    if hl:
        hl[0].attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor'] = bookmark_name
        
    # Update PAGEREF
    instr = new_p_xml.xpath('.//w:instrText')
    for inst in instr:
        if 'PAGEREF' in inst.text:
            inst.text = f" PAGEREF {bookmark_name} \\h "
            
    # Update text nodes
    t_nodes = new_p_xml.xpath('.//w:hyperlink//w:t')
    if len(t_nodes) >= 2:
        t_nodes[0].text = new_caption_text
        t_nodes[-1].text = fig_page
        
    # Insert new_p_xml right after last_toc_p._p
    last_toc_p._p.addnext(new_p_xml)
    last_toc_p = docx.text.paragraph.Paragraph(new_p_xml, doc)

print(f"Successfully updated/appended all 73 entries in Índice de figuras!")

# 5. Save modified document to advance and mirror files
print("\n=== SAVING DOCUMENTS ===")
doc.save(DOC_ADVANCE)
print(f"Saved: {DOC_ADVANCE}")

doc.save(DOC_MIRROR)
print(f"Saved: {DOC_MIRROR}")

print("\n=== VERIFICATION RELOAD ===")
doc_check = docx.Document(DOC_ADVANCE)
check_body = doc_check.paragraphs[105:]
verified_captions = [p.text.strip() for p in check_body if p.text.strip().startswith("Figura ")]
print(f"Verified body figure captions count: {len(verified_captions)}")
print(f"First caption: {verified_captions[0]}")
print(f"Fig 28: {verified_captions[27]}")
print(f"Fig 31: {verified_captions[30]}")
print(f"Fig 34: {verified_captions[33]}")
print(f"Fig 35: {verified_captions[34]}")
print(f"Fig 37: {verified_captions[36]}")
print(f"Fig 38: {verified_captions[37]}")
print(f"Fig 39: {verified_captions[38]}")
print(f"Last caption: {verified_captions[-1]}")

check_toc = []
for p in doc_check.paragraphs[idx_title_p+1:idx_title_p+85]:
    txt = p.text.strip()
    if txt.startswith("Figura "):
        check_toc.append(txt)
print(f"Verified TOC entries count: {len(check_toc)}")
print(f"First TOC: {check_toc[0]}")
print(f"TOC 28: {check_toc[27]}")
print(f"TOC 31: {check_toc[30]}")
print(f"TOC 34: {check_toc[33]}")
print(f"TOC 35: {check_toc[34]}")
print(f"TOC 37: {check_toc[36]}")
print(f"TOC 38: {check_toc[37]}")
print(f"TOC 39: {check_toc[38]}")
print(f"Last TOC: {check_toc[-1]}")

print("\n=== COMPLETED SUCCESSFULLY! ===")
