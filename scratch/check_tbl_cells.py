import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')
for idx in [2, 13, 15]:
    t = doc.tables[idx]
    shds = t._tbl.xpath('.//w:shd')
    print(f"Table {idx} ({t.style.name}): shd count={len(shds)}, vals={[s.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') for s in shds[:5]]}")
    for r_i in [0, 1]:
        p = t.rows[r_i].cells[0].paragraphs[0]
        runs_info = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs]
        print(f"  Row {r_i} Cell 0: style={p.style.name}, align={p.alignment}, runs={runs_info}")

