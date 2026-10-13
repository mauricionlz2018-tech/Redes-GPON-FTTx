import docx

doc_06 = docx.Document('docs/Documentacion_Residencias_avance_BACKUP_06OCT.docx')
c4_idx = None
c5_idx = None
for i, p in enumerate(doc_06.paragraphs):
    if 'CAPÍTULO IV' in p.text.upper() and c4_idx is None:
        c4_idx = i
    if 'CAPÍTULO V' in p.text.upper() and c5_idx is None:
        c5_idx = i

print(f"BACKUP_06OCT: Cap IV en P{c4_idx}, Cap V en P{c5_idx}")
print(f"Total párrafos en Cap IV de BACKUP_06OCT: {c5_idx - c4_idx}")
for k in range(c4_idx, min(c4_idx + 20, c5_idx)):
    p = doc_06.paragraphs[k]
    draws = len(p._element.xpath('.//w:drawing'))
    print(f"P{k} [{p.style.name}, draws={draws}]: '{p.text[:70]}'")

