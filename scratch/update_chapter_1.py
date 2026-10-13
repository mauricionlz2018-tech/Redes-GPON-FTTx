import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import shutil

def update_chapter1():
    doc_path = 'docs/Documentacion_Residencias_avance.docx'
    print(f"Abriendo {doc_path}...")
    doc = docx.Document(doc_path)
    body = doc._body._element

    # 1. P151: Introducción de Capítulo 1
    p151 = doc.paragraphs[151]
    p151.text = 'En el presente capítulo se describen los antecedentes organizacionales, el contexto operativo y la estructura institucional de GPON TELECOM S.A. de C.V., delimitando su marco de actuación, filosofía corporativa y la infraestructura en la que se fundamenta el desarrollo del proyecto.'
    p151.runs[0].font.name = 'Arial'
    p151.runs[0].font.size = Pt(11)
    p151.paragraph_format.line_spacing = 1.5
    p151.paragraph_format.space_after = Pt(6)
    print("  -> Actualizada introducción de Capítulo 1.")

    # 2. P152: 1.1 Contexto de la organización
    p152 = doc.paragraphs[152]
    p152.text = '1.1 Contexto de la organización'
    p152.style = 'Heading 2'
    p152.runs[0].font.name = 'Arial'
    p152.runs[0].font.size = Pt(12)
    p152.runs[0].font.bold = True
    p152.paragraph_format.space_before = Pt(12)
    p152.paragraph_format.space_after = Pt(6)
    print("  -> Actualizado título 1.1 Contexto de la organización.")

    # 3. P153: Texto 1.1 con referencia formal y cita corta
    p153 = doc.paragraphs[153]
    p153.text = ''
    r1 = p153.add_run('La residencia profesional se desarrolla en la empresa GPON TELECOM S.A. de C.V., ubicada en Ejido San José del Rincón, Estado de México. De acuerdo con la información corporativa publicada en su portal web oficial ')
    r2 = p153.add_run('Infraestructura para la vida')
    r2.font.italic = True
    r3 = p153.add_run(' (https://gpontelecom.com.mx/#div_nosotros), la organización se consolida como «una empresa Carrier líder en la región norte del Estado de México. Nos consolidamos a partir de una empresa Wisp para evolucionar a una empresa con tecnología GPON» (GPON Telecom, 2026, sección «Nosotros»). Asimismo, la empresa respalda su presencia en el sector mediante una trayectoria consolidada, destacando que «tenemos más de 30 años de experiencia en el ámbito de las Telecomunicaciones, mejoramos la productividad de tu empresa y de tu vida diaria» (GPON Telecom, 2026).')
    
    for r in [r1, r2, r3]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    p153.paragraph_format.line_spacing = 1.5
    p153.paragraph_format.space_after = Pt(6)
    print("  -> Actualizado texto 1.1 con citas cortas y portal web.")

    # 4. P154: 1.2 Misión, Visión y Valores (antes 1.4)
    p154 = doc.paragraphs[154]
    p154.text = '1.2 Misión, Visión y Valores'
    p154.style = 'Heading 2'
    p154.runs[0].font.name = 'Arial'
    p154.runs[0].font.size = Pt(12)
    p154.runs[0].font.bold = True
    p154.paragraph_format.space_before = Pt(12)
    p154.paragraph_format.space_after = Pt(6)
    print("  -> Actualizado título 1.2 Misión, Visión y Valores.")

    # 5. P155: 1.2.1 Misión (antes 1.4.1)
    p155 = doc.paragraphs[155]
    p155.text = '1.2.1 Misión'
    p155.style = 'Heading 3'
    p155.runs[0].font.name = 'Arial'
    p155.runs[0].font.size = Pt(11)
    p155.runs[0].font.bold = True
    p155.paragraph_format.space_before = Pt(6)
    p155.paragraph_format.space_after = Pt(4)
    print("  -> Actualizado título 1.2.1 Misión.")

    # 6. P156: Texto Misión con cita corta entre comillas y referencia
    p156 = doc.paragraphs[156]
    p156.text = ''
    r_m1 = p156.add_run('La misión institucional de la empresa declara como compromiso fundamental con la sociedad: ')
    r_m2 = p156.add_run('«Brindar acceso a nuestros clientes con telecomunicaciones innovadoras a fin de impactar positivamente a comunidades rezagadas en comunicación para así contribuir a cerrar la brecha digital»')
    r_m3 = p156.add_run(' (GPON Telecom, 2026).')
    for r in [r_m1, r_m2, r_m3]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    p156.paragraph_format.line_spacing = 1.5
    p156.paragraph_format.space_after = Pt(6)
    print("  -> Actualizado texto 1.2.1 Misión con cita corta.")

    # 7. P157: 1.2.2 Visión (antes 1.4.2)
    p157 = doc.paragraphs[157]
    p157.text = '1.2.2 Visión'
    p157.style = 'Heading 3'
    p157.runs[0].font.name = 'Arial'
    p157.runs[0].font.size = Pt(11)
    p157.runs[0].font.bold = True
    p157.paragraph_format.space_before = Pt(6)
    p157.paragraph_format.space_after = Pt(4)
    print("  -> Actualizado título 1.2.2 Visión.")

    # 8. P158: Texto Visión con cita corta entre comillas y referencia
    p158 = doc.paragraphs[158]
    p158.text = ''
    r_v1 = p158.add_run('Como meta de proyección y liderazgo en el despliegue de redes ópticas a escala regional y nacional, la visión corporativa establece: ')
    r_v2 = p158.add_run('«Posicionarnos como una empresa Carrier de calidad sobresaliente en telecomunicaciones a nivel nacional, con el mejor sentido de filantropía por las comunidades donde desplegamos nuestra infraestructura»')
    r_v3 = p158.add_run(' (GPON Telecom, 2026).')
    for r in [r_v1, r_v2, r_v3]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    p158.paragraph_format.line_spacing = 1.5
    p158.paragraph_format.space_after = Pt(6)
    print("  -> Actualizado texto 1.2.2 Visión con cita corta.")

    # 9. P159: 1.2.3 Valores (antes 1.4.3)
    p159 = doc.paragraphs[159]
    p159.text = '1.2.3 Valores'
    p159.style = 'Heading 3'
    p159.runs[0].font.name = 'Arial'
    p159.runs[0].font.size = Pt(11)
    p159.runs[0].font.bold = True
    p159.paragraph_format.space_before = Pt(6)
    p159.paragraph_format.space_after = Pt(4)
    print("  -> Actualizado título 1.2.3 Valores.")

    # 10. P160: Introducción a Valores (estaba vacío)
    p160 = doc.paragraphs[160]
    p160.text = 'La cultura organizacional y el desempeño operativo del personal de GPON TELECOM se rigen por cuatro valores fundamentales formalizados en su portal corporativo (GPON Telecom, 2026):'
    p160.runs[0].font.name = 'Arial'
    p160.runs[0].font.size = Pt(11)
    p160.paragraph_format.line_spacing = 1.5
    p160.paragraph_format.space_after = Pt(6)

    # P161: Respeto
    p161 = doc.paragraphs[161]
    p161.text = ''
    r_val1_b = p161.add_run('Respeto: ')
    r_val1_b.font.bold = True
    r_val1_t = p161.add_run('«Hacia nuestro trabajo, nuestro entorno y principalmente dar un trato amable a nuestros clientes» (GPON Telecom, 2026).')
    for r in [r_val1_b, r_val1_t]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    p161.paragraph_format.line_spacing = 1.5
    p161.paragraph_format.space_after = Pt(6)

    # P162: Responsabilidad
    p162 = doc.paragraphs[162]
    p162.text = ''
    r_val2_b = p162.add_run('Responsabilidad: ')
    r_val2_b.font.bold = True
    r_val2_t = p162.add_run('«Social contribuimos a mejorar el lugar donde trabajamos, realizando acciones positivas. Brindamos servicio de Internet gratuito a escuelas y centros de salud» (GPON Telecom, 2026).')
    for r in [r_val2_b, r_val2_t]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    p162.paragraph_format.line_spacing = 1.5
    p162.paragraph_format.space_after = Pt(6)

    # P163: Innovación
    p163 = doc.paragraphs[163]
    p163.text = ''
    r_val3_b = p163.add_run('Innovación: ')
    r_val3_b.font.bold = True
    r_val3_t = p163.add_run('«En la adopción constante de nuevas tecnologías para ofrecer soluciones que se adapten a cada una de las necesidades de nuestros clientes» (GPON Telecom, 2026).')
    for r in [r_val3_b, r_val3_t]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    p163.paragraph_format.line_spacing = 1.5
    p163.paragraph_format.space_after = Pt(6)

    # P164: Inclusión
    p164 = doc.paragraphs[164]
    p164.text = ''
    r_val4_b = p164.add_run('Inclusión: ')
    r_val4_b.font.bold = True
    r_val4_t = p164.add_run('«Dentro de esta organización no haremos discriminación alguna, desde la contratación de capital humano hasta la atención a nuestros clientes, sin importar su origen étnico, cultural, social» (GPON Telecom, 2026).')
    for r in [r_val4_b, r_val4_t]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    p164.paragraph_format.line_spacing = 1.5
    p164.paragraph_format.space_after = Pt(6)
    print("  -> Actualizados los 4 valores institucionales con citas cortas y comillas.")

    # 11. P165: 1.3 Ubicación (antes 1.5)
    p165 = doc.paragraphs[165]
    p165.text = '1.3 Ubicación'
    p165.style = 'Heading 2'
    p165.runs[0].font.name = 'Arial'
    p165.runs[0].font.size = Pt(12)
    p165.runs[0].font.bold = True
    p165.paragraph_format.space_before = Pt(12)
    p165.paragraph_format.space_after = Pt(6)

    # P166: Texto Ubicación
    p166 = doc.paragraphs[166]
    p166.text = 'Las instalaciones principales y el centro de operaciones de red (NOC) de GPON TELECOM S.A. de C.V. se localizan en Ejido San José del Rincón, Estado de México, C.P. 50558.'
    p166.runs[0].font.name = 'Arial'
    p166.runs[0].font.size = Pt(11)
    p166.paragraph_format.line_spacing = 1.5
    p166.paragraph_format.space_after = Pt(6)

    # P170: Pie de Figura 1 (P167 e imagen permanecen intactos)
    p170 = doc.paragraphs[170]
    p170.text = 'Fuente: Tomado de Google Maps (2026) y portal web oficial GPON Telecom (https://gpontelecom.com.mx).'
    p170.runs[0].font.name = 'Arial'
    p170.runs[0].font.size = Pt(9)
    p170.runs[0].font.bold = True
    print("  -> Actualizada sección 1.3 Ubicación y fuente de Figura 1.")

    # 12. P172: 1.4 Generalidades de GPON TELECOM (antes 1.6)
    p172 = doc.paragraphs[172]
    p172.text = '1.4 Generalidades de GPON TELECOM'
    p172.style = 'Heading 2'
    p172.runs[0].font.name = 'Arial'
    p172.runs[0].font.size = Pt(12)
    p172.runs[0].font.bold = True
    p172.paragraph_format.space_before = Pt(12)
    p172.paragraph_format.space_after = Pt(6)

    # P173-P175: Texto Generalidades
    p173 = doc.paragraphs[173]
    p173.text = 'GPON Telecom S.A. de C.V. es una organización proveedora de servicios de conectividad y telecomunicaciones enfocada en la transmisión de datos a alta velocidad mediante fibra óptica hasta el hogar (FTTH/FTTx) bajo estándares internacionales GPON de la ITU-T. La empresa cuenta con experiencia técnica en el despliegue, operación y mantenimiento de planta externa e interna, atendiendo a hogares, comercios, escuelas y organizaciones de la región norte mexiquense (GPON Telecom, 2026).'
    p173.runs[0].font.name = 'Arial'
    p173.runs[0].font.size = Pt(11)
    p173.paragraph_format.line_spacing = 1.5
    p173.paragraph_format.space_after = Pt(6)

    p174 = doc.paragraphs[174]
    p174.text = 'A través del uso de tecnología GPON y topologías pasivas de árbol (punto a multipunto), la empresa provee enlaces estables y de alta velocidad con relaciones de división óptica calculadas formalmente para mitigar atenuaciones y preservar la calidad del servicio.'
    p174.runs[0].font.name = 'Arial'
    p174.runs[0].font.size = Pt(11)
    p174.paragraph_format.line_spacing = 1.5
    p174.paragraph_format.space_after = Pt(6)

    p175 = doc.paragraphs[175]
    p175.text = 'Entre sus principales actividades y líneas de servicio se destacan las siguientes:'
    p175.runs[0].font.name = 'Arial'
    p175.runs[0].font.size = Pt(11)
    p175.paragraph_format.line_spacing = 1.5
    p175.paragraph_format.space_after = Pt(6)

    # Actividades numeradas P176 a P183
    acts = [
        "1.- Provisión y distribución de servicios de Internet de banda ancha mediante enlaces de fibra óptica.",
        "2.- Diseño, tendido y mantenimiento preventivo y correctivo de infraestructura de redes ópticas (cables ADSS y canalizaciones).",
        "3.- Implementación e interconexión de infraestructura pasiva (cajas terminales NAP, splitters PLC y cierres de empalme FOSC).",
        "4.- Aprovisionamiento, calibración de potencia óptica y configuración de terminales ONT en predios de suscriptores.",
        "5.- Soporte técnico especializado y atención reactiva y proactiva a reportes de incidencias en campo.",
        "6.- Mantenimiento preventivo y correctivo de empalmes por fusión y monitoreo de atenuación en planta externa.",
        "7.- Desarrollo de soluciones de conectividad comunitaria para instituciones educativas y de salud pública.",
        "8.- Administración, supervisión y mapeo del inventario lógico y físico de la red de acceso FTTx."
    ]
    for act_idx, act_text in enumerate(acts):
        p_act = doc.paragraphs[176 + act_idx]
        p_act.text = act_text
        p_act.runs[0].font.name = 'Arial'
        p_act.runs[0].font.size = Pt(11)
        p_act.paragraph_format.line_spacing = 1.5
        p_act.paragraph_format.space_after = Pt(4)
    print("  -> Actualizada sección 1.4 Generalidades de GPON TELECOM.")

    # 13. P184: 1.5 Organigrama de la institución (antes 1.7)
    p184 = doc.paragraphs[184]
    p184.text = '1.5 Organigrama de la institución'
    p184.style = 'Heading 2'
    p184.runs[0].font.name = 'Arial'
    p184.runs[0].font.size = Pt(12)
    p184.runs[0].font.bold = True
    p184.paragraph_format.space_before = Pt(12)
    p184.paragraph_format.space_after = Pt(6)

    # P187: Pie de Figura 2 (P185 e imagen permanecen intactos)
    p187 = doc.paragraphs[187]
    p187.text = 'Fuente: Tomado de GPON TELECOM S.A. de C.V. (2026).'
    p187.runs[0].font.name = 'Arial'
    p187.runs[0].font.size = Pt(9)
    p187.runs[0].font.bold = True
    print("  -> Actualizada sección 1.5 Organigrama de la institución y fuente de Figura 2.")

    # 14. Agregar referencia formal en Referencias Bibliográficas
    # Localizar posición entre Gilbert y Grigorik
    target_ref_idx = None
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith("Gilbert, S., & Lynch, N."):
            target_ref_idx = i
            break

    if target_ref_idx is not None:
        target_p = doc.paragraphs[target_ref_idx]._element
        new_p = parse_xml(f'<w:p {nsdecls("w")}/>')
        target_p.addnext(new_p)
        p_ref = docx.text.paragraph.Paragraph(new_p, doc)
        p_ref.paragraph_format.line_spacing = 1.5
        p_ref.paragraph_format.space_after = Pt(6)
        
        # Referencia APA 7 con cursiva en el título del portal
        r_ref1 = p_ref.add_run('GPON Telecom S.A. de C.V. (2026). ')
        r_ref2 = p_ref.add_run('Infraestructura para la vida: Nosotros. ')
        r_ref2.font.italic = True
        r_ref3 = p_ref.add_run('https://gpontelecom.com.mx/#div_nosotros')
        for r in [r_ref1, r_ref2, r_ref3]:
            r.font.name = 'Arial'
            r.font.size = Pt(11)
        print("  -> Insertada referencia formal APA 7 de GPON Telecom en la sección de Referencias Bibliográficas.")

    # Guardar documento principal
    print(f"Guardando cambios en {doc_path}...")
    doc.save(doc_path)
    print("Guardado exitoso de Documentacion_Residencias_avance.docx.")

    # 15. Sincronizar Documentacion_Residencias (4).docx
    alt_path = 'docs/Documentacion_Residencias (4).docx'
    shutil.copyfile(doc_path, alt_path)
    print(f"Sincronizado {alt_path} exitosamente.")

if __name__ == '__main__':
    update_chapter1()

