import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

ch3_start = None
ch4_start = None
for i, p in enumerate(doc.paragraphs):
    txt_u = p.text.upper()
    if 'CAP' in txt_u and 'III' in txt_u:
        ch3_start = i
    if 'CAP' in txt_u and 'IV' in txt_u:
        ch4_start = i
        break

print(f"Chapter III range: P[{ch3_start}] to P[{ch4_start}]")

headings_ch3 = []
for i in range(ch3_start, ch4_start):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if t.startswith('3.') or p.style.name.startswith('Heading'):
        headings_ch3.append((i, p.style.name, t))

for h in headings_ch3:
    print(f"P[{h[0]}] ({h[1]}): {h[2]}")

