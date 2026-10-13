import docx

doc_dev = docx.Document('docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx')
doc_main = docx.Document('docs/Documentacion_Residencias_avance.docx')

dev_rels = doc_dev.part.rels
main_rels = doc_main.part.rels

cap4_rids = []
in_c4 = False
for p in doc_dev.paragraphs:
    t = p.text.strip()
    if 'CAPÍTULO IV' in t:
        in_c4 = True
    elif 'CAPÍTULO V' in t:
        in_c4 = False
        break
    if in_c4:
        for blip in p._element.xpath('.//a:blip/@r:embed'):
            cap4_rids.append((p.text[:30], blip))

for txt, rid in cap4_rids[:10]:
    t_dev = dev_rels[rid].target_ref if rid in dev_rels else 'None'
    t_main = main_rels[rid].target_ref if rid in main_rels else 'None'
    print(f"{rid}: dev='{t_dev}' vs main='{t_main}' (match={t_dev == t_main})")

