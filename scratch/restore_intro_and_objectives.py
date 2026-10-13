import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from copy import deepcopy
import shutil

pristine_path = 'docs/Documentacion_Residencias_avance_BACKUP_08OCT_FINAL_PAG.docx'
target_avance = 'docs/Documentacion_Residencias_avance.docx'
target_mirror = 'docs/Documentacion_Residencias (4).docx'

doc_pristine = docx.Document(pristine_path)
doc = docx.Document(target_avance)

# 1. Extraer los párrafos íntegros P160 a P186 desde BACKUP_08OCT_FINAL_PAG.docx
# P160: Introducción ... P165: Planteamiento del problema ... P173: Por consiguiente...
# P174: Justificación ... P179: Objetivos ... P186: 3. Programar el backend...
pristine_elements = []
for i in range(160, 187):
    p_orig = doc_pristine.paragraphs[i]
    p_copy = deepcopy(p_orig._element)
    # Limpiar cualquier sectPr previo dentro de la copia
    for sp in p_copy.xpath('./w:pPr/w:sectPr'):
        sp.getparent().remove(sp)
    pristine_elements.append((i, p_orig.text.strip(), p_copy))

# Crear el párrafo de la pregunta del Planteamiento del problema clonando el formato de P173
p173_copy = deepcopy(doc_pristine.paragraphs[173]._element)
# Reemplazar el texto dentro de p173_copy por la pregunta solicitada
runs = p173_copy.xpath('./w:r')
if runs:
    # Conservar el primer run y ponerle el texto de la pregunta, eliminar runs adicionales
    t_els = runs[0].xpath('./w:t')
    if t_els:
        t_els[0].text = "¿El desarrollo de un sistema web integral de inventario y mapeo lógico permitirá modernizar la gestión operativa y optimizar la administración de la red de fibra óptica GPON/FTTx en la empresa GPON TELECOM S.A. de C.V.?"
    for extra_r in runs[1:]:
        extra_r.getparent().remove(extra_r)

# 2. En doc (target_avance), localizar desde 'Introducción' (P113) hasta '3. Programar el backend...' (P126)
start_del = None
end_del = None
for idx, p in enumerate(doc.paragraphs):
    if start_del is None and p.text.strip() == 'Introducción':
        start_del = idx
    if start_del is not None and '3. Programar el backend del sistema' in p.text:
        end_del = idx
        break

print(f"Rango actual dañado en doc: P{start_del} a P{end_del}")

# Conservar el sectPr que cierra la Sección 2 (que estaba en end_del)
sectPr_sec2 = None
sp_list = doc.paragraphs[end_del]._element.xpath('./w:pPr/w:sectPr')
if sp_list:
    sectPr_sec2 = deepcopy(sp_list[0])
else:
    sectPr_sec2 = parse_xml(f'''
    <w:sectPr {nsdecls("w", "r")}>
        <w:headerReference w:type="default" r:id="rId17"/>
        <w:footerReference w:type="default" r:id="rId18"/>
        <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
        <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
        <w:pgNumType w:fmt="upperRoman"/>
        <w:cols w:space="708"/>
        <w:type w:val="nextPage"/>
        <w:docGrid w:linePitch="360"/>
    </w:sectPr>
    ''')

# Punto de anclaje previo a start_del (es el párrafo que cierra la Sección 1 con sectPr)
anchor_el = doc.paragraphs[start_del - 1]._element

# Eliminar los párrafos incompletos (recolectando referencias antes de borrar)
to_delete = [doc.paragraphs[k]._element for k in range(start_del, end_del + 1)]
for el in to_delete:
    el.getparent().remove(el)

# Insertar todos los párrafos íntegros en orden después de anchor_el, incluyendo la pregunta después de P173
curr_el = anchor_el
last_inserted_el = None
for orig_idx, txt, el_copy in pristine_elements:
    curr_el.addnext(el_copy)
    curr_el = el_copy
    last_inserted_el = el_copy
    if orig_idx == 173:
        # Insertar la pregunta del Planteamiento del problema inmediatamente después de P173
        curr_el.addnext(p173_copy)
        curr_el = p173_copy
        last_inserted_el = p173_copy

# Adjuntar sectPr_sec2 al último párrafo de Objetivos (P186: '3. Programar el backend...')
pPr_last = last_inserted_el.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
if pPr_last is None:
    pPr_last = parse_xml(f'<w:pPr {nsdecls("w")}/>')
    last_inserted_el.insert(0, pPr_last)
pPr_last.append(sectPr_sec2)

doc.save(target_avance)
shutil.copyfile(target_avance, target_mirror)
print("¡Introducción, Planteamiento del problema (con la pregunta), Justificación y Objetivos restaurados íntegramente!")

