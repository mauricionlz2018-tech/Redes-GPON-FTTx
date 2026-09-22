import os, subprocess, pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

src_backup = r'docs/Plantilla presentacion de residencia_BACKUP.pptx'
prs = pptx.Presentation(src_backup)

GREEN_COLOR = RGBColor(0x76, 0x92, 0x3C)
BLACK_COLOR = RGBColor(0x00, 0x00, 0x00)
GRAY_COLOR = RGBColor(0x50, 0x50, 0x50)

def set_font(run, name="Arial", size=18, bold=False, italic=False, color=BLACK_COLOR):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def strip_pPr(para):
    pPr = para._element.find('{http://schemas.openxmlformats.org/drawingml/2006/main}pPr')
    if pPr is not None:
        para._element.remove(pPr)

# ==========================================
# SLIDE 1: PORTADA
# ==========================================
s1 = prs.slides[0]
for shp in s1.shapes:
    if shp.shape_id == 93:
        tf = shp.text_frame
        tf.clear()
        p1 = tf.paragraphs[0]
        strip_pPr(p1)
        p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run()
        r1.text = "UNIVERSIDAD MEXIQUENSE DEL BICENTENARIO\n"
        set_font(r1, name="Calibri", size=24, bold=True, color=BLACK_COLOR)
        r2 = p1.add_run()
        r2.text = "UNIDAD DE ESTUDIOS SUPERIORES SAN JOSÉ DEL RINCÓN"
        set_font(r2, name="Calibri", size=20, bold=True, color=BLACK_COLOR)
    elif shp.shape_id == 92:
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        lines = [
            ("Carrera:", 17, False, False, GRAY_COLOR, 0),
            ("Ingeniería en Sistemas Computacionales", 19, True, False, BLACK_COLOR, 8),
            ("“Sistema de Inventario y Mapeo Lógico de Redes GPON / FTTx”", 21, True, False, GREEN_COLOR, 12),
            ("Presenta:", 17, False, False, GRAY_COLOR, 0),
            ("Mauricio Nolazco Lonjino", 19, True, False, BLACK_COLOR, 8),
            ("Asesor Interno:", 17, False, False, GRAY_COLOR, 0),
            ("I.S.C. Leonardo Becerril Sánchez", 19, True, False, BLACK_COLOR, 0)
        ]
        for i, (txt, sz, b, it, col, sp_after) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            strip_pPr(p)
            p.alignment = PP_ALIGN.CENTER
            if sp_after > 0:
                p.space_after = Pt(sp_after)
            r = p.add_run()
            r.text = txt
            set_font(r, name="Arial", size=sz, bold=b, italic=it, color=col)
    elif shp.shape_id == 91:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Septiembre de 2026"
        set_font(r, name="Calibri", size=18, bold=False, color=GRAY_COLOR)

# ==========================================
# SLIDE 2: AGENDA
# ==========================================
s2 = prs.slides[1]
for shp in s2.shapes:
    if shp.shape_id == 99:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Agenda"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 100:
        shp.top = Inches(1.1)
        shp.height = Inches(4.3)
        shp.left = Inches(1.0)
        shp.width = Inches(8.2)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        items = [
            "1. Empresa (Ubicación Geográfica y Organigrama)",
            "2. Introducción al Proyecto",
            "3. Planteamiento del Problema y Pregunta Generadora",
            "4. Justificación Técnica y Operativa",
            "5. Objetivos del Proyecto (General y Específicos)",
            "6. Herramientas Tecnológicas (Stack de Software)",
            "7. Evidencias de Avance Técnico",
            "8. Cronograma de Actividades",
            "9. Conclusiones Preliminares (Avance)",
            "10. Bibliografía y Referencias"
        ]
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(5)
            r = p.add_run()
            r.text = it
            set_font(r, name="Arial", size=16, bold=True, color=BLACK_COLOR)
    elif shp.shape_id == 101:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "2/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 3: EMPRESA - MAPA
# ==========================================
s3 = prs.slides[2]
if os.path.exists('scratch/fig1_map_extracted.png'):
    with open('scratch/fig1_map_extracted.png', 'rb') as f:
        s3.part.related_part('rId4')._blob = f.read()

for shp in s3.shapes:
    if shp.shape_id == 108:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Empresa"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 106:
        shp.top = Inches(1.1)
        shp.height = Inches(1.6)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        bullets = [
            ("• GPON TELECOM S.A. de C.V.: ", "Empresa proveedora de servicios de telecomunicaciones e Internet de alta velocidad mediante fibra óptica (ISP regional).\n"),
            ("• Ubicación Geográfica: ", "Sede operativa y nodo central (NOC) en San José del Rincón, Estado de México, abasteciendo conectividad GPON/FTTx urbana y rural.")
        ]
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.space_after = Pt(4)
        for bld, body in bullets:
            r1 = p.add_run()
            r1.text = bld
            set_font(r1, name="Arial", size=15, bold=True, color=BLACK_COLOR)
            r2 = p.add_run()
            r2.text = body
            set_font(r2, name="Arial", size=15, bold=False, color=BLACK_COLOR)
    elif shp.shape_id == 110:
        shp.top = Inches(2.8)
        shp.left = Inches(1.8)
        shp.width = Inches(5.5)
        shp.height = Inches(2.45)
    elif shp.shape_id == 111:
        shp.top = Inches(5.35)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "Figura 1. Localización geográfica de la empresa GPON TELECOM S.A. de C.V."
        set_font(r, name="Calibri", size=14, bold=False, color=GRAY_COLOR)
    elif shp.shape_id == 109:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "3/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 4: EMPRESA - ORGANIGRAMA
# ==========================================
s4 = prs.slides[3]
if os.path.exists('scratch/fig_2_bw_final.png'):
    with open('scratch/fig_2_bw_final.png', 'rb') as f:
        s4.part.related_part('rId4')._blob = f.read()

for shp in list(s4.shapes):
    if shp.shape_id == 120:
        sp = shp._element
        sp.getparent().remove(sp)

for shp in s4.shapes:
    if shp.shape_id == 117:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Empresa"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 119:
        shp.top = Inches(1.0)
        shp.left = Inches(1.2)
        shp.width = Inches(6.8)
        shp.height = Inches(4.3)
    elif shp.shape_id == 121:
        shp.top = Inches(5.42)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "Figura 2. Organigrama de la empresa GPON TELECOM S.A. de C.V."
        set_font(r, name="Calibri", size=14, bold=False, color=GRAY_COLOR)
    elif shp.shape_id == 118:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "4/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 5: INTRODUCCIÓN
# ==========================================
s5 = prs.slides[4]
for shp in s5.shapes:
    if shp.shape_id == 128:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Introducción"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 127:
        shp.top = Inches(1.1)
        shp.height = Inches(3.8)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        bullets = [
            ("• Propósito General: ", "Desarrollo de una plataforma web para el inventario centralizado y mapeo lógico de la red de fibra óptica GPON/FTTx en GPON TELECOM S.A. de C.V."),
            ("• Problemática Atendida: ", "Reemplazar la gestión manual y desarticulada en campo, previniendo la colisión de puertos y la saturación imprevista en cajas NAP."),
            ("• Arquitectura Desacoplada: ", "Integración de visor cartográfico geoespacial (React 18 + Leaflet), backend seguro (Node.js + Express) y base de datos relacional (PostgreSQL 16)."),
            ("• Resiliencia Operativa: ", "Modelo Offline-First (IndexedDB con Dexie.js) que faculta a los técnicos para continuar registrando operaciones en zonas sin cobertura móvil.")
        ]
        for i, (bold_prefix, text_body) in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(6)
            r1 = p.add_run()
            r1.text = bold_prefix
            set_font(r1, name="Arial", size=15, bold=True, color=BLACK_COLOR)
            r2 = p.add_run()
            r2.text = text_body
            set_font(r2, name="Arial", size=15, bold=False, color=BLACK_COLOR)
    elif shp.shape_id == 129:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "5/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 6: PLANTEAMIENTO DEL PROBLEMA
# ==========================================
s6 = prs.slides[5]
for shp in s6.shapes:
    if shp.shape_id == 136:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Planteamiento del problema"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 135:
        shp.top = Inches(1.05)
        shp.height = Inches(3.9)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        bullets = [
            ("• Carencia de Inventario Centralizado: ", "Falta de control físico y lógico de splitters, cajas NAP y puertos SC-APC, provocando retrasos en la atención de nuevas altas."),
            ("• Colisión por Asignación Concurrente: ", "La asignación simultánea y manual de puertos por diferentes técnicos genera conflictos de servicio y saturación invisible para el NOC."),
            ("• Aislamiento en Campo: ", "Pérdida de conectividad celular en cuadrillas rurales, imposibilitando la consulta o actualización de la red en tiempo real.")
        ]
        for i, (bold_prefix, text_body) in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(4)
            r1 = p.add_run()
            r1.text = bold_prefix
            set_font(r1, name="Arial", size=14.5, bold=True, color=BLACK_COLOR)
            r2 = p.add_run()
            r2.text = text_body
            set_font(r2, name="Arial", size=14.5, bold=False, color=BLACK_COLOR)
            
        p_q = tf.add_paragraph()
        strip_pPr(p_q)
        p_q.space_before = Pt(8)
        p_q.space_after = Pt(2)
        r_q_lbl = p_q.add_run()
        r_q_lbl.text = "Pregunta Generadora:\n"
        set_font(r_q_lbl, name="Arial", size=15, bold=True, color=GREEN_COLOR)
        
        r_q = p_q.add_run()
        r_q.text = "“¿Es viable optimizar la gestión de inventario y la trazabilidad técnica de la red GPON/FTTx mediante una plataforma web georreferenciada con arquitectura Offline-First?”\n"
        set_font(r_q, name="Arial", size=14.5, bold=True, italic=True, color=BLACK_COLOR)
        
        r_ans = p_q.add_run()
        r_ans.text = "Respuesta: Sí."
        set_font(r_ans, name="Arial", size=17, bold=True, color=GREEN_COLOR)
        
    elif shp.shape_id == 137:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "6/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 7: JUSTIFICACIÓN
# ==========================================
s7 = prs.slides[6]
for shp in s7.shapes:
    if shp.shape_id == 143:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Justificación"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 145:
        shp.top = Inches(1.1)
        shp.height = Inches(3.8)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        bullets = [
            ("• Eficiencia Operativa: ", "Automatiza la asignación y liberación de puertos SC-APC, reduciendo el tiempo de atención técnica de 45 a 15 minutos por suscriptor."),
            ("• Trazabilidad Geoespacial: ", "Faculta al personal de soporte y NOC para visualizar la red mediante cartografía satelital y semáforos de saturación en tiempo real."),
            ("• Resiliencia en Terreno: ", "El esquema Offline-First con IndexedDB garantiza continuidad operativa en zonas sin cobertura, sincronizando datos al reconectar."),
            ("• Reducción de Costos: ", "Disminuye en un 80% los traslados técnicos fallidos por saturación de cajas NAP y protege la inversión en infraestructura pasiva.")
        ]
        for i, (bold_prefix, text_body) in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(6)
            r1 = p.add_run()
            r1.text = bold_prefix
            set_font(r1, name="Arial", size=15, bold=True, color=BLACK_COLOR)
            r2 = p.add_run()
            r2.text = text_body
            set_font(r2, name="Arial", size=15, bold=False, color=BLACK_COLOR)
    elif shp.shape_id == 144:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "7/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 8: OBJETIVO GENERAL
# ==========================================
s8 = prs.slides[7]
for shp in s8.shapes:
    if shp.shape_id == 152:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Objetivos"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 151:
        shp.top = Inches(1.1)
        shp.height = Inches(3.8)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        
        p0 = tf.paragraphs[0]
        strip_pPr(p0)
        p0.space_after = Pt(6)
        r0 = p0.add_run()
        r0.text = "Objetivo General:"
        set_font(r0, name="Arial", size=20, bold=True, color=GREEN_COLOR)
        
        p1 = tf.add_paragraph()
        strip_pPr(p1)
        p1.space_after = Pt(10)
        r1 = p1.add_run()
        r1.text = "Desarrollar e implementar un sistema web para el inventario y mapeo lógico de la infraestructura de red GPON/FTTx de la empresa GPON TELECOM S.A. de C.V., mediante arquitectura cliente-servidor, bases de datos relacionales y visualización geoespacial, optimizando la asignación de puertos y la trazabilidad técnica en campo."
        set_font(r1, name="Arial", size=15.5, bold=False, color=BLACK_COLOR)
        
        pillars = [
            ("• Alcance de Dominio: ", "Control pasivo integral desde el ODF central hasta cajas NAP terminales de 16 puertos."),
            ("• Concurrencia Robusta: ", "Bloqueo pesimista de registros para garantizar consistencia transaccional ACID."),
            ("• Disponibilidad Móvil: ", "Sincronización diferida y operación continua en cuadrillas técnicas de campo.")
        ]
        for bold_pfx, body in pillars:
            p = tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(4)
            r_b = p.add_run()
            r_b.text = bold_pfx
            set_font(r_b, name="Arial", size=14.5, bold=True, color=BLACK_COLOR)
            r_t = p.add_run()
            r_t.text = body
            set_font(r_t, name="Arial", size=14.5, bold=False, color=BLACK_COLOR)
            
    elif shp.shape_id == 153:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "8/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 9: OBJETIVOS ESPECÍFICOS
# ==========================================
s9 = prs.slides[8]
for shp in s9.shapes:
    if shp.shape_id == 160:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Objetivos"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 159:
        shp.top = Inches(1.05)
        shp.height = Inches(3.9)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        
        p0 = tf.paragraphs[0]
        strip_pPr(p0)
        p0.space_after = Pt(4)
        r0 = p0.add_run()
        r0.text = "Objetivos Específicos:"
        set_font(r0, name="Arial", size=19, bold=True, color=GREEN_COLOR)
        
        specs = [
            ("1. Levantamiento e Infraestructura: ", "Caracterizar la red GPON pasiva (ODF, splitters PLC 1:4 y 1:16, cajas NAP y acometidas de fibra)."),
            ("2. Modelado de Datos y Arquitectura: ", "Diseñar la base de datos relacional y los diagramas estructurales (DER notación Chen, casos de uso y secuencia UML)."),
            ("3. Cartografía Web Responsiva: ", "Maquetar e implementar interfaces de usuario optimizadas para móviles con mapas satelitales en React Leaflet."),
            ("4. Backend Transaccional: ", "Programar la API REST en Node.js/PostgreSQL con bloqueo pesimista (SELECT ... FOR UPDATE) y autenticación JWT."),
            ("5. Motor Offline-First: ", "Desarrollar persistencia local en Dexie.js (IndexedDB) para sincronización automática de operaciones en campo."),
            ("6. Validación y Reportes: ", "Ejecutar pruebas de consistencia concurrente, contrastes WCAG 2.1 y generación de reportes técnicos PDF.")
        ]
        for bold_pfx, body in specs:
            p = tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(3)
            r_b = p.add_run()
            r_b.text = bold_pfx
            set_font(r_b, name="Arial", size=14, bold=True, color=BLACK_COLOR)
            r_t = p.add_run()
            r_t.text = body
            set_font(r_t, name="Arial", size=14, bold=False, color=BLACK_COLOR)
            
    elif shp.shape_id == 161:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "9/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 10: HERRAMIENTAS
# ==========================================
s10 = prs.slides[9]
for shp in s10.shapes:
    if shp.shape_id == 168:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Herramientas"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 167:
        shp.top = Inches(1.05)
        shp.left = Inches(0.7)
        shp.width = Inches(8.6)
        shp.height = Inches(0.6)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "Para el desarrollo de la plataforma web se seleccionó un stack tecnológico desacoplado y orientado a resiliencia:"
        set_font(r, name="Arial", size=14.5, bold=True, color=BLACK_COLOR)
    elif shp.shape_id == 170:
        shp.top = Inches(4.55)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        shp.height = Inches(0.45)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "Figura 3. Stack tecnológico desacoplado (Frontend PWA, Backend REST, Base de Datos y DevOps)."
        set_font(r, name="Calibri", size=14, bold=False, color=GRAY_COLOR)
    elif shp.shape_id == 169:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "10/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

if os.path.exists('scratch/fig_herramientas_gpon.png'):
    for shp in list(s10.shapes):
        if shp.shape_id in [171, 184]:
            sp = shp._element
            sp.getparent().remove(sp)
    s10.shapes.add_picture('scratch/fig_herramientas_gpon.png', Inches(0.8), Inches(1.75), Inches(8.4), Inches(2.52))

# ==========================================
# SLIDE 11: EVIDENCIAS
# ==========================================
s11 = prs.slides[10]
for shp in s11.shapes:
    if shp.shape_id == 190:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Evidencias"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 192:
        shp.top = Inches(1.1)
        shp.height = Inches(3.8)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        
        p0 = tf.paragraphs[0]
        strip_pPr(p0)
        p0.space_after = Pt(6)
        r0 = p0.add_run()
        r0.text = "Principales entregables y componentes técnicos implementados:"
        set_font(r0, name="Arial", size=18, bold=True, color=GREEN_COLOR)
        
        evidences = [
            ("• Documentación Técnica Formal: ", "Memoria de residencia completa con 66 figuras estandarizadas en blanco y negro, especificación de requerimientos y diagrama de flujo global ANSI/ISO."),
            ("• Mapeo Cartográfico GIS Interactivo: ", "Prototipo funcional en React Leaflet con polilíneas de fibra, trazado de rutas, semáforo de saturación en cajas NAP y calibración satelital GPS."),
            ("• Matriz de Puertos y Bloqueo Concurrente: ", "Chasis interactivo de 16 puertos SC-APC con resolución de colisiones mediante transacciones pesimistas (SELECT ... FOR UPDATE) en PostgreSQL."),
            ("• Repositorio de Código y Motor Offline: ", "Base de código estructurada en TypeScript, modelos Sequelize y base de datos local Dexie.js (IndexedDB) para operación continua en campo.")
        ]
        for bold_pfx, body in evidences:
            p = tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(5)
            r_b = p.add_run()
            r_b.text = bold_pfx
            set_font(r_b, name="Arial", size=14.5, bold=True, color=BLACK_COLOR)
            r_t = p.add_run()
            r_t.text = body
            set_font(r_t, name="Arial", size=14.5, bold=False, color=BLACK_COLOR)
            
    elif shp.shape_id == 191:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "11/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 12: CRONOGRAMA
# ==========================================
s12 = prs.slides[11]
cronograma_img_path = r'C:/Users/karen/.gemini/antigravity/brain/a33b5b22-c89b-4ff3-a6cc-733a0c69ceac/cronograma_actividades_gpon_20sep.png'
if os.path.exists(cronograma_img_path):
    with open(cronograma_img_path, 'rb') as f:
        s12.part.related_part('rId4')._blob = f.read()

for shp in list(s12.shapes):
    if shp.shape_id in [200, 202, 203, 204, 205, 207, 208, 209]:
        sp = shp._element
        sp.getparent().remove(sp)

for shp in s12.shapes:
    if shp.shape_id == 198:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Cronograma"
        set_font(r, name="Calibri", size=44, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 201:
        shp.top = Inches(1.05)
        shp.left = Inches(0.87)
        shp.width = Inches(7.4)
        shp.height = Inches(4.25)
    elif shp.shape_id == 206:
        shp.top = Inches(5.38)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        shp.height = Inches(0.45)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "Figura 4. Cronograma de actividades del proyecto de residencia profesional (Corte al 20 de septiembre de 2026)."
        set_font(r, name="Calibri", size=13.5, bold=False, color=GRAY_COLOR)
    elif shp.shape_id == 199:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "12/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 13: CONCLUSIONES PRELIMINARES
# ==========================================
s13 = prs.slides[12]
for shp in s13.shapes:
    if shp.shape_id == 216:
        shp.left = Inches(3.0)
        shp.width = Inches(6.0)
        shp.top = Inches(0.0)
        shp.height = Inches(0.70)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Conclusiones Preliminares"
        set_font(r, name="Calibri", size=38, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 215:
        shp.top = Inches(1.1)
        shp.height = Inches(3.8)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        
        p0 = tf.paragraphs[0]
        strip_pPr(p0)
        p0.space_after = Pt(5)
        r0 = p0.add_run()
        r0.text = "Evaluación de avance conforme al cronograma de actividades:"
        set_font(r0, name="Arial", size=17, bold=True, color=GREEN_COLOR)
        
        bullets = [
            ("• Fases Cumplidas al 100% (Etapas 1 a 6): ", "Se concluyeron exitosamente en tiempo y forma la inmersión técnica, levantamiento de requerimientos, diseño de arquitectura cliente-servidor, modelado DER y diseño de wireframes responsivos."),
            ("• Fase Actual en Desarrollo (~70% de avance): ", "Se encuentra en ejecución la codificación de la API REST, la configuración de transacciones pesimistas en PostgreSQL y la persistencia local en Dexie.js para cuadrillas en campo."),
            ("• Fases Programadas por Ejecutar: ", "Pruebas piloto con personal de cuadrilla en San José del Rincón, simulación de alta concurrencia y empaquetado final multicontenedor con Docker Compose."),
            ("• Diagnóstico Preliminar: ", "El proyecto cumple rigurosamente con los plazos estipulados en el anteproyecto, demostrando la viabilidad técnica y operativa de la solución para GPON TELECOM S.A. de C.V.")
        ]
        for bold_pfx, body in bullets:
            p = tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(5)
            r_b = p.add_run()
            r_b.text = bold_pfx
            set_font(r_b, name="Arial", size=14, bold=True, color=BLACK_COLOR)
            r_t = p.add_run()
            r_t.text = body
            set_font(r_t, name="Arial", size=14, bold=False, color=BLACK_COLOR)
            
    elif shp.shape_id == 217:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "13/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 14: BIBLIOGRAFÍA Y REFERENCIAS
# ==========================================
s14 = prs.slides[13]
for shp in s14.shapes:
    if shp.shape_id == 224:
        shp.left = Inches(3.0)
        shp.width = Inches(6.0)
        shp.top = Inches(0.0)
        shp.height = Inches(0.70)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Bibliografía y Referencias"
        set_font(r, name="Calibri", size=38, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 223:
        shp.top = Inches(1.15)
        shp.left = Inches(0.8)
        shp.width = Inches(8.4)
        shp.height = Inches(3.8)
        tf = shp.text_frame
        tf.clear()
        strip_pPr(tf.paragraphs[0])
        
        refs = [
            "• International Telecommunication Union. (2003). Recomendación ITU-T G.984.1 / G.984.2: Redes ópticas pasivas con capacidad de gigabit (GPON). Ginebra: ITU.",
            "• Tanenbaum, A. S., & Wetherall, D. J. (2014). Redes de computadoras (5.ª ed.). Ciudad de México: Pearson Educación.",
            "• Silberschatz, A., Korth, H. F., & Sudarshan, S. (2019). Database System Concepts (7th ed.). Nueva York: McGraw-Hill.",
            "• World Wide Web Consortium (W3C). (2018). Web Content Accessibility Guidelines (WCAG) 2.1. W3C Recommendation. https://www.w3.org/TR/WCAG21/",
            "• PostgreSQL Global Development Group. (2024). PostgreSQL 16.0 Documentation: Explicit Locking and Transaction Isolation. https://www.postgresql.org/docs/16/"
        ]
        for i, ref in enumerate(refs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            strip_pPr(p)
            p.space_after = Pt(5)
            r = p.add_run()
            r.text = ref
            set_font(r, name="Arial", size=13.5, bold=False, color=BLACK_COLOR)
            
    elif shp.shape_id == 225:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "14/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# ==========================================
# SLIDE 15: CONTRAPORTADA
# ==========================================
s15 = prs.slides[14]
for shp in s15.shapes:
    if shp.shape_id == 234:
        shp.top = Inches(1.8)
        shp.left = Inches(1.5)
        shp.width = Inches(6.1)
        shp.height = Inches(1.2)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "¡Muchas Gracias!"
        set_font(r, name="Calibri", size=56, bold=False, color=GREEN_COLOR)
    elif shp.shape_id == 236:
        shp.top = Inches(3.1)
        shp.left = Inches(1.0)
        shp.width = Inches(7.1)
        shp.height = Inches(1.2)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.CENTER
        r1 = p.add_run()
        r1.text = "¿Preguntas o Comentarios?\n"
        set_font(r1, name="Calibri", size=42, bold=False, color=GREEN_COLOR)
        p2 = tf.add_paragraph()
        strip_pPr(p2)
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = "Espacio de retroalimentación técnica"
        set_font(r2, name="Arial", size=18, bold=False, italic=True, color=GRAY_COLOR)
    elif shp.shape_id == 233:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = "Septiembre de 2026"
        set_font(r, name="Calibri", size=18, bold=False, color=GRAY_COLOR)
    elif shp.shape_id == 235:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        strip_pPr(p)
        r = p.add_run()
        r.text = "15/15"
        set_font(r, name="Calibri", size=16, color=GREEN_COLOR)

# Save directly to target
target_final = r'docs/Plantilla presentación de residencia.pptx'
prs.save(target_final)
print("SUCCESS: Saved updated presentation to:", target_final)

final_copy = r'docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx'
prs.save(final_copy)
print("SUCCESS: Saved copy to:", final_copy)

