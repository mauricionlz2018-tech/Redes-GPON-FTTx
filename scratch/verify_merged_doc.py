import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

print(f"Total párrafos: {len(doc.paragraphs)}")
print(f"Total secciones: {len(doc.sections)}")

print("\n--- ENCABEZADOS POR SECCIÓN ---")
for i, s in enumerate(doc.sections):
    h_draw = len(s.header._element.xpath('.//w:drawing'))
    print(f"  Sec {i}: header_drawings={h_draw}")

print("\n--- ÍNDICE DE TABLAS (ÚLTIMAS ENTRADAS) ---")
for p in doc.paragraphs[:200]:
    if p.text.strip().startswith('Tabla 3') and '\t' in p.text:
        print(f"  {p.text.strip()}")

print("\n--- FIGURAS DE CAPÍTULO IV (MUESTRA) ---")
for p in doc.paragraphs[:200]:
    t = p.text.strip()
    if ('Figura 66.' in t or 'Figura 80.' in t or 'Figura 96.' in t) and '\t' in t:
        print(f"  {t}")

print("\n--- CAPÍTULO IV EN ÍNDICE DE CONTENIDO (SDT) ---")
sdts = doc._body._element.xpath('./w:sdt')
if sdts:
    for p in sdts[0].findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
        t = "".join(p.itertext()).strip()
        if any(k in t for k in ['CAPÍTULO IV', '4.1', '4.2', '4.3', '4.4', '4.5', 'CAPÍTULO V', 'Planteamiento']):
            print(f"  {t[:80]}")
