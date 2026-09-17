import os
import sys

sys.path.insert(0, os.getcwd())
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import docx
from copy import deepcopy

# Import our modular builders
from scratch.builder_part1_prelims import build_preliminaries_and_chapter_1
from scratch.builder_part2_theory import build_chapter_2
from scratch.builder_part3_design import build_chapter_3
from scratch.builder_part4_coding import build_chapter_4
from scratch.builder_part5_deployment_results import build_chapters_5_6_7_annex
from scratch.common_docx import sanitize_text

def assemble():
    print("=== INICIANDO ENSAMBLADO FINAL DE LA TESIS REESTRUCTURADA ===")
    
    source_path = 'docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx'
    output_path = 'docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx'
    
    if not os.path.exists(source_path):
        print(f"Error: {source_path} no existe.")
        sys.exit(1)
        
    print(f"1. Cargando documento base y capturando tablas desde: {source_path}")
    source_doc = docx.Document(source_path)
    tables_source = [deepcopy(t) for t in source_doc.tables]
    print(f"   Tablas capturadas: {len(tables_source)}")
    
    print("2. Inicializando documento destino y limpiando cuerpo preservando estilos y márgenes...")
    target_doc = docx.Document(source_path)
    body = target_doc._body._element
    
    # Preservar sectPr
    for child in list(body):
        if child.tag.endswith('sectPr'):
            continue
        body.remove(child)
        
    print("3. Construyendo Preliminares, Índices y Capítulo I...")
    build_preliminaries_and_chapter_1(target_doc, tables_source, None)
    
    print("4. Construyendo Capítulo II: Marco Teórico o Estado del Arte (con 15 Figuras y Tablas 1-3)...")
    build_chapter_2(target_doc, tables_source)
    
    print("5. Construyendo Capítulo III: Diseño, Modelado y Maquetado del Sistema (con Figuras 16-22 y Tablas 4-12)...")
    build_chapter_3(target_doc, tables_source)
    
    print("6. Construyendo Capítulo IV: Codificación del Sistema (con código real, Figuras 23-26 y Tabla 13)...")
    build_chapter_4(target_doc, tables_source)
    
    print("7. Construyendo Capítulos V, VI, VII y Anexo A (con Figuras 27-29)...")
    build_chapters_5_6_7_annex(target_doc)
    
    print("8. Ejecutando barrido estricto de sanitización (eliminación de guiones largos, puntos y coma, emojis)...")
    em_dash_count = 0
    semicolon_count = 0
    
    for p in target_doc.paragraphs:
        if '—' in p.text or '–' in p.text or ';' in p.text:
            for r in p.runs:
                if '—' in r.text or '–' in r.text:
                    em_dash_count += r.text.count('—') + r.text.count('–')
                    r.text = r.text.replace('—', '-').replace('–', '-')
                if ';' in r.text:
                    semicolon_count += r.text.count(';')
                    r.text = r.text.replace(';', ',')
                r.text = sanitize_text(r.text)
                
    for t in target_doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    if '—' in cp.text or '–' in cp.text or ';' in cp.text:
                        for cr in cp.runs:
                            if '—' in cr.text or '–' in cr.text:
                                em_dash_count += cr.text.count('—') + cr.text.count('–')
                                cr.text = cr.text.replace('—', '-').replace('–', '-')
                            if ';' in cr.text:
                                semicolon_count += cr.text.count(';')
                                cr.text = cr.text.replace(';', ',')
                            cr.text = sanitize_text(cr.text)

    print(f"   Guiones largos erradicados: {em_dash_count}")
    print(f"   Puntos y coma erradicados: {semicolon_count}")
    
    print(f"9. Guardando documento final en: {output_path}")
    target_doc.save(output_path)
    print("=== PROCESO COMPLETADO EXITOSAMENTE ===")

if __name__ == '__main__':
    assemble()
