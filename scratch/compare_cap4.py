import docx

doc_dev = docx.Document('docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx')
doc_06 = docx.Document('docs/Documentacion_Residencias_avance_BACKUP_06OCT.docx')

dev_cap4 = [p.text.strip() for p in doc_dev.paragraphs if p.text.strip()]
# filter from CAPÍTULO IV to before CAPÍTULO V
dev_texts = []
in_c4 = False
for p in doc_dev.paragraphs:
    t = p.text.strip()
    if 'CAPÍTULO IV' in t:
        in_c4 = True
    elif 'CAPÍTULO V' in t:
        in_c4 = False
        break
    if in_c4 and t:
        dev_texts.append(t)

b06_texts = []
in_c4 = False
for p in doc_06.paragraphs:
    t = p.text.strip()
    if 'CAPÍTULO IV' in t:
        in_c4 = True
    elif 'CAPÍTULO V' in t:
        in_c4 = False
        break
    if in_c4 and t:
        b06_texts.append(t)

print(f"Non-empty paras in dev_doc: {len(dev_texts)}")
print(f"Non-empty paras in backup_06: {len(b06_texts)}")
print(f"Are they identical? {dev_texts == b06_texts}")
if dev_texts != b06_texts:
    for idx, (t1, t2) in enumerate(zip(dev_texts, b06_texts)):
        if t1 != t2:
            print(f"Diff at {idx}: dev='{t1[:50]}' vs b06='{t2[:50]}'")
            break

