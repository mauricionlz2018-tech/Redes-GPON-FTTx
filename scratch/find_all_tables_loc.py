import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')

for t_idx, tbl in enumerate(doc.tables):
    parent_idx = doc._body._element.index(tbl._element)
    prev_text = ""
    next_text = ""
    # find previous paragraph text
    for i in range(parent_idx - 1, -1, -1):
        el = doc._body._element[i]
        if el.tag.endswith('p'):
            prev_text = docx.text.paragraph.Paragraph(el, doc).text.strip()
            if prev_text:
                break
    for i in range(parent_idx + 1, len(doc._body._element)):
        el = doc._body._element[i]
        if el.tag.endswith('p'):
            next_text = docx.text.paragraph.Paragraph(el, doc).text.strip()
            if next_text:
                break
    c00 = tbl.cell(0, 0).text.strip().replace('\n', ' ')[:35]
    print(f"Tbl[{t_idx:2d}] at elem[{parent_idx:3d}]: first_cell='{c00}' | PREV='{prev_text[:50]}' | NEXT='{next_text[:50]}'")

