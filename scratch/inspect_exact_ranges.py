import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

def print_section_range(title_start, title_end):
    start_idx = None
    end_idx = None
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if start_idx is None and title_start in t:
            start_idx = i
        elif start_idx is not None and title_end in t:
            end_idx = i
            break
    print(f"\n=== {title_start} (P[{start_idx}] to P[{end_idx}]) ===")
    if start_idx is not None:
        end = end_idx if end_idx is not None else start_idx + 15
        for j in range(start_idx, min(end + 1, len(doc.paragraphs))):
            print(f"P[{j}] ({doc.paragraphs[j].style.name}): {doc.paragraphs[j].text[:80]}")

print_section_range('3.2.2 Diagrama General de Casos de Uso', '3.2.3 Requerimientos Funcionales')
print_section_range('3.2.6 Diagrama de Flujo', '3.3 Arquitectura topol')
print_section_range('3.4.2 Diagrama Entidad-Relaci', '3.5 Ingenier')

