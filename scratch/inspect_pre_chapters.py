import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')

print("=== INSPECTING PRE_RESTRUCTURE ===")
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
print("Sections:", len(doc.sections))

# Let's inspect where the chapters are
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if any(k in t.upper() for k in ['CAPITULO', 'CAPÍTULO', 'CONCLUSIONES', 'ANEXOS', 'REFERENCIAS']):
        print(f"P[{i}] ({p.style.name}): '{t[:80]}'")

