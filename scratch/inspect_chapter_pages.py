import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

chapters_info = [
    (195, "CAPITULO I"),
    (250, "CAPITULO II"),
    (759, "CAPITULO III"),
    (1198, "CAPITULO IV"),
    (1205, "CAPITULO V")
]

for p_idx, name in chapters_info:
    print(f"=== {name} at P{p_idx} ===")
    for i in range(p_idx, p_idx + 18):
        if i >= len(doc.paragraphs): break
        p = doc.paragraphs[i]
        has_br = 'w:br' in p._element.xml
        has_pb = 'pageBreakBefore' in p._element.xml
        print(f"  P{i}: {repr(p.text[:40])} (br={has_br}, pb_before={has_pb})")
        if p.style.name.startswith('Heading 2') or (p.text.strip() and not p.text.startswith('CAP')):
            # First text
            pass

