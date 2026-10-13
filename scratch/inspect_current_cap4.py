import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

c4_idx = None
c5_idx = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO IV' in p.text.upper():
        c4_idx = i
    if 'CAPÍTULO V' in p.text.upper():
        c5_idx = i
        break

print(f"Capítulo IV en P{c4_idx}, Capítulo V en P{c5_idx}")
for k in range(c4_idx - 5, min(c5_idx + 5, len(doc.paragraphs))):
    p = doc.paragraphs[k]
    has_sp = bool(p._element.xpath('./w:pPr/w:sectPr'))
    print(f"P{k} [{p.style.name}, sectPr={has_sp}]: '{p.text.strip()[:60]}'")

