import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

print("=== CHECKING ALL FIGURES IN CHAPTER III ===")
for i in range(728, 1015):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if t.startswith('Figura '):
        # find drawing in prev paragraph
        prev_p = doc.paragraphs[i-1]
        drawings = prev_p._p.xpath('.//w:drawing')
        blips = prev_p._p.xpath('.//a:blip/@r:embed')
        # check relationship
        rId = blips[0] if blips else 'None'
        target = 'None'
        if rId != 'None' and rId in doc.part.rels:
            target = doc.part.rels[rId].target_ref
        print(f"P[{i}]: {t[:70]} | rId={rId} -> {target}")

