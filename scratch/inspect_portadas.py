import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')
print("Total elements in body:", len(doc._body._element))
print("Tables:", len(doc.tables))

for idx, child in enumerate(doc._body._element[:25]):
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = docx.text.paragraph.Paragraph(child, doc)
        t = p.text.strip()
        print(f"[{idx}] <p> style='{p.style.name}': {t[:60]}")
    elif tag == 'tbl':
        tbl = docx.table.Table(child, doc)
        print(f"[{idx}] <tbl> rows={len(tbl.rows)}, cols={len(tbl.columns)}, cell00='{tbl.cell(0,0).text.strip().replace(chr(10), ' ')[:40]}'")
    else:
        print(f"[{idx}] <{tag}>")

