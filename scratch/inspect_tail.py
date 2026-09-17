import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')
for i in range(840, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    t = p.text.strip()
    xml = p._p.xml
    is_drawing = '<w:drawing' in xml or '<w:pict' in xml
    print(f"P[{i}] [DRAWING={is_drawing}]: '{t[:80]}'")

