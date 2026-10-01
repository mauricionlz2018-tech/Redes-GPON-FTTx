import os
import re
import copy
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_ADVANCE = 'docs/Documentacion_Residencias_avance.docx'
DOC_MIRROR = 'docs/Documentacion_Residencias (4).docx'

print("=== STARTING MOCKUPS INSERTION AND 82 FIGURES RENUMBERING ===")

# Copy images to artifacts directory as well
ARTIFACT_DIR = r'C:/Users/karen/.gemini/antigravity/brain/a33b5b22-c89b-4ff3-a6cc-733a0c69ceac'
for fn in [
    'mockup_vista_mapa_gis_puertos.png',
    'mockup_vista_padron_abonados.png',
    'mockup_vista_reportes_saturacion.png',
    'mockup_vista_gestion_personal.png'
]:
    src = os.path.join('scratch', fn)
    dst = os.path.join(ARTIFACT_DIR, fn)
    shutil.copyfile(src, dst)
    print(f"Copied {src} -> {dst}")

doc = docx.Document(DOC_ADVANCE)

# 1. Define complete 82 Figures metadata
FIGURES_METADATA = [
    # Cap I & II (Figures 1 to 26)
    (1, "Ubicación de la empresa GPON TELECOM S.A de C.V", "8"),
    (2, "Organigrama de la empresa GPON TELECOM S.A de C.V.", "9"),
    (3, "Modelo de referencia OSI de 7 capas frente a la arquitectura TCP/IP.", "13"),
    (4, "Principio físico de propagación: Refracción, ángulo crítico y reflexión interna total (TIR).", "15"),
    (5, "Protocolo UDP.", "17"),
    (6, "Paradigma de FTT existente actualmente.", "19"),
    (7, "Espectro y plan de longitudes de onda en GPON (ITU-T G.984.2 WDM).", "21"),
    (8, "Arquitectura de superposición de capas en un Sistema de Información Geográfica (GIS).", "26"),
    (9, "Librería Leaflet para mapas en react.", "28"),
    (10, "Base de datos PostgreSQL.", "30"),
    (11, "¿Qué es un ORM?", "32"),
    (12, "Logo de Sequelize.js", "32"),
    (13, "Arquitectura Cliente-Servidor.", "37"),
    (14, "Integraciones y entregas continuas.", "55"),
    (15, "Metodología de cascada.", "57"),
    (16, "Ejemplo de diagrama de casos de uso.", "59"),
    (17, "Arquitectura Modelo-Vista-Controlador.", "61"),
    (18, "Ejemplo de Responsive Web Design y Mobile-First Web Design.", "64"),
    (19, "Interfaz de Adobe Color.", "66"),
    (20, "Tipos de modelos de caja (Caja negra y Caja blanca).", "68"),
    (21, "Pruebas de API REST.", "69"),
    (22, "Pasos para realizar un test de usabilidad (ejemplo).", "70"),
    (23, "Diferencia en BSS y OSS.", "71"),
    (24, "Ejemplo de sistema de inventario de redes de telecomunicaciones.", "72"),
    (25, "Cuestionario en Google Forms sobre redes GPON/FTTx para el levantamiento de requerimientos.", "76"),
    (26, "Modelado de actores y sus funciones.", "78"),

    # Cap III - Requerimientos & Diagramas de Casos de Uso (27 to 30)
    (27, "Diagrama de casos de uso - Módulo 1: Seguridad, autenticación y control de acceso (RBAC).", "79"),
    (28, "Diagrama de casos de uso - Módulo 2: Cartografía GIS, trazado troncal y cajas terminales NAP.", "80"),
    (29, "Diagrama de casos de uso - Módulo 3: Operación de puertos físicos, concurrencia ACID y abonados.", "81"),
    (30, "Diagrama de casos de uso - Módulo 4: Operación móvil Offline-First y generación de reportes ejecutivos PDF.", "82"),

    # Cap III - Robustez, Flujo, Actividades (31 to 33)
    (31, "Diagrama de robustez (V-O-C de Jacobson) para la asignación concurrente de puertos y control transaccional.", "84"),
    (32, "Diagrama de flujo general de operación y procesos del sistema GPON.", "89"),
    (33, "Diagrama de actividades UML (Swimlanes) para el proceso de asignación y provisión concurrente de puertos.", "91"),

    # Cap III - Topología, Clases, DER (34 to 37)
    (34, "Cadena de distribución de la red GPON desde la cabecera central (NOC) hasta la acometida domiciliaria (ONT).", "93"),
    (35, "Topología jerárquica de la red óptica GPON/FTTx y matriz de presupuesto óptico de potencia (ITU-T G.984.2).", "95"),
    (36, "Diagrama de clases del dominio y entidades del sistema bajo estándar UML 2.5.", "97"),
    (37, "Diagrama Entidad-Relación (DER) conceptual con notación Chen del sistema de inventario GPON.", "99"),

    # Cap III - Secuencia concurrencia, FSM estados, Secuencia offline (38 to 40)
    (38, "Diagrama de secuencia transaccional de concurrencia con SELECT ... FOR UPDATE.", "112"),
    (39, "Diagrama de máquina de estados finitos (FSM) del ciclo de vida del puerto óptico SC-APC.", "113"),
    (40, "Diagrama de secuencia UML para sincronización diferida y arquitectura móvil Offline-First.", "114"),

    # Cap III - SECTION 3.6.1 Maquetado Wireframes Figma (41)
    (41, "Maquetado de interfaces (Wireframes) y arquitectura visual del sistema en Figma.", "115"),

    # Cap III - SECTION 3.6.1 NUEVAS VISTAS DE MAQUETADO DE ALTA FIDELIDAD (42 to 45)
    (42, "Maquetado de interfaz — Visor cartográfico GIS interactivo y modal transaccional de inspección de puertos NAP.", "116"),
    (43, "Maquetado de interfaz — Padrón consolidado de abonados FTTx y gestión de expedientes técnicos.", "117"),
    (44, "Maquetado de interfaz — Consola analítica de reportes ejecutivos e indicadores de saturación de red GPON.", "118"),
    (45, "Maquetado de interfaz — Módulo de gestión centralizada de personal operativo y control de acceso (RBAC).", "119"),

    # Cap III - SECTION 3.6.1 Diagramas de Navegación, IA, Paquetes (Old 42, 43, 44 -> 46, 47, 48)
    (46, "Diagrama de navegación del sistema y flujo heurístico de interfaces de usuario.", "120"),
    (47, "Diagrama jerárquico de arquitectura de información (IA) del sistema web y móvil FTTx.", "121"),
    (48, "Diagrama de paquetes y componentes de software bajo estándar UML.", "122"),

    # Cap III - SECTION 3.6.2 Pruebas de Contraste Adobe Color (Old 45 to 53 -> 49 to 57)
    (49, "Prueba de colores #FFFFFF y #4F46ES", "124"),
    (50, "Prueba de colores para Botones Soporte / Técnico", "124"),
    (51, "Pruebas de color para Badge \"Datos de Prueba\"", "125"),
    (52, "Contraste para Pestaña \"Mapa de Red\" (Activa)", "125"),
    (53, "Contraste de pestañas \"Abonados\" / \"Reportes\"", "125"),
    (54, "Contraste de Botón \"APK\"", "126"),
    (55, "Contraste de color Tag Rol \"ADMIN\"", "126"),
    (56, "Contraste de Botón \"+ Troncal / Ramal\"", "127"),
    (57, "Contraste de Botón \"+ Mufa\" (Empalme)", "127"),

    # Cap IV - Implementación y Desarrollo de Código (Old 54 to 61 -> 58 to 65)
    (58, "Configuración de la conexión a PostgreSQL con Sequelize (database.ts).", "129"),
    (59, "Modelo Client.ts y estructura de la capa de modelos (backend/src/models/).", "130"),
    (60, "Middleware de autenticación con verificación de token JWT (auth.ts).", "131"),
    (61, "Controlador de asignación concurrente de puertos con transacción pesimista ACID (portController.ts).", "132"),
    (62, "Definición de rutas y endpoints de la API REST con middlewares de seguridad y RBAC (api.ts).", "133"),
    (63, "Interceptor de Axios para la inyección del token JWT en el frontend (client.ts).", "135"),
    (64, "Inicialización del servidor Express, configuración CORS y conexión Sequelize (index.ts).", "137"),
    (65, "Cliente HTTP Axios con interceptor de autorización Bearer JWT y manejo de sesión (client.ts).", "139"),

    # Cap IV - Cartografía Leaflet y Estados de Puertos (Old 62 to 71 -> 66 to 75)
    (66, "Marcador de caja NAP con ocupación inferior al 80% (NAP-SJR-26, 0/8).", "141"),
    (67, "Marcador de caja NAP en umbral preventivo, con ocupación entre 80% y 99% (NAP-SJR-01, 14/16).", "141"),
    (68, "Marcador de caja NAP saturada, con ocupación del 100% (NAP-SJR-02, 16/16).", "141"),
    (69, "Ventana emergente con el detalle de un cable troncal trazado como polilínea en el mapa.", "142"),
    (70, "Modal de captura de coordenadas GPS en campo (GpsCaptureModal.tsx).", "143"),
    (71, "Puertos de la matriz NAP en estado Libre.", "143"),
    (72, "Puertos de la matriz NAP en estado Ocupado.", "144"),
    (73, "Puerto de la matriz NAP en estado Dañado.", "144"),
    (74, "Puerto de la matriz NAP en estado Reservado.", "144"),
    (75, "Puertos libres y ocupados.", "145"),

    # Cap IV - Frontend, Offline PWA, Reportes PDF, Docker (Old 72 to 78 -> 76 to 82)
    (76, "Arquitectura de componentes frontend React y visor cartográfico.", "145"),
    (77, "Flujo de decisión y sincronización diferida de la arquitectura móvil Offline-First.", "148"),
    (78, "Esquema de persistencia local IndexedDB con Dexie.js para trabajo sin conexión (offlineDb.ts).", "149"),
    (79, "Encabezado institucional y resumen del reporte ejecutivo de auditoría de red en PDF.", "150"),
    (80, "Inventario y nivel de saturación por caja NAP en el reporte ejecutivo en PDF.", "151"),
    (81, "Flujo de generación de reportes técnicos ejecutivos en PDF mediante streaming en memoria.", "152"),
    (82, "Arquitectura de contenerización multicontenedor con Docker Compose y red aislada.", "154"),
]

assert len(FIGURES_METADATA) == 82, f"Expected 82 figures, got {len(FIGURES_METADATA)}"

# Locate the target paragraphs in Section 3.6.1
p_fig41_cap = None
p_bullets = []
p_nav_intro = None

for i, p in enumerate(doc.paragraphs[117:], start=117):
    if 'Figura 41. Maquetado de interfaces' in p.text:
        p_fig41_cap = p
    elif '1.- Wireframe 1: Visor Cartogr' in p.text:
        p_bullets.append(p)
    elif '2.- Wireframe 2: Matriz Isom' in p.text:
        p_bullets.append(p)
    elif '3.- Wireframe 3: Arquitectura PWA' in p.text:
        p_bullets.append(p)
    elif '4.- Wireframe 4: Padr' in p.text:
        p_bullets.append(p)
    elif 'A partir de los esquemas de maquetado aprobados, se formaliz' in p.text:
        p_nav_intro = p
        break

print(f"Target Figura 41: {p_fig41_cap.text[:50]}")
print(f"Found {len(p_bullets)} bullet paragraphs to replace.")
print(f"Target Navigation intro: {p_nav_intro.text[:50]}")

# Helper to build figure elements attached to doc
def build_mockup_block(intro_text, img_path, caption_title):
    # 1. Intro description
    p_intro = doc.add_paragraph()
    p_intro.style = 'Normal'
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.space_before = Pt(6)
    p_intro.paragraph_format.space_after = Pt(6)
    r_intro = p_intro.add_run(intro_text)
    r_intro.font.name = 'Arial'
    r_intro.font.size = Pt(11)

    # 2. Image
    p_img = doc.add_paragraph()
    p_img.style = 'Normal'
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(4)
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=Inches(6.05))

    # 3. Caption
    p_cap = doc.add_paragraph()
    p_cap.style = 'Figuras'
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(2)
    r_cap = p_cap.add_run(caption_title)
    r_cap.font.name = 'Arial'
    r_cap.font.size = Pt(9)
    r_cap.bold = True

    # 4. Note
    p_note = doc.add_paragraph()
    p_note.style = 'Normal'
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(12)
    r_note = p_note.add_run("Elaboración propia.")
    r_note.font.name = 'Arial'
    r_note.font.size = Pt(9)
    r_note.bold = True

    return [p_intro, p_img, p_cap, p_note]

# The 4 Mockups technical descriptions
desc_mockup_1 = (
    "La interfaz del Visor Cartográfico GIS constituye el núcleo geoespacial operativo de la plataforma, "
    "diseñada bajo principios de interacción directa y visualización por capas temáticas (OpenStreetMap y cartografía vectorial). "
    "En este entorno, las cuadrillas y operadores visualizan la traza georreferenciada de los cables troncales ADSS y ramales "
    "secundarios de fibra óptica que interconectan la cabecera ODF municipal con los diferentes cierres y cajas terminales. "
    "Los nodos NAP se representan mediante marcadores circulares vectoriales cuyo color varía dinámicamente de acuerdo con el nivel "
    "de saturación de capacidad: verde para ocupación operativa menor al 80%, ámbar preventivo para ocupación entre 80% y 99%, y rojo "
    "crítico para saturación total (100%). Al interactuar sobre cualquier marcador de caja NAP en el mapa, el sistema abre de manera no intrusiva "
    "el modal transaccional de inspección física, el cual despliega una matriz isomórfica 2×8 fiel al chasis de 16 puertos SC-APC. "
    "En el maquetado presentado se ilustra la inspección en tiempo real de la caja 'NAP-SJR-01' con el Puerto #3 en estado 'Ocupado', "
    "visualizando de forma inmediata el expediente del abonado asignado (ID de contrato, nombre completo, modelo de ONT Huawei HG8245H, "
    "dirección MAC física y nivel de potencia óptica calibrado en -19.4 dBm). Para puertos disponibles, la interfaz habilita el formulario "
    "de aprovisionamiento inmediato gobernado por la cláusula atómica SELECT ... FOR UPDATE a nivel de base de datos."
)

desc_mockup_2 = (
    "La interfaz del Padrón de Abonados Conectados concentra la administración centralizada de la cartera de suscriptores activos, "
    "suspendidos y en proceso de instalación sobre la infraestructura de fibra óptica. El diseño implementa una tabla de datos responsiva "
    "de alta densidad informacional que integra un motor de búsqueda reactivo en tiempo real con capacidad de filtrado multicriterio "
    "por número de contrato, nombre o razón social del cliente, caja NAP asignada, identificador de puerto físico y dirección MAC de la "
    "terminal de red óptica (ONT). Cada registro de la grilla expone el estado operativo del servicio mediante insignias cromáticas de alto contraste, "
    "fecha de alta transaccional, nivel de potencia óptica registrada (-dBm) y atajos de acción rápida para la edición del expediente, consulta "
    "de ubicación geográfica en mapa o inicio de procedimientos de baja y mantenimiento de enlace. Asimismo, la cabecera del maquetado "
    "incorpora controles de exportación rápida y advertencias de entorno de auditoría, optimizando la gestión de inventario y eliminando "
    "discrepancias entre el padrón comercial y el estado físico de los puertos en planta externa."
)

desc_mockup_3 = (
    "La Consola de Reportes e Indicadores de Saturación ofrece un tablero de mando integral (Dashboard analítico) para la supervisión directiva "
    "y técnica del despliegue FTTx en San José del Rincón. La interfaz estructura métricas ejecutivas de primer orden mediante tarjetas de indicadores "
    "clave de rendimiento (KPIs), cuantificando en tiempo real el total de terminales NAP desplegadas, la capacidad nominal en puertos ópticos, "
    "los puertos efectivamente aprovisionados y el porcentaje global de saturación de la infraestructura. En el cuerpo central, el maquetado "
    "dispone una matriz de criticidad semafórica que clasifica las cajas terminales en tres cuadrantes operativos: terminales con holgura de servicio "
    "(<80% de ocupación en verde), terminales en umbral preventivo de ampliación (80% a 99% en amarillo) y cajas con saturación absoluta (100% en rojo, "
    "sin capacidad de venta adicional). Esta visualización jerárquica previene la sobreventa de acometidas en zonas de alta densidad poblacional y "
    "se complementa con un módulo de generación de reportes técnicos auditables exportables a formato PDF mediante streaming asíncrono en memoria."
)

desc_mockup_4 = (
    "El módulo de Gestión de Personal y Cuadrillas Técnicas formaliza la administración de identidades, credenciales y privilegios operativos "
    "bajo el modelo de Control de Acceso Basado en Roles (Role-Based Access Control - RBAC). La interfaz organiza a los operadores en una "
    "cuadrícula administrativa que desglosa el identificador de usuario, nombre del técnico, correo electrónico corporativo, perfil funcional "
    "asignado (Administrador General, Supervisor de Planta Externa o Técnico Instalador de Campo) y estado de activación de la cuenta. A través "
    "de este panel, los supervisores asignan cuadrillas de trabajo a zonas geográficas específicas, gestionan el restablecimiento seguro de claves "
    "mediante mecanismos de cifrado robusto y auditan las sesiones activas en la plataforma. Esta interfaz asegura que únicamente el personal debidamente "
    "certificado pueda ejecutar mutaciones transaccionales en la topología de la red óptica o autorizar reconexiones de fibra, salvaguardando la "
    "integridad de los datos de inventario y cumpliendo con las políticas de ciberseguridad institucional."
)

# Build blocks
b1 = build_mockup_block(desc_mockup_1, 'scratch/mockup_vista_mapa_gis_puertos.png', 'FIGURA_DUMMY_42')
b2 = build_mockup_block(desc_mockup_2, 'scratch/mockup_vista_padron_abonados.png', 'FIGURA_DUMMY_43')
b3 = build_mockup_block(desc_mockup_3, 'scratch/mockup_vista_reportes_saturacion.png', 'FIGURA_DUMMY_44')
b4 = build_mockup_block(desc_mockup_4, 'scratch/mockup_vista_gestion_personal.png', 'FIGURA_DUMMY_45')

all_mockup_elems = b1 + b2 + b3 + b4

# Replace the text of P[989] (the intro to wireframes)
intro_p = None
for i, p in enumerate(doc.paragraphs[117:], start=117):
    if 'A partir de estos esquemas se consolidaron las siguientes interfaces funcionales' in p.text:
        intro_p = p
        break

if intro_p is not None:
    intro_p.text = (
        "A partir de estos esquemas conceptuales y del análisis de requerimientos en campo, se desarrollaron los "
        "maquetados interactivos de alta fidelidad correspondientes a las vistas neurálgicas de la plataforma, "
        "detalladas individualmente a continuación:"
    )
    intro_p.style = 'Normal'
    intro_p.paragraph_format.line_spacing = 1.5
    intro_p.paragraph_format.space_before = Pt(6)
    intro_p.paragraph_format.space_after = Pt(6)
    for r in intro_p.runs:
        r.font.name = 'Arial'
        r.font.size = Pt(11)

# Remove old bullet paragraphs
for p in p_bullets:
    p._p.getparent().remove(p._p)

# Insert all mockup blocks right before p_nav_intro
curr_xml = p_nav_intro._p
for el in reversed(all_mockup_elems):
    curr_xml.addprevious(el._p)

print("Inserted all 4 Mockup blocks successfully into Section 3.6.1!")

# 2. Identify all figure captions in the document body (from p[117] onwards)
print("\n=== IDENTIFYING ALL FIGURE CAPTIONS IN BODY ===")
body_after_insert = doc.paragraphs[117:]
figure_captions = []
for p in body_after_insert:
    txt = p.text.strip()
    if p.style.name == 'Figuras' and '\t' not in txt:
        figure_captions.append(p)
    elif txt.startswith("FIGURA_DUMMY_"):
        figure_captions.append(p)

print(f"Total figure captions detected in body: {len(figure_captions)}")
if len(figure_captions) != 82:
    print(f"ERROR: Expected 82 captions, found {len(figure_captions)}")
    for i, p in enumerate(figure_captions):
        print(f"  [{i}]: {p.text[:60]}")
    raise RuntimeError("Figure caption count mismatch!")

# 3. Clean and renumber all 82 figure captions in the body
print("\n=== RENUMBERING ALL 82 FIGURE CAPTIONS IN BODY ===")
for k, p_cap in enumerate(figure_captions):
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_id = str(400 + k)
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    # Clean the paragraph XML
    p_elem = p_cap._p
    pPr = p_elem.find(qn('w:pPr'))
    for child in list(p_elem):
        if child != pPr:
            p_elem.remove(child)
            
    # Set style to Figuras
    p_cap.style = 'Figuras'
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(2)
    
    # Add bookmarkStart
    bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{bookmark_name}"/>')
    p_elem.append(bm_start)
    
    # Add text run
    r_elem = parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>'
        f'<w:b/><w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>'
        f'<w:t xml:space="preserve">{new_caption_text}</w:t></w:r>'
    )
    p_elem.append(r_elem)
    
    # Add bookmarkEnd
    bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>')
    p_elem.append(bm_end)

print("All 82 figure captions in body successfully renumbered and bookmarked!")

# 4. Update the Índice de figuras
print("\n=== UPDATING ÍNDICE DE FIGURAS ===")
idx_title_p = None
idx_end_p = None
for i, p in enumerate(doc.paragraphs[:120]):
    if 'ndice de figuras' in p.text:
        idx_title_p = i
    if 'ndice de tablas' in p.text:
        idx_end_p = i
        break

print(f"Índice de figuras starts at P[{idx_title_p}], Índice de tablas at P[{idx_end_p}]")

existing_toc_p = []
for i in range(idx_title_p + 1, idx_end_p):
    p = doc.paragraphs[i]
    if p.text.strip().startswith("Figura "):
        existing_toc_p.append(p)

print(f"Found {len(existing_toc_p)} existing TOC figure paragraphs.")

# Update existing entries
for k in range(min(len(existing_toc_p), 82)):
    p = existing_toc_p[k]
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    # Update hyperlink anchor
    hl = p._p.xpath('.//w:hyperlink')
    if hl:
        hl[0].attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor'] = bookmark_name
        
    # Update PAGEREF instruction
    instr = p._p.xpath('.//w:instrText')
    for inst in instr:
        if 'PAGEREF' in inst.text:
            inst.text = f" PAGEREF {bookmark_name} \\h "
            
    # Update text nodes
    t_nodes = p._p.xpath('.//w:hyperlink//w:t')
    if len(t_nodes) >= 2:
        t_nodes[0].text = new_caption_text
        t_nodes[-1].text = fig_page
    elif len(t_nodes) == 1:
        t_nodes[0].text = new_caption_text

# If there are more figures than existing TOC paragraphs (82 vs 78, difference 4)
last_toc_p = existing_toc_p[-1]
for k in range(len(existing_toc_p), 82):
    fig_num, fig_title, fig_page = FIGURES_METADATA[k]
    new_caption_text = f"Figura {fig_num}. {fig_title}"
    bookmark_name = f"_Toc240964{71 + k:03d}"
    
    new_p_xml = copy.deepcopy(last_toc_p._p)
    
    # Update anchor
    hl = new_p_xml.xpath('.//w:hyperlink')
    if hl:
        hl[0].attrib['{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor'] = bookmark_name
        
    # Update PAGEREF
    instr = new_p_xml.xpath('.//w:instrText')
    for inst in instr:
        if 'PAGEREF' in inst.text:
            inst.text = f" PAGEREF {bookmark_name} \\h "
            
    # Update text nodes
    t_nodes = new_p_xml.xpath('.//w:hyperlink//w:t')
    if len(t_nodes) >= 2:
        t_nodes[0].text = new_caption_text
        t_nodes[-1].text = fig_page
        
    last_toc_p._p.addnext(new_p_xml)
    last_toc_p = docx.text.paragraph.Paragraph(new_p_xml, doc)

print("Successfully updated and appended all 82 entries in Índice de figuras!")

# 5. Save modified document to advance and mirror files
print("\n=== SAVING DOCUMENTS ===")
doc.save(DOC_ADVANCE)
print(f"Saved: {DOC_ADVANCE}")

doc.save(DOC_MIRROR)
print(f"Saved: {DOC_MIRROR}")

print("\n=== VERIFICATION RELOAD ===")
doc_check = docx.Document(DOC_ADVANCE)
check_body = doc_check.paragraphs[117:]
verified_captions = [p.text.strip() for p in check_body if p.style.name == 'Figuras' and '\t' not in p.text]
print(f"Verified body figure captions count: {len(verified_captions)}")
print(f"Fig 1: {verified_captions[0]}")
print(f"Fig 41: {verified_captions[40]}")
print(f"Fig 42: {verified_captions[41]}")
print(f"Fig 43: {verified_captions[42]}")
print(f"Fig 44: {verified_captions[43]}")
print(f"Fig 45: {verified_captions[44]}")
print(f"Fig 46: {verified_captions[45]}")
print(f"Fig 47: {verified_captions[46]}")
print(f"Fig 48: {verified_captions[47]}")
print(f"Fig 82: {verified_captions[-1]}")

check_toc = []
for p in doc_check.paragraphs[idx_title_p+1:idx_title_p+95]:
    txt = p.text.strip()
    if txt.startswith("Figura "):
        check_toc.append(txt.split('\t')[0])

print(f"Verified TOC entries count: {len(check_toc)}")
print(f"TOC 1: {check_toc[0]}")
print(f"TOC 41: {check_toc[40]}")
print(f"TOC 42: {check_toc[41]}")
print(f"TOC 43: {check_toc[42]}")
print(f"TOC 44: {check_toc[43]}")
print(f"TOC 45: {check_toc[44]}")
print(f"TOC 46: {check_toc[45]}")
print(f"TOC 82: {check_toc[-1]}")

print("\n=== SUCCESSFUL COMPLETION! ===")

