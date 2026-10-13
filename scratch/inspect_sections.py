import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')
print("Total paragraphs:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    sp = p._element.xpath('./w:pPr/w:sectPr')
    if sp:
        print(f"P{i}: '{p.text[:60]}'")
print("Body sectPr:", bool(doc._body._element.xpath('./w:sectPr')))

