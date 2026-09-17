import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
for i, p in enumerate(doc.paragraphs):
    if '<w:drawing' in p._p.xml or '<w:pict' in p._p.xml:
        next_t = doc.paragraphs[i+1].text if i+1 < len(doc.paragraphs) else ''
        print(f"P[{i}]: next='{next_t[:80]}'")

