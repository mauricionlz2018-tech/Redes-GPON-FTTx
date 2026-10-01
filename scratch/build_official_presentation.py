import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import io

def set_font(run, name="Arial", size_pt=12, bold=False, italic=False, color_rgb=(30, 41, 59)):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def update_presentation():
    src_ppt = "docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx"
    prs = pptx.Presentation(src_ppt)
    print(f"Loaded {src_ppt} with {len(prs.slides)} slides.")

    # Colors
    c_guinda = (122, 28, 53)    # #7A1C35
    c_dark = (30, 41, 59)       # #1E293B
    c_gray = (100, 116, 139)    # #64748B
    c_gold = (163, 131, 64)     # #A38340
    c_white = (255, 255, 255)

    # ==========================================
    # SLIDE 1: Carátula
    # ==========================================
    s1 = prs.slides[0]
    for shp in s1.shapes:
        if shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                if "Septiembre de 2026" in p.text:
                    p.text = "Octubre de 2026"
                    for r in p.runs:
                        set_font(r, name="Arial", size_pt=14, bold=True, color_rgb=c_guinda)
                elif "Presenta:" in p.text or "Mauricio" in p.text:
                    pass
    print("Slide 1 updated: Date set to Octubre de 2026.")

    # ==========================================
    # SLIDE 2: Agenda
    # ==========================================
    s2 = prs.slides[1]
    for shp in s2.shapes:
        if shp.has_text_frame and "Institución" in shp.text:
            tf = shp.text_frame
            tf.clear()
            agenda_items = [
                "1. Datos de la Empresa (Giro y Ubicación)",
                "2. Organigrama y Área del Residente",
                "3. Introducción",
                "4. Planteamiento del Problema",
                "5. Justificación",
                "6. Objetivos del Proyecto (General y Específicos)",
                "7. Herramientas de Desarrollo",
                "8. Evidencias de Avance: Documentación Formal (PDF)",
                "9. Evidencias de Avance: Proyecto y Aplicación Web",
                "10. Cronograma de Actividades (Avance al Corte)",
                "11. Conclusiones Preliminares",
                "12. Bibliografía, Referencias y Preguntas"
            ]
            for idx, item in enumerate(agenda_items):
                p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
                p.text = item
                p.space_after = Pt(4)
                for r in p.runs:
                    set_font(r, name="Arial", size_pt=13, bold=False, color_rgb=c_dark)
    print("Slide 2 updated: Agenda aligned to 12 official elements.")

    # ==========================================
    # SLIDE 8: Objetivos (General y Específicos combinados)
    # ==========================================
    s8 = prs.slides[7]
    for shp in s8.shapes:
        if shp.has_text_frame and "Objetivo General:" in shp.text:
            tf = shp.text_frame
            tf.clear()

            # Objetivo General
            p0 = tf.paragraphs[0]
            p0.text = "Objetivo General:"
            p0.space_after = Pt(2)
            for r in p0.runs:
                set_font(r, name="Arial", size_pt=13, bold=True, color_rgb=c_guinda)

            p1 = tf.add_paragraph()
            p1.text = (
                "Desarrollar una aplicación web para el inventario y mapeo lógico de redes GPON/FTTx, "
                "mediante el diseño de vistas tabulares y diagramas lógicos simples, para optimizar "
                "el control de la infraestructura por parte del equipo técnico y de soporte."
            )
            p1.space_after = Pt(12)
            for r in p1.runs:
                set_font(r, name="Arial", size_pt=11.5, bold=False, color_rgb=c_dark)

            # Objetivos Específicos
            p2 = tf.add_paragraph()
            p2.text = "Objetivos Específicos:"
            p2.space_after = Pt(2)
            for r in p2.runs:
                set_font(r, name="Arial", size_pt=13, bold=True, color_rgb=c_guinda)

            especificos = [
                "1. Analizar los requerimientos operativos de la red de fibra óptica mediante el levantamiento de información sobre el hardware (NAP/ODF), para diseñar la arquitectura del sistema y el modelo de la base de datos.",
                "2. Construir una interfaz web responsiva a través del uso de tecnologías de desarrollo frontend, para que los técnicos de campo actualicen desde su celular el estado de los puertos (libres, ocupados, dañados).",
                "3. Programar el backend del sistema mediante la creación de una API centralizada, para gestionar la información de los activos y vincular los puertos de red correspondientes a los datos de los clientes."
            ]
            for esp in especificos:
                p_esp = tf.add_paragraph()
                p_esp.text = esp
                p_esp.space_after = Pt(5)
                for r in p_esp.runs:
                    set_font(r, name="Arial", size_pt=10.5, bold=False, color_rgb=c_dark)
    print("Slide 8 updated: Objectives consolidated.")

    # ==========================================
    # SLIDE 9: Herramientas de Desarrollo (Moved from 10)
    # ==========================================
    # In the original final PPT:
    # Slide 9 had "Objetivos Específicos"
    # Slide 10 had "Herramientas"
    # Slide 11 had "Evidencias" (text only)
    #
    # We will repurpose Slide 9 to be HERRAMIENTAS!
    # Copy shapes/pictures from Slide 10 to Slide 9, or repurpose:
    # Actually, let's make Slide 9 hold Herramientas,
    # Slide 10 hold Evidencia A (Documentación PDF),
    # Slide 11 hold Evidencia B (Proyecto y Aplicación Web)!
    s9 = prs.slides[8]
    s10 = prs.slides[9]
    s11 = prs.slides[10]

    # Let's inspect shapes on s10 (Herramientas) and transfer them or swap
    # Instead of manual copying, let's copy s10's contents into s9,
    # then make s10 hold Evidencia A, and s11 hold Evidencia B!
    
    # Let's clear s9's specific content box
    for shp in s9.shapes:
        if shp.has_text_frame and "Objetivos Específicos" in shp.text:
            s9.shapes._spTree.remove(shp._element)

    # Update s9 Title banner to "Herramientas"
    for shp in s9.shapes:
        if shp.has_text_frame and "Objetivos" in shp.text:
            shp.text_frame.text = "Herramientas"
            for r in shp.text_frame.paragraphs[0].runs:
                set_font(r, name="Arial", size_pt=18, bold=True, color_rgb=c_white)
        elif shp.has_text_frame and "9/15" in shp.text:
            shp.text_frame.text = "9/15"

    # Add description in s9 for Herramientas
    tb_h = s9.shapes.add_textbox(Inches(0.70), Inches(1.05), Inches(8.60), Inches(0.60))
    p_h = tb_h.text_frame.paragraphs[0]
    p_h.text = "Para el desarrollo de la plataforma web se seleccionó un stack tecnológico desacoplado, moderno y de alto rendimiento:"
    for r in p_h.runs:
        set_font(r, name="Arial", size_pt=11.5, bold=False, color_rgb=c_dark)

    # Move pictures from s10 into s9
    for shp in list(s10.shapes):
        if shp.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE and shp.name != "Google Shape;166;p10":
            # Add to s9
            try:
                img_blob = shp.image.blob
                s9.shapes.add_picture(
                    io.BytesIO(img_blob),
                    shp.left, shp.top, shp.width, shp.height
                )
            except Exception as e:
                print("Error copying pic to s9:", e)

    print("Slide 9 updated: Now holds Herramientas de Desarrollo.")

    # ==========================================
    # SLIDE 10: Evidencias de Avance: a. Documentación Formal (PDF)
    # ==========================================
    # Clear s10 and rebuild for Evidence A
    for shp in list(s10.shapes):
        if shp.name != "Google Shape;166;p10": # keep background ribbon
            s10.shapes._spTree.remove(shp._element)

    # Add Title banner
    title_box_10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.80), Inches(0.01), Inches(4.20), Inches(0.91))
    title_box_10.fill.solid()
    title_box_10.fill.fore_color.rgb = RGBColor(*c_guinda)
    title_box_10.line.color.rgb = RGBColor(*c_guinda)
    p_t10 = title_box_10.text_frame.paragraphs[0]
    p_t10.text = "Evidencias: Documentación"
    p_t10.alignment = PP_ALIGN.CENTER
    for r in p_t10.runs:
        set_font(r, name="Arial", size_pt=16, bold=True, color_rgb=c_white)

    # Badge 10/15
    badge_10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.25), Inches(7.13), Inches(0.75), Inches(0.37))
    badge_10.fill.solid()
    badge_10.fill.fore_color.rgb = RGBColor(*c_guinda)
    badge_10.line.color.rgb = RGBColor(*c_guinda)
    badge_10.text_frame.text = "10/15"
    badge_10.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for r in badge_10.text_frame.paragraphs[0].runs:
        set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=c_white)

    # Insert Evidence A Image (docs/evidencia_documentacion_formal_pdf.png)
    img_ev_doc = "docs/evidencia_documentacion_formal_pdf.png"
    if os.path.exists(img_ev_doc):
        s10.shapes.add_picture(img_ev_doc, Inches(0.70), Inches(1.10), Inches(8.60), Inches(4.65))

    # Caption / Footnote
    tb_c10 = s10.shapes.add_textbox(Inches(0.70), Inches(5.82), Inches(8.60), Inches(0.40))
    p_c10 = tb_c10.text_frame.paragraphs[0]
    p_c10.text = "Avance Formal en PDF: Memoria técnica oficial en 5 capítulos, 96 figuras consecutivas y 28 tablas normativas de ingeniería."
    p_c10.alignment = PP_ALIGN.CENTER
    for r in p_c10.runs:
        set_font(r, name="Arial", size_pt=10, bold=True, italic=True, color_rgb=c_guinda)
    print("Slide 10 updated: Holds Evidence A (Documentación Formal PDF).")

    # ==========================================
    # SLIDE 11: Evidencias de Avance: b. Proyecto y Software Funcional
    # ==========================================
    # Clear s11 and rebuild for Evidence B
    for shp in list(s11.shapes):
        if shp.name != "Google Shape;189;p11": # keep background ribbon
            s11.shapes._spTree.remove(shp._element)

    # Add Title banner
    title_box_11 = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.00), Inches(0.01), Inches(4.00), Inches(0.91))
    title_box_11.fill.solid()
    title_box_11.fill.fore_color.rgb = RGBColor(*c_guinda)
    title_box_11.line.color.rgb = RGBColor(*c_guinda)
    p_t11 = title_box_11.text_frame.paragraphs[0]
    p_t11.text = "Evidencias: Software"
    p_t11.alignment = PP_ALIGN.CENTER
    for r in p_t11.runs:
        set_font(r, name="Arial", size_pt=17, bold=True, color_rgb=c_white)

    # Badge 11/15
    badge_11 = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.25), Inches(7.13), Inches(0.75), Inches(0.37))
    badge_11.fill.solid()
    badge_11.fill.fore_color.rgb = RGBColor(*c_guinda)
    badge_11.line.color.rgb = RGBColor(*c_guinda)
    badge_11.text_frame.text = "11/15"
    badge_11.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for r in badge_11.text_frame.paragraphs[0].runs:
        set_font(r, name="Arial", size_pt=10, bold=True, color_rgb=c_white)

    # Insert Evidence B Image (docs/evidencia_proyecto_aplicacion_web.png)
    img_ev_app = "docs/evidencia_proyecto_aplicacion_web.png"
    if os.path.exists(img_ev_app):
        s11.shapes.add_picture(img_ev_app, Inches(0.70), Inches(1.10), Inches(8.60), Inches(4.65))

    # Caption / Footnote
    tb_c11 = s11.shapes.add_textbox(Inches(0.70), Inches(5.82), Inches(8.60), Inches(0.40))
    p_c11 = tb_c11.text_frame.paragraphs[0]
    p_c11.text = "Prototipo y Aplicación Web: Visor GIS Leaflet, matriz 2×8 de puertos NAP (1:16), padrón de abonados y consola de reportes ejecutivos."
    p_c11.alignment = PP_ALIGN.CENTER
    for r in p_c11.runs:
        set_font(r, name="Arial", size_pt=10, bold=True, italic=True, color_rgb=c_guinda)
    print("Slide 11 updated: Holds Evidence B (Proyecto y Software Funcional).")

    # ==========================================
    # SLIDE 12: Cronograma
    # ==========================================
    s12 = prs.slides[11]
    for shp in s12.shapes:
        if shp.has_text_frame and "Cronograma de actividades" in shp.text:
            shp.text_frame.text = "Figura 4. Cronograma de actividades del proyecto de residencia profesional (Corte: 40% de avance físico y metodológico)."
            for r in shp.text_frame.paragraphs[0].runs:
                set_font(r, name="Arial", size_pt=10, bold=True, italic=True, color_rgb=c_guinda)
    print("Slide 12 updated: Cronograma description refined.")

    # ==========================================
    # SLIDE 13: Conclusiones
    # ==========================================
    s13 = prs.slides[12]
    for shp in s13.shapes:
        if shp.has_text_frame and "Con base en el cronograma" in shp.text:
            tf = shp.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = (
                "Con base en el cronograma de actividades planificado, el proyecto registra actualmente un avance físico "
                "y metodológico del 40%, cumpliendo satisfactoriamente en tiempo y forma con las fases de levantamiento de "
                "infraestructura, análisis formal de requerimientos, diseño arquitectónico del software y prototipado del "
                "visor cartográfico y del backend transaccional. Los resultados preliminares confirman la viabilidad técnica "
                "y operativa de la plataforma web, resolviendo la trazabilidad geoespacial de planta externa y sentando las "
                "bases sólidas para culminar con éxito la implementación en campo y la validación con las cuadrillas técnicas "
                "en las semanas restantes."
            )
            for r in p.runs:
                set_font(r, name="Arial", size_pt=12.5, bold=False, color_rgb=c_dark)
    print("Slide 13 updated: Conclusions text confirmed.")

    # ==========================================
    # SLIDE 15: Cierre / Preguntas
    # ==========================================
    s15 = prs.slides[14]
    for shp in s15.shapes:
        if shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                if "Septiembre de 2026" in p.text:
                    p.text = "Octubre de 2026"
                    for r in p.runs:
                        set_font(r, name="Arial", size_pt=14, bold=True, color_rgb=c_guinda)
    print("Slide 15 updated: Date set to Octubre de 2026.")

    # Save to all target files!
    targets = [
        "docs/Plantilla presentación de residencia oficial.pptx",
        "docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx",
        "docs/PRESENTACION_AVANCE_RESIDENCIA_GPON.pptx"
    ]
    for t in targets:
        prs.save(t)
        print(f"Saved: {t}")

if __name__ == "__main__":
    update_presentation()
