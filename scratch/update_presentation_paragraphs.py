import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.xmlchemy import OxmlElement

COLOR_PRIMARY = RGBColor(31, 41, 55)   # #1F2937 dark slate
COLOR_MUTED = RGBColor(75, 85, 99)     # #4B5563 gray
FONT_FAMILY = "Arial"

def set_clean_paragraph(p, text, font_size_pt=15, bold=False, italic=False, color=COLOR_PRIMARY, space_after_pt=10, line_spacing=1.3, align=PP_ALIGN.LEFT):
    p.text = text
    p.font.name = FONT_FAMILY
    p.font.size = Pt(font_size_pt)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.line_spacing = line_spacing
    p.space_after = Pt(space_after_pt)
    p.space_before = Pt(0)
    p.alignment = align
    
    # Remove any XML bullet formatting
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith('buChar') or child.tag.endswith('buAutoNum') or child.tag.endswith('buSzPct') or child.tag.endswith('buFont'):
            pPr.remove(child)
    if pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buNone') is None:
        buNone = OxmlElement('a:buNone')
        pPr.append(buNone)

def apply_presentation_revisions(pptx_path):
    print(f"\n==========================================")
    print(f"Applying revisions to: {pptx_path}")
    print(f"==========================================")
    
    prs = pptx.Presentation(pptx_path)
    
    # ----------------------------------------------------
    # SLIDE 2: Agenda
    # ----------------------------------------------------
    slide2 = prs.slides[1]
    for s in slide2.shapes:
        if s.has_text_frame and 'Planteamiento del Problema' in s.text_frame.text:
            tf = s.text_frame
            tf.word_wrap = True
            lines = [
                "1. Empresa (Ubicación Geográfica y Organigrama)",
                "2. Introducción al Proyecto",
                "3. Planteamiento del Problema y Pregunta de Investigación",
                "4. Justificación Técnica y Operativa",
                "5. Objetivos del Proyecto (General y Específicos)",
                "6. Herramientas Tecnológicas (Stack de Software)",
                "7. Evidencias de Avance Técnico",
                "8. Cronograma de Actividades",
                "9. Conclusiones Preliminares (Avance)",
                "10. Bibliografía y Referencias"
            ]
            tf.clear()
            for idx, line in enumerate(lines):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                set_clean_paragraph(p, line, font_size_pt=14, bold=False, space_after_pt=4, line_spacing=1.2)
            print("Updated Slide 2 (Agenda)")
            break

    # ----------------------------------------------------
    # SLIDE 3: Empresa
    # ----------------------------------------------------
    slide3 = prs.slides[2]
    for s in slide3.shapes:
        if s.has_text_frame and ('GPON TELECOM S.A. de C.V.' in s.text_frame.text or 'Misión Operativa' in s.text_frame.text):
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            text_s3 = (
                "GPON TELECOM S.A. de C.V. es una empresa proveedora de servicios de telecomunicaciones e Internet "
                "de alta velocidad mediante fibra óptica (FTTx), ubicada en San José del Rincón, Estado de México, "
                "cuya misión operativa se enfoca en brindar conectividad digital confiable a comunidades urbanas y "
                "rurales mediante infraestructura pasiva GPON de última generación."
            )
            set_clean_paragraph(p, text_s3, font_size_pt=15, bold=False, space_after_pt=8, line_spacing=1.3)
            print("Updated Slide 3 (Empresa)")
            break

    # ----------------------------------------------------
    # SLIDE 5: Introducción (Single cohesive general paragraph)
    # ----------------------------------------------------
    slide5 = prs.slides[4]
    for s in slide5.shapes:
        if s.has_text_frame and ('Propósito General' in s.text_frame.text or 'Problemática Atendida' in s.text_frame.text or 'plataforma web' in s.text_frame.text.lower()):
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            text_s5 = (
                "El presente proyecto de residencia comprende el desarrollo de una plataforma web integral para el "
                "inventario centralizado y mapeo geoespacial de la red de fibra óptica GPON/FTTx de la empresa GPON "
                "TELECOM S.A. de C.V. Su propósito fundamental es modernizar y reemplazar la gestión manual desarticulada "
                "en campo, previniendo colisiones en la asignación de puertos y saturación imprevista en cajas NAP. "
                "La solución articula una arquitectura desacoplada con visor cartográfico interactivo en React y Leaflet, "
                "un backend transaccional robusto en Node.js con Express y PostgreSQL con control de concurrencia ACID, "
                "respaldada por un modelo de resiliencia Offline-First (IndexedDB con Dexie.js) que garantiza la "
                "continuidad operativa de las cuadrillas en localidades rurales con cobertura móvil nula."
            )
            set_clean_paragraph(p, text_s5, font_size_pt=15, bold=False, space_after_pt=10, line_spacing=1.35)
            print("Updated Slide 5 (Introducción - Single general paragraph)")
            break

    # ----------------------------------------------------
    # SLIDE 6: Planteamiento del Problema & Pregunta de Investigación
    # ----------------------------------------------------
    slide6 = prs.slides[5]
    for s in slide6.shapes:
        if s.has_text_frame and ('Carencia de Inventario' in s.text_frame.text or 'Colisión por Asignación' in s.text_frame.text or 'Pregunta Generadora' in s.text_frame.text or 'deficiencias operativas' in s.text_frame.text):
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            # Paragraph 1: Problem statement
            p1 = tf.paragraphs[0]
            text_prob = (
                "Actualmente, la empresa GPON TELECOM S.A. de C.V. enfrenta severas dificultades operativas derivadas "
                "de la gestión manual y desarticulada en el inventario de su red de fibra óptica pasiva. La carencia "
                "de un registro digital centralizado provoca frecuentes colisiones por asignación concurrente de puertos "
                "entre cuadrillas de campo y personal de NOC, saturación inadvertida en cajas terminales NAP y demoras "
                "críticas en la provisión de nuevos suscriptores, escenario agravado por el aislamiento que experimentan "
                "los técnicos al trabajar en localidades rurales sin conectividad móvil."
            )
            set_clean_paragraph(p1, text_prob, font_size_pt=14.5, bold=False, space_after_pt=12, line_spacing=1.3)
            
            # Paragraph 2: Question Header
            p2 = tf.add_paragraph()
            set_clean_paragraph(p2, "Pregunta de Investigación:", font_size_pt=14.5, bold=True, color=RGBColor(17, 24, 39), space_after_pt=4, line_spacing=1.2)
            
            # Paragraph 3: Research Question (proper, professional, academic, NO "Respuesta: Sí")
            p3 = tf.add_paragraph()
            text_preg = (
                "¿De qué manera el desarrollo e implementación de una plataforma web georreferenciada con arquitectura "
                "transaccional desacoplada y soporte Offline-First optimiza el control de inventario y previene la "
                "saturación y colisión de puertos en la red de fibra óptica GPON de la empresa GPON TELECOM S.A. de C.V.?"
            )
            set_clean_paragraph(p3, text_preg, font_size_pt=14.5, bold=False, italic=True, color=RGBColor(31, 41, 55), space_after_pt=8, line_spacing=1.3)
            print("Updated Slide 6 (Planteamiento del Problema & Pregunta de Investigación)")
            break

    # ----------------------------------------------------
    # SLIDE 7: Justificación
    # ----------------------------------------------------
    slide7 = prs.slides[6]
    for s in slide7.shapes:
        if s.has_text_frame and ('Eficiencia Operativa' in s.text_frame.text or 'Trazabilidad Geoespacial' in s.text_frame.text or 'transformar integralmente' in s.text_frame.text):
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            text_s7 = (
                "La implementación de este sistema se justifica al transformar integralmente la operatividad y rentabilidad "
                "de GPON TELECOM S.A. de C.V., automatizando la asignación y liberación de puertos SC-APC para reducir los "
                "tiempos de atención técnica de 45 a tan solo 15 minutos por suscriptor. Asimismo, faculta al centro de "
                "operaciones y cuadrillas para disponer de trazabilidad geoespacial en tiempo real con semaforización de "
                "saturación, mientras que su arquitectura resiliente Offline-First garantiza la continuidad laboral en campo "
                "sin cobertura celular; disminuyendo en un 80% los traslados técnicos fallidos por cajas saturadas y protegiendo "
                "de forma sustentable la inversión en infraestructura pasiva de fibra óptica."
            )
            set_clean_paragraph(p, text_s7, font_size_pt=15, bold=False, space_after_pt=10, line_spacing=1.35)
            print("Updated Slide 7 (Justificación - Single paragraph)")
            break

    # ----------------------------------------------------
    # SLIDE 8: Objetivos - Objetivo General y Alcance
    # ----------------------------------------------------
    slide8 = prs.slides[7]
    for s in slide8.shapes:
        if s.has_text_frame and ('Objetivo General' in s.text_frame.text or 'Alcance de Dominio' in s.text_frame.text or 'infraestructura de red GPON' in s.text_frame.text):
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            # Header
            p_head = tf.paragraphs[0]
            set_clean_paragraph(p_head, "Objetivo General y Alcance:", font_size_pt=15, bold=True, space_after_pt=6, line_spacing=1.2)
            
            # Single unified paragraph
            p_body = tf.add_paragraph()
            text_s8 = (
                "Desarrollar e implementar una plataforma web para el inventario centralizado y mapeo lógico de la "
                "infraestructura de red de fibra óptica GPON/FTTx en GPON TELECOM S.A. de C.V., mediante una arquitectura "
                "desacoplada cliente-servidor, bases de datos relacionales y visualización geoespacial interactiva, con la "
                "finalidad de optimizar la asignación transaccional de puertos, eliminar colisiones concurrentes y garantizar "
                "la trazabilidad técnica en campo. El alcance abarca el control integral de la planta externa pasiva desde el "
                "ODF de cabecera hasta cajas NAP de 16 puertos, incorporando bloqueo pesimista ACID y sincronización diferida "
                "para asegurar la disponibilidad ininterrumpida de las cuadrillas técnicas móviles."
            )
            set_clean_paragraph(p_body, text_s8, font_size_pt=14.5, bold=False, space_after_pt=10, line_spacing=1.3)
            print("Updated Slide 8 (Objetivo General y Alcance - Single paragraph)")
            break

    # ----------------------------------------------------
    # SLIDE 9: Objetivos Específicos
    # ----------------------------------------------------
    slide9 = prs.slides[8]
    for s in slide9.shapes:
        if s.has_text_frame and ('Objetivos Específicos' in s.text_frame.text or 'Levantamiento e Infraestructura' in s.text_frame.text or 'Para dar cumplimiento al objetivo' in s.text_frame.text):
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            # Header
            p_head = tf.paragraphs[0]
            set_clean_paragraph(p_head, "Objetivos Específicos:", font_size_pt=15, bold=True, space_after_pt=6, line_spacing=1.2)
            
            # Single narrative paragraph of specific goals
            p_body = tf.add_paragraph()
            text_s9 = (
                "Para dar cumplimiento al objetivo general, el proyecto articula las siguientes metas específicas: "
                "caracterizar la infraestructura física de la red GPON pasiva (ODF, splitters PLC y cajas NAP); diseñar el "
                "modelo de datos relacional y los diagramas de arquitectura de software bajo estándar UML; maquetar e implementar "
                "interfaces web adaptativas con cartografía interactiva en React Leaflet; programar una API REST robusta en Node.js "
                "y PostgreSQL con control de concurrencia pesimista (SELECT ... FOR UPDATE); desarrollar un motor de persistencia "
                "local Offline-First en Dexie.js para la sincronización automática de operaciones en campo; y validar el sistema "
                "mediante auditorías de consistencia transaccional, normativas de accesibilidad WCAG 2.1 y generación de reportes "
                "ejecutivos en PDF."
            )
            set_clean_paragraph(p_body, text_s9, font_size_pt=14, bold=False, space_after_pt=10, line_spacing=1.3)
            print("Updated Slide 9 (Objetivos Específicos - Single narrative paragraph)")
            break

    # ----------------------------------------------------
    # SLIDE 11: Evidencias y Entregables
    # ----------------------------------------------------
    slide11 = prs.slides[10]
    for s in slide11.shapes:
        if s.has_text_frame and ('Principales entregables' in s.text_frame.text or 'Mapeo Cartográfico GIS' in s.text_frame.text or 'componentes técnicos desarrollados' in s.text_frame.text):
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            # Header
            p_head = tf.paragraphs[0]
            set_clean_paragraph(p_head, "Componentes Técnicos y Entregables Desarrollados:", font_size_pt=15, bold=True, space_after_pt=6, line_spacing=1.2)
            
            # Single unified paragraph (updated with 82 figures)
            p_body = tf.add_paragraph()
            text_s11 = (
                "Los principales componentes técnicos desarrollados a la fecha comprenden la documentación técnica formal de "
                "residencia integrada por 82 figuras y 26 tablas estandarizadas de ingeniería, el prototipo funcional del visor "
                "cartográfico en React Leaflet con trazado vectorial de fibra y semaforización de saturación en cajas NAP, la matriz "
                "interactiva de chasis de 16 puertos SC-APC gobernada por transacciones pesimistas en PostgreSQL para prevenir "
                "colisiones concurrentes, así como la base de código modular en TypeScript con persistencia local Offline-First en "
                "Dexie.js para garantizar la operación ininterrumpida de las cuadrillas técnicas en campo."
            )
            set_clean_paragraph(p_body, text_s11, font_size_pt=14, bold=False, space_after_pt=10, line_spacing=1.3)
            print("Updated Slide 11 (Evidencias y Entregables - Single paragraph, 82 figures)")
            break

    # ----------------------------------------------------
    # SLIDE 13: Conclusiones Preliminares
    # ----------------------------------------------------
    slide13 = prs.slides[12]
    for s in slide13.shapes:
        if s.has_text_frame and s.name == 'Google Shape;215;p13':
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            text_s13 = (
                "Con base en el cronograma de actividades planificado, el proyecto registra actualmente un avance físico y "
                "metodológico estimado del 40%, cumpliendo satisfactoriamente en tiempo y forma con las fases de levantamiento "
                "de infraestructura, análisis formal de requerimientos, diseño arquitectónico del software y prototipado inicial "
                "del visor cartográfico y del backend transaccional, consolidando las bases técnicas para el cierre de la "
                "implementación y pruebas operativas en campo durante las semanas subsecuentes."
            )
            set_clean_paragraph(p, text_s13, font_size_pt=15, bold=False, space_after_pt=10, line_spacing=1.35)
            print("Updated Slide 13 (Conclusiones Preliminares - Single cohesive paragraph)")
            break

    # ----------------------------------------------------
    # SLIDE 14: Bibliografía y Referencias (APA format, no bullet character)
    # ----------------------------------------------------
    slide14 = prs.slides[13]
    for s in slide14.shapes:
        if s.has_text_frame and ('International Telecommunication Union' in s.text_frame.text or 'Tanenbaum' in s.text_frame.text):
            tf = s.text_frame
            tf.word_wrap = True
            refs = [
                "International Telecommunication Union. (2003). Recomendación ITU-T G.984.1 / G.984.2: Redes ópticas pasivas con capacidad de gigabit (GPON). Ginebra: ITU.",
                "Tanenbaum, A. S., & Wetherall, D. J. (2014). Redes de computadoras (5.ª ed.). Ciudad de México: Pearson Educación.",
                "Silberschatz, A., Korth, H. F., & Sudarshan, S. (2019). Database System Concepts (7th ed.). Nueva York: McGraw-Hill.",
                "World Wide Web Consortium (W3C). (2018). Web Content Accessibility Guidelines (WCAG) 2.1. W3C Recommendation. https://www.w3.org/TR/WCAG21/",
                "PostgreSQL Global Development Group. (2024). PostgreSQL 16.0 Documentation: Explicit Locking and Transaction Isolation. https://www.postgresql.org/docs/16/"
            ]
            tf.clear()
            for idx, ref in enumerate(refs):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                set_clean_paragraph(p, ref, font_size_pt=11.5, bold=False, space_after_pt=6, line_spacing=1.15)
            print("Updated Slide 14 (Bibliografía - Clean APA paragraphs)")
            break

    # Save presentation
    prs.save(pptx_path)
    print(f"Successfully saved: {pptx_path}")

# Run for both target presentations
files_to_update = [
    'docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx',
    'docs/Plantilla presentación de residencia.pptx'
]

for fp in files_to_update:
    if os.path.exists(fp):
        apply_presentation_revisions(fp)

# Also save a canonical copy named docs/PRESENTACION_AVANCE_RESIDENCIA_GPON.pptx
canonical_path = 'docs/PRESENTACION_AVANCE_RESIDENCIA_GPON.pptx'
import shutil
shutil.copyfile('docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx', canonical_path)
print(f"\nCreated/Updated canonical copy: {canonical_path}")
