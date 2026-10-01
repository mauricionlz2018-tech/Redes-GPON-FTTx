import docx
import re

def verify_full_document(doc_path):
    print(f"\n=======================================================")
    print(f"VERIFYING: {doc_path}")
    print(f"=======================================================")
    doc = docx.Document(doc_path)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total sections: {len(doc.sections)}")
    print(f"Total tables: {len(doc.tables)}")

    # 1. Section margins
    for i, s in enumerate(doc.sections):
        print(f"Section {i}: {s.page_width.cm:.2f} x {s.page_height.cm:.2f} cm | Margins: L={s.left_margin.cm:.2f}, R={s.right_margin.cm:.2f}, T={s.top_margin.cm:.2f}, B={s.bottom_margin.cm:.2f}")

    # 2. Locate Introducción
    intro_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Introducci\u00f3n":
            intro_idx = i
            break
    print(f"Introducci\u00f3n located at P[{intro_idx}]")
    assert intro_idx is not None, "Introducci\u00f3n not found!"

    # 3. Figure Index check
    idx_figs = []
    in_fig_idx = False
    for i in range(intro_idx):
        t = doc.paragraphs[i].text.strip()
        if 'ndice de figuras' in t:
            in_fig_idx = True
            continue
        if 'ndice de tablas' in t:
            in_fig_idx = False
            break
        if in_fig_idx and t.startswith("Figura "):
            idx_figs.append((i, t))

    print(f"Figure index entries found: {len(idx_figs)}")
    assert len(idx_figs) == 78, f"Expected 78 figure index entries, got {len(idx_figs)}"
    for k, (p_idx, t) in enumerate(idx_figs):
        expected_num = k + 1
        m = re.match(r'^Figura\s+(\d+)\.', t)
        assert m is not None, f"TOC Figure {p_idx} format error: {t}"
        num = int(m.group(1))
        assert num == expected_num, f"TOC Figure {k} num {num} != expected {expected_num}: {t}"
    print("ALL 78 FIGURE INDEX ENTRIES STRICTLY SEQUENTIAL (1 to 78)!")

    # 4. Table Index check
    idx_tbls = []
    in_tbl_idx = False
    for i in range(intro_idx):
        t = doc.paragraphs[i].text.strip()
        if 'ndice de tablas' in t:
            in_tbl_idx = True
            continue
        if in_tbl_idx and t.startswith("Tabla "):
            idx_tbls.append((i, t))

    print(f"Table index entries found: {len(idx_tbls)}")
    assert len(idx_tbls) == 26, f"Expected 26 table index entries, got {len(idx_tbls)}"
    for k, (p_idx, t) in enumerate(idx_tbls):
        expected_num = k + 1
        m = re.match(r'^Tabla\s+(\d+)\.', t)
        assert m is not None, f"TOC Table {p_idx} format error: {t}"
        num = int(m.group(1))
        assert num == expected_num, f"TOC Table {k} num {num} != expected {expected_num}: {t}"
    print("ALL 26 TABLE INDEX ENTRIES STRICTLY SEQUENTIAL (1 to 26)!")

    # 5. Body Figures check
    body_figs = []
    for i, p in enumerate(doc.paragraphs[intro_idx:], start=intro_idx):
        t = p.text.strip()
        if t.startswith("Figura ") and '\t' not in t:
            # check previous 2 paragraphs for drawing
            p_prev1 = doc.paragraphs[i-1]
            p_prev2 = doc.paragraphs[i-2] if i >= 2 else None
            dw = len(p_prev1._p.xpath('.//w:drawing')) + len(p_prev1._p.xpath('.//w:pict'))
            if p_prev2:
                dw += len(p_prev2._p.xpath('.//w:drawing')) + len(p_prev2._p.xpath('.//w:pict'))
            p_next = doc.paragraphs[i+1]
            body_figs.append({
                'idx': i,
                'text': t,
                'has_drawing': dw > 0,
                'next_text': p_next.text.strip(),
                'style': p.style.name
            })

    print(f"Body figures found: {len(body_figs)}")
    assert len(body_figs) == 78, f"Expected 78 body figures, got {len(body_figs)}"
    for k, fig in enumerate(body_figs):
        expected_num = k + 1
        m = re.match(r'^Figura\s+(\d+)\.', fig['text'])
        assert m is not None, f"Body Figure {fig['idx']} format error: {fig['text']}"
        num = int(m.group(1))
        assert num == expected_num, f"Body Figure {k} num {num} != expected {expected_num}: {fig['text']}"
        assert fig['has_drawing'], f"Figure {num} at P[{fig['idx']}] missing drawing!"
        assert fig['style'] == 'Figuras', f"Figure {num} style {fig['style']} != Figuras"
    print("ALL 78 BODY FIGURES STRICTLY SEQUENTIAL (1 to 78) WITH DRAWINGS & STYLES!")

    # 6. Body Tables check
    body_tbls = []
    for i, p in enumerate(doc.paragraphs[intro_idx:], start=intro_idx):
        t = p.text.strip()
        if t.startswith("Tabla ") and '\t' not in t:
            body_tbls.append((i, t, p.style.name))

    print(f"Body tables found: {len(body_tbls)}")
    assert len(body_tbls) == 26, f"Expected 26 body tables, got {len(body_tbls)}"
    for k, (p_idx, t, st) in enumerate(body_tbls):
        expected_num = k + 1
        m = re.match(r'^Tabla\s+(\d+)\.', t)
        assert m is not None, f"Body Table {p_idx} format error: {t}"
        num = int(m.group(1))
        assert num == expected_num, f"Body Table {k} num {num} != expected {expected_num}: {t}"
    print("ALL 26 BODY TABLES STRICTLY SEQUENTIAL (1 to 26)!")

    # 7. Print all Chapter III items requested by teacher
    print("\n--- TEACHER'S CHECKLIST VERIFICATION IN CHAPTER III ---")
    ch3_figs = [f for f in body_figs if 27 <= int(re.match(r'^Figura\s+(\d+)\.', f['text']).group(1)) <= 53]
    for f in ch3_figs:
        print(f"  P[{f['idx']}]: {f['text']}")

    print("\n--- CARDINALITY TABLE & DATA DICTIONARIES VERIFICATION ---")
    ch3_tbls = [t for t in body_tbls if 12 <= int(re.match(r'^Tabla\s+(\d+)\.', t[1]).group(1)) <= 25]
    for t in ch3_tbls:
        print(f"  P[{t[0]}]: {t[1]}")

    print("\n>>> ALL VALIDATIONS PASSED 100%! <<<")

verify_full_document('docs/Documentacion_Residencias_avance.docx')
verify_full_document('docs/Documentacion_Residencias (4).docx')

