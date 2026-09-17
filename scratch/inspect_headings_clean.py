import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
print("Total paragraphs:", len(doc.paragraphs))

# Locate key headings
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if p.style.name.startswith("Heading") or any(k in t.upper() for k in ["CAPÍTULO", "CAPITULO", "CONCLUSIONES", "REFERENCIAS", "ANEXOS"]):
        print(f"P[{i}] ({p.style.name}): '{t[:70]}'")

