import pptx

prs = pptx.Presentation(r'docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx')
slide_w = prs.slide_width
slide_h = prs.slide_height
print(f'Slide dimensions: {slide_w} x {slide_h}')

for idx in [2, 3, 9, 10, 11]:
    s = prs.slides[idx]
    print(f'=== Slide {idx+1} ===')
    for shp in s.shapes:
        txt = shp.text_frame.text.strip().replace('\n', ' ') if shp.has_text_frame else ''
        print(f'  Shape {shp.shape_id:<5} ({shp.name:<22}): L={shp.left:<8} T={shp.top:<8} W={shp.width:<8} H={shp.height:<8} | Bottom={shp.top+shp.height} | Right={shp.left+shp.width} | {txt[:40]}')

