import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = 'scratch/test_sections_setup.docx'
doc = docx.Document(doc_path)

def set_page_field_in_footer(footer, num_format=None):
    # Clear footer paragraphs
    for p in footer.paragraphs:
        p.text = ''
    # Create single paragraph with PAGE field aligned to right
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._element.get_or_add_pPr()
    
    # Run with PAGE field
    fld_xml = parse_xml(f'''
    <w:p {nsdecls("w")}>
        <w:pPr>
            <w:jc w:val="right"/>
            <w:rPr>
                <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                <w:sz w:val="22"/>
            </w:rPr>
        </w:pPr>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                <w:sz w:val="22"/>
            </w:rPr>
            <w:fldChar w:fldCharType="begin"/>
        </w:r>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                <w:sz w:val="22"/>
            </w:rPr>
            <w:instrText xml:space="preserve"> PAGE </w:instrText>
        </w:r>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                <w:sz w:val="22"/>
            </w:rPr>
            <w:fldChar w:fldCharType="separate"/>
        </w:r>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                <w:sz w:val="22"/>
            </w:rPr>
            <w:t>1</w:t>
        </w:r>
        <w:r>
            <w:rPr>
                <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                <w:sz w:val="22"/>
            </w:rPr>
            <w:fldChar w:fldCharType="end"/>
        </w:r>
    </w:p>
    ''')
    p._element.getparent().replace(p._element, fld_xml)

def clear_footer(footer):
    footer.is_linked_to_previous = False
    for p in list(footer.paragraphs):
        p.text = ''

print(f"Total secciones a configurar: {len(doc.sections)}")

for i, s in enumerate(doc.sections):
    sectPr = s._sectPr
    # Limpiar pgNumType previo
    for el in sectPr.xpath('./w:pgNumType'):
        sectPr.remove(el)
        
    if i == 0:
        # Seccion 0: Caratula
        s.different_first_page_header_footer = True
        clear_footer(s.first_page_footer)
        clear_footer(s.footer)
        sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="decimal"/>'))
        print("  -> Sec 0 (Carátula): configurada.")
        
    elif i == 1:
        # Seccion 1: Preliminares (Indices, Introduccion, Planteamiento, Justificacion, Objetivos)
        s.different_first_page_header_footer = False
        s.footer.is_linked_to_previous = False
        set_page_field_in_footer(s.footer)
        sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="3" w:fmt="upperRoman"/>'))
        print("  -> Sec 1 (Preliminares: Índices hasta Objetivos en ROMANOS): configurada.")
        
    elif 2 <= i <= 6:
        # Secciones 2 a 6: Capitulos I, II, III, IV, V
        ch_num = i - 1
        s.different_first_page_header_footer = True
        clear_footer(s.first_page_footer)
        s.footer.is_linked_to_previous = False
        set_page_field_in_footer(s.footer)
        # Decimal consecutivo a la seccion previa (sin start)
        sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal"/>'))
        print(f"  -> Sec {i} (Capítulo {ch_num}): portada sin número, páginas siguientes en arábigos consecutivos.")
        
    elif i == 7:
        # Seccion 7: Conclusiones y Glosario
        s.different_first_page_header_footer = False
        s.footer.is_linked_to_previous = False
        set_page_field_in_footer(s.footer)
        # Romanos consecutivos a la paginacion general (sin start)
        sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="upperRoman"/>'))
        print("  -> Sec 7 (Conclusiones y Glosario): upperRoman consecutivos.")
        
    elif i == 8:
        # Seccion 8: REFERENCIAS
        s.different_first_page_header_footer = False
        clear_footer(s.footer)
        print("  -> Sec 8 (REFERENCIAS): sin paginación.")
        
    elif i == 9:
        # Seccion 9: Anexos
        s.different_first_page_header_footer = False
        clear_footer(s.footer)
        print("  -> Sec 9 (Anexos): sin paginación.")

doc.save(doc_path)
print("Archivo test guardado exitosamente.")

