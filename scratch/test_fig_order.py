import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
target_p = doc.paragraphs[433]

p_src = doc.add_paragraph()
p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_src = p_src.add_run("Fuente: Recuperado de Neon Docs, 2026.")
r_src.font.name = 'Arial'
r_src.font.size = Pt(8.5)
r_src.italic = True
r_src.font.color.rgb = RGBColor(0, 0, 0)

p_cap = doc.add_paragraph()
p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_cap = p_cap.add_run("Figura 15. Arquitectura de base de datos Serverless desacoplada en Neon Database (Postgres).")
r_cap.font.name = 'Arial'
r_cap.font.size = Pt(9)
r_cap.bold = True
r_cap.font.color.rgb = RGBColor(0, 0, 0)

p_img = doc.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_img = p_img.add_run()
r_img.add_picture('scratch/teoria_neon_serverless_architecture.png', width=Inches(5.5))

target_p._p.addnext(p_src._p)
target_p._p.addnext(p_cap._p)
target_p._p.addnext(p_img._p)

doc.save('scratch/test_fig_order.docx')
print("Saved test_fig_order.docx successfully!")

doc2 = docx.Document('scratch/test_fig_order.docx')
for i in range(432, 438):
    xml = doc2.paragraphs[i]._p.xml
    is_img = '<w:drawing' in xml
    print(f"P[{i}] [IMG={is_img}]: '{doc2.paragraphs[i].text[:60]}'")

