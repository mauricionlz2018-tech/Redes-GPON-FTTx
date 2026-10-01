import pptx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

template_path = os.path.join('docs', 'Plantilla presentación de residencia oficial.pptx')
prs = pptx.Presentation(template_path)

print(f"Presentation has {len(prs.slides)} slides.")
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for s in slide.shapes:
        txt = s.text.strip() if s.has_text_frame else ""
        fill_info = ""
        try:
            if s.fill.type:
                fill_info = f"fill={s.fill.type}"
                if s.fill.type == pptx.enum.dml.MSO_FILL_TYPE.SOLID:
                    fill_info += f", rgb={s.fill.fore_color.rgb}"
        except:
            pass
        print(f"  Shape '{s.name}' [{s.shape_type}], left={s.left/914400:.2f}, top={s.top/914400:.2f}, w={s.width/914400:.2f}, h={s.height/914400:.2f}, {fill_info}")
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    fn = p.font.name if p.font else None
                    fs = p.font.size.pt if p.font and p.font.size else None
                    print(f"    Text: '{p_text}' | font={fn}, size={fs}")
