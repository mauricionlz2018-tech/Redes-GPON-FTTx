import docx

for path in ['docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx', 'docs/Documentacion_Residencias_avance_BACKUP_06OCT.docx']:
    d = docx.Document(path)
    print(f"=== {path} ===")
    print(f"  Párrafos: {len(d.paragraphs)}, Tablas: {len(d.tables)}")
    if len(d.tables) > 0:
        for t_idx, t in enumerate(d.tables):
            if any('4.' in p.text for p in d.paragraphs if p._element in [t._element.getprevious()]):
                pass
            print(f"    Tabla {t_idx}: {len(t.rows)}x{len(t.columns)} | primer celda: {t.rows[0].cells[0].text[:30].strip()}")

