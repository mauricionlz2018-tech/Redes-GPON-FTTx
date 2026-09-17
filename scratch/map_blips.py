import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')

for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    if 'r:embed' in xml:
        # find blip
        import xml.etree.ElementTree as ET
        root = ET.fromstring(xml)
        for elem in root.iter():
            if elem.tag.endswith('blip'):
                r_id = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                target = doc.part.rels[r_id].target_ref if r_id in doc.part.rels else 'UNKNOWN'
                # Find caption
                cap = ""
                for off in range(1, 4):
                    if i + off < len(doc.paragraphs):
                        t = doc.paragraphs[i+off].text.strip()
                        if 'figura' in t.lower() or 'tabla' in t.lower():
                            cap = t
                            break
                print(f"P[{i}] -> r_id={r_id} ({target}) | cap='{cap[:60]}'")

