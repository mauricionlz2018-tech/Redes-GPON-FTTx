import os
import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

COLOR_PRIMARY = RGBColor(31, 41, 55)   # #1F2937 dark slate
COLOR_HEADING = RGBColor(17, 24, 39)   # #111827 deep black/slate
FONT_FAMILY = "Arial"

def set_clean_paragraph(p, text, font_size_pt=15, bold=False, italic=False, color=COLOR_PRIMARY, space_after_pt=10, line_spacing=1.3, align=PP_ALIGN.LEFT):
    p.text = text
    p.font.name = FONT_FAMILY
    p.font.size = Pt(font_size_pt)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.line_spacing = line_spacing
    p.space_after = Pt(space_after_pt)
    p.space_before = Pt(0)
    p.alignment = align
    
    # Remove XML bullets
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith('buChar') or child.tag.endswith('buAutoNum') or child.tag.endswith('buSzPct') or child.tag.endswith('buFont'):
            pPr.remove(child)
    if pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buNone') is None:
        buNone = OxmlElement('a:buNone')
        pPr.append(buNone)

def apply_anteproyecto_and_general_updates(pptx_path):
    print(f"\n==========================================")
    print(f"Applying general & anteproyecto revisions to: {pptx_path}")
    print(f"==========================================")
    
    prs = pptx.Presentation(pptx_path)
    
    # ----------------------------------------------------
    # SLIDE 5: Introducción (General del Proyecto, sin Offline)
    # ----------------------------------------------------
    slide5 = prs.slides[4]
    for s in slide5.shapes:
        if s.has_text_frame and s.height > 1000000:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            text_s5 = (
                "El presente proyecto de residencia comprende el desarrollo de una plataforma web integral para el "
                "inventario centralizado y mapeo geoespacial de la red de fibra óptica GPON/FTTx de la empresa GPON "
                "TELECOM S.A. de C.V. Su propósito fundamental es modernizar y reemplazar la gestión manual desarticulada "
                "en campo, previniendo colisiones en la asignación de puertos y saturación imprevista en cajas NAP. "
                "La solución articula una arquitectura desacoplada con visor cartográfico interactivo en React y Leaflet, "
                "un backend transaccional robusto en Node.js con Express y una base de datos relacional en PostgreSQL "
                "con control de concurrencia ACID, permitiendo al personal técnico y de soporte consultar, administrar y "
                "actualizar en tiempo real el estado físico y lógico de toda la infraestructura pasiva."
            )
            set_clean_paragraph(p, text_s5, font_size_pt=15, bold=False, space_after_pt=10, line_spacing=1.35)
            print("Updated Slide 5 (Introducción - General, sin offline)")
            break

    # ----------------------------------------------------
    # SLIDE 6: Planteamiento del Problema & Pregunta (General, sin Offline)
    # ----------------------------------------------------
    slide6 = prs.slides[5]
    for s in slide6.shapes:
        if s.has_text_frame and s.height > 1000000:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            # Paragraph 1: Problem statement
            p1 = tf.paragraphs[0]
            text_prob = (
                "Actualmente, la empresa GPON TELECOM S.A. de C.V. enfrenta severas dificultades operativas derivadas "
                "de la gestión manual y desarticulada en el inventario de su red de fibra óptica pasiva. La carencia "
                "de un registro digital centralizado provoca frecuentes colisiones por asignación concurrente de puertos "
                "entre cuadrillas de campo y personal de NOC, saturación inadvertida en cajas terminales NAP y demoras "
                "críticas en la provisión de nuevos suscriptores, dificultando la supervisión técnica y la toma de "
                "decisiones sobre la infraestructura de telecomunicaciones."
            )
            set_clean_paragraph(p1, text_prob, font_size_pt=14.5, bold=False, space_after_pt=12, line_spacing=1.3)
            
            # Paragraph 2: Question Header
            p2 = tf.add_paragraph()
            set_clean_paragraph(p2, "Pregunta de Investigación:", font_size_pt=14.5, bold=True, color=COLOR_HEADING, space_after_pt=4, line_spacing=1.2)
            
            # Paragraph 3: General Question (Alternativa 1 seleccionada por el usuario)
            p3 = tf.add_paragraph()
            text_preg = (
                "¿El desarrollo e implementación de un sistema web integral de inventario y mapeo lógico permite "
                "modernizar la gestión operativa y optimizar la administración de la red de fibra óptica GPON/FTTx "
                "en la empresa GPON TELECOM S.A. de C.V.?"
            )
            set_clean_paragraph(p3, text_preg, font_size_pt=14.5, bold=False, italic=True, space_after_pt=8, line_spacing=1.3)
            
            # Paragraph 4: Respuesta directa que valida la hipótesis
            p4 = tf.add_paragraph()
            set_clean_paragraph(p4, "Respuesta: Sí.", font_size_pt=14.5, bold=True, color=RGBColor(37, 99, 235), space_after_pt=4, line_spacing=1.2)
            print("Updated Slide 6 (Planteamiento del Problema & Pregunta General con respuesta Sí)")
            break

    # ----------------------------------------------------
    # SLIDE 7: Justificación (General, sin Offline)
    # ----------------------------------------------------
    slide7 = prs.slides[6]
    for s in slide7.shapes:
        if s.has_text_frame and s.height > 1000000:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            text_s7 = (
                "La implementación de este sistema se justifica al transformar integralmente la operatividad y rentabilidad "
                "de GPON TELECOM S.A. de C.V., automatizando la asignación y liberación de puertos SC-APC para reducir los "
                "tiempos de atención técnica de 45 a tan solo 15 minutos por suscriptor. Asimismo, faculta al centro de "
                "operaciones y cuadrillas para disponer de trazabilidad geoespacial en tiempo real con semaforización de "
                "saturación, disminuyendo en un 80% los traslados técnicos fallidos por cajas saturadas y protegiendo de "
                "forma sustentable la inversión en la infraestructura pasiva de fibra óptica."
            )
            set_clean_paragraph(p, text_s7, font_size_pt=15, bold=False, space_after_pt=10, line_spacing=1.35)
            print("Updated Slide 7 (Justificación - General, sin offline)")
            break

    # ----------------------------------------------------
    # SLIDE 8: Objetivo General (Literal del Anteproyecto)
    # ----------------------------------------------------
    slide8 = prs.slides[7]
    for s in slide8.shapes:
        if s.has_text_frame and s.height > 1000000:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            p_head = tf.paragraphs[0]
            set_clean_paragraph(p_head, "Objetivo General:", font_size_pt=16, bold=True, color=COLOR_HEADING, space_after_pt=10, line_spacing=1.2)
            
            p_body = tf.add_paragraph()
            text_s8 = (
                "Desarrollar una aplicación web para el inventario y mapeo lógico de redes GPON/FTTx, "
                "mediante el diseño de vistas tabulares y diagramas lógicos simples, para optimizar "
                "el control de la infraestructura por parte del equipo técnico y de soporte."
            )
            set_clean_paragraph(p_body, text_s8, font_size_pt=15, bold=False, space_after_pt=12, line_spacing=1.35)
            print("Updated Slide 8 (Objetivo General - Extraído del Anteproyecto)")
            break

    # ----------------------------------------------------
    # SLIDE 9: Objetivos Específicos (Literales del Anteproyecto)
    # ----------------------------------------------------
    slide9 = prs.slides[8]
    for s in slide9.shapes:
        if s.has_text_frame and s.height > 1000000:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            p_head = tf.paragraphs[0]
            set_clean_paragraph(p_head, "Objetivos Específicos:", font_size_pt=16, bold=True, color=COLOR_HEADING, space_after_pt=10, line_spacing=1.2)
            
            specs = [
                "1. Analizar los requerimientos operativos de la red de fibra óptica mediante el levantamiento de información sobre el hardware (NAP/ODF), para diseñar la arquitectura del sistema y el modelo de la base de datos.",
                "2. Construir una interfaz web responsiva a través del uso de tecnologías de desarrollo frontend, para que los técnicos de campo actualicen desde su celular el estado de los puertos (libres, ocupados, dañados).",
                "3. Programar el backend del sistema mediante la creación de una API centralizada, para gestionar la información de los activos y vincular los puertos de red correspondientes a los datos de los clientes."
            ]
            for idx, spec in enumerate(specs):
                p_item = tf.add_paragraph()
                set_clean_paragraph(p_item, spec, font_size_pt=14, bold=False, space_after_pt=10, line_spacing=1.3)
            print("Updated Slide 9 (Objetivos Específicos - Extraídos del Anteproyecto)")
            break

    # ----------------------------------------------------
    # SLIDE 10: Herramientas (General)
    # ----------------------------------------------------
    slide10 = prs.slides[9]
    for s in slide10.shapes:
        if s.has_text_frame and s.name == 'Google Shape;167;p10':
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            set_clean_paragraph(p, "Para el desarrollo de la plataforma web se seleccionó un stack tecnológico desacoplado, moderno y de alto rendimiento:", font_size_pt=14.5, bold=False, space_after_pt=6, line_spacing=1.2)
            print("Updated Slide 10 (Herramientas - Texto general)")
            break

    # ----------------------------------------------------
    # SLIDE 11: Evidencias y Entregables (General, sin Offline)
    # ----------------------------------------------------
    slide11 = prs.slides[10]
    for s in slide11.shapes:
        if s.has_text_frame and s.height > 1000000:
            tf = s.text_frame
            tf.word_wrap = True
            tf.clear()
            
            p_head = tf.paragraphs[0]
            set_clean_paragraph(p_head, "Componentes Técnicos y Entregables Desarrollados:", font_size_pt=15, bold=True, color=COLOR_HEADING, space_after_pt=6, line_spacing=1.2)
            
            p_body = tf.add_paragraph()
            text_s11 = (
                "Los principales componentes técnicos desarrollados a la fecha comprenden la documentación técnica formal de "
                "residencia integrada por 82 figuras y 26 tablas estandarizadas de ingeniería, el prototipo funcional del visor "
                "cartográfico en React Leaflet con trazado vectorial de fibra y semaforización de saturación en cajas NAP, la matriz "
                "interactiva de chasis de 16 puertos SC-APC gobernada por transacciones pesimistas en PostgreSQL para prevenir "
                "colisiones concurrentes, así como la base de código modular estructurada en TypeScript con modelos Sequelize y "
                "endpoints REST para la administración centralizada de abonados y puertos."
            )
            set_clean_paragraph(p_body, text_s11, font_size_pt=14, bold=False, space_after_pt=10, line_spacing=1.3)
            print("Updated Slide 11 (Evidencias y Entregables - General, sin offline)")
            break

    prs.save(pptx_path)
    print(f"Successfully saved: {pptx_path}")

files_to_update = [
    'docs/PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx',
    'docs/Plantilla presentación de residencia.pptx',
    'docs/PRESENTACION_AVANCE_RESIDENCIA_GPON.pptx'
]

for fp in files_to_update:
    if os.path.exists(fp):
        apply_anteproyecto_and_general_updates(fp)

print("\n=== ALL PRESENTATIONS UPDATED SUCCESSFULLY! ===")

