import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import shutil

src_path = 'docs/Documentacion_Residencias_avance_BACKUP_08OCT_FINAL_PAG.docx'
test_path = 'scratch/test_clean_solution.docx'
shutil.copyfile(src_path, test_path)

doc = docx.Document(test_path)

def remove_page_break_before(p):
    pPr = p._element.find(qn('w:pPr'))
    if pPr is not None:
        pbb = pPr.find(qn('w:pageBreakBefore'))
        if pbb is not None:
            pPr.remove(pbb)

print(f"Párrafos totales: {len(doc.paragraphs)}")

# 1. Verificar P159 (Fin de Indices)
# En BACKUP_08OCT_FINAL_PAG, P159 tiene el corte de seccion.
# Vamos a dejar el corte de sección en P159, pero asegurar que:
# Sec 1 (Indices): upperRoman, start=3, header rId13 (vacio), footer rId14 (PAGE).
# P160 inicia Sec 2 (Introduccion, Planteamiento, Justificacion, Objetivos).
p159 = doc.paragraphs[159]
p159_sp = p159._element.xpath('./w:pPr/w:sectPr')[0]
# Limpiar y reconfigurar sectPr de P159
for el in list(p159_sp):
    p159_sp.remove(el)

p159_sp.append(parse_xml(f'<w:headerReference {nsdecls("w")} w:type="default" r:id="rId13"/>'))
p159_sp.append(parse_xml(f'<w:footerReference {nsdecls("w")} w:type="default" r:id="rId14"/>'))
p159_sp.append(parse_xml(f'<w:pgSz {nsdecls("w")} w:w="12242" w:h="15842" w:code="1"/>'))
p159_sp.append(parse_xml(f'<w:pgMar {nsdecls("w")} w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>'))
p159_sp.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="upperRoman" w:start="3"/>'))
p159_sp.append(parse_xml(f'<w:cols {nsdecls("w")} w:space="708"/>'))
p159_sp.append(parse_xml(f'<w:docGrid {nsdecls("w")} w:linePitch="360"/>'))
print("  -> Reconfigurada Sección 1 (Índices): limpia, sin logos, upperRoman start=3.")

# 2. Localizar fin de Objetivos (P186) para cerrar Sección 2 (Introducción hasta Objetivos)
p_end_obj = None
for i, p in enumerate(doc.paragraphs):
    if '3. Programar el backend del sistema' in p.text:
        p_end_obj = p
        break

# P_end_obj recibe el sectPr que cierra la Sección 2 (Introduccion, Planteamiento, Justificacion, Objetivos)
# Esta sección debe ser upperRoman continua de la sección previa (sin start), header vacio (rId13), footer PAGE (rId14)
pPr_obj = p_end_obj._element.get_or_add_pPr()
for sp in pPr_obj.xpath('./w:sectPr'):
    pPr_obj.remove(sp)

new_sp_obj = parse_xml(f'''
<w:sectPr {nsdecls("w", "r")}>
    <w:headerReference w:type="default" r:id="rId13"/>
    <w:footerReference w:type="default" r:id="rId14"/>
    <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
    <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
    <w:pgNumType w:fmt="upperRoman"/>
    <w:cols w:space="708"/>
    <w:type w:val="nextPage"/>
    <w:docGrid w:linePitch="360"/>
</w:sectPr>
''')
pPr_obj.append(new_sp_obj)
print("  -> Creada Sección 2 (Introducción, Planteamiento, Justificación, Objetivos): upperRoman continua, header limpio sin logos.")

# 3. Quitar pageBreakBefore del párrafo inicial de Cap I
idx_obj = doc.paragraphs.index(p_end_obj)
remove_page_break_before(doc.paragraphs[idx_obj + 1])

# 4. Configurar cortes para cada capítulo con portada sin número:
# Función constructora de sectPr para capítulos
def make_chapter_sectPr():
    return parse_xml(f'''
    <w:sectPr {nsdecls("w", "r")}>
        <w:headerReference w:type="default" r:id="rId13"/>
        <w:headerReference w:type="first" r:id="rId13"/>
        <w:footerReference w:type="default" r:id="rId92"/>
        <w:footerReference w:type="first" r:id="rId10"/>
        <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
        <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
        <w:pgNumType w:fmt="decimal"/>
        <w:cols w:space="708"/>
        <w:type w:val="nextPage"/>
        <w:titlePg/>
        <w:docGrid w:linePitch="360"/>
    </w:sectPr>
    ''')

# Fin de Cap I: antes de Cap II
p_end_cap1 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO II. MARCO' in p.text.upper():
        p_end_cap1 = doc.paragraphs[i-8]
        remove_page_break_before(doc.paragraphs[i-7])
        break

pPr_c1 = p_end_cap1._element.get_or_add_pPr()
for sp in pPr_c1.xpath('./w:sectPr'): pPr_c1.remove(sp)
pPr_c1.append(make_chapter_sectPr())
print("  -> Capítulo I configurado: portadilla sin número, páginas siguientes en arábigos, cero logos.")

# Fin de Cap II: antes de Cap III
p_end_cap2 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO III. DISEÑO' in p.text.upper():
        p_end_cap2 = doc.paragraphs[i-7]
        remove_page_break_before(doc.paragraphs[i-6])
        break

pPr_c2 = p_end_cap2._element.get_or_add_pPr()
for sp in pPr_c2.xpath('./w:sectPr'): pPr_c2.remove(sp)
pPr_c2.append(make_chapter_sectPr())
print("  -> Capítulo II configurado: portadilla sin número, páginas siguientes en arábigos, cero logos.")

# Fin de Cap III: antes de Cap IV
p_end_cap3 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO IV. DESARROLLO' in p.text.upper():
        p_end_cap3 = doc.paragraphs[i-7]
        remove_page_break_before(doc.paragraphs[i-6])
        break

pPr_c3 = p_end_cap3._element.get_or_add_pPr()
for sp in pPr_c3.xpath('./w:sectPr'): pPr_c3.remove(sp)
pPr_c3.append(make_chapter_sectPr())
print("  -> Capítulo III configurado: portadilla sin número, páginas siguientes en arábigos, cero logos.")

# Fin de Cap IV: antes de Cap V
p_end_cap4 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO V. PRUEBAS' in p.text.upper():
        p_end_cap4 = doc.paragraphs[i-7]
        remove_page_break_before(doc.paragraphs[i-6])
        break

pPr_c4 = p_end_cap4._element.get_or_add_pPr()
for sp in pPr_c4.xpath('./w:sectPr'): pPr_c4.remove(sp)
pPr_c4.append(make_chapter_sectPr())
print("  -> Capítulo IV configurado: portadilla sin número, cero logos.")

# Fin de Cap V: antes de Conclusiones (P1209)
p_end_cap5 = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'Conclusiones':
        p_end_cap5 = doc.paragraphs[i-1]
        remove_page_break_before(p)
        break

pPr_c5 = p_end_cap5._element.get_or_add_pPr()
for sp in pPr_c5.xpath('./w:sectPr'): pPr_c5.remove(sp)
pPr_c5.append(make_chapter_sectPr())
print("  -> Capítulo V configurado: portadilla sin número, cero logos.")

# 5. Configurar Conclusiones y Glosario (Sección que termina antes de REFERENCIAS)
p_end_concl = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'REFERENCIAS':
        p_end_concl = doc.paragraphs[i-1]
        break

pPr_concl = p_end_concl._element.get_or_add_pPr()
for sp in pPr_concl.xpath('./w:sectPr'): pPr_concl.remove(sp)
new_sp_concl = parse_xml(f'''
<w:sectPr {nsdecls("w", "r")}>
    <w:headerReference w:type="default" r:id="rId13"/>
    <w:footerReference w:type="default" r:id="rId14"/>
    <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
    <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
    <w:pgNumType w:fmt="upperRoman"/>
    <w:cols w:space="708"/>
    <w:type w:val="nextPage"/>
    <w:docGrid w:linePitch="360"/>
</w:sectPr>
''')
pPr_concl.append(new_sp_concl)
print("  -> Conclusiones y Glosario configurados: upperRoman continuos, cero logos.")

# 6. Configurar REFERENCIAS (termina antes de Anexos)
p_end_ref = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'Anexos':
        p_end_ref = doc.paragraphs[i-1]
        break

pPr_ref = p_end_ref._element.get_or_add_pPr()
for sp in pPr_ref.xpath('./w:sectPr'): pPr_ref.remove(sp)
new_sp_ref = parse_xml(f'''
<w:sectPr {nsdecls("w")}>
    <w:headerReference w:type="default" r:id="rId13"/>
    <w:footerReference w:type="default" r:id="rId10"/>
    <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
    <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
    <w:cols w:space="708"/>
    <w:type w:val="nextPage"/>
    <w:docGrid w:linePitch="360"/>
</w:sectPr>
''')
pPr_ref.append(new_sp_ref)
print("  -> REFERENCIAS configuradas: sin pie de página, cero logos.")

# 7. Configurar Anexos (última sección en body)
body_sp = doc._body._element.xpath('./w:sectPr')[0]
for el in list(body_sp): body_sp.remove(el)
body_sp.append(parse_xml(f'<w:headerReference {nsdecls("w")} w:type="default" r:id="rId13"/>'))
body_sp.append(parse_xml(f'<w:footerReference {nsdecls("w")} w:type="default" r:id="rId10"/>'))
body_sp.append(parse_xml(f'<w:pgSz {nsdecls("w")} w:w="12242" w:h="15842" w:code="1"/>'))
body_sp.append(parse_xml(f'<w:pgMar {nsdecls("w")} w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>'))
body_sp.append(parse_xml(f'<w:cols {nsdecls("w")} w:space="708"/>'))
body_sp.append(parse_xml(f'<w:docGrid {nsdecls("w")} w:linePitch="360"/>'))
print("  -> Anexos configurados: sin pie de página, cero logos.")

doc.save(test_path)
print("\nDocumento de prueba guardado exitosamente.")
