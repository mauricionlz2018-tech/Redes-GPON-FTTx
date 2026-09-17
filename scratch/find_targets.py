import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
queries = [
    '2.5.2 Entorno',
    '2.5.5 Biblioteca',
    '2.6.2 Service Workers',
    '2.7.3 Autenticaci',
    '2.7.4 Control de Acceso',
    '2.8.1 Virtualizaci',
    '2.8.4 Arquitectura de Base de Datos Serverless'
]

for q in queries:
    found = False
    for i, p in enumerate(doc.paragraphs):
        if q.lower() in p.text.lower():
            print(f"Match for '{q}': P[{i}] -> '{p.text[:70]}'")
            # Show next 3 paragraphs
            for off in range(1, 4):
                if i + off < len(doc.paragraphs):
                    print(f"   P[{i+off}]: '{doc.paragraphs[i+off].text[:60]}'")
            found = True
            break
    if not found:
        print(f"NOT FOUND: '{q}'")

