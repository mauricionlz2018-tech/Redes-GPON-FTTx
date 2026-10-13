import docx
import zipfile
import io

doc_dev = docx.Document('docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx')
doc_main = docx.Document('docs/Documentacion_Residencias_avance.docx')

with zipfile.ZipFile('docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx') as z_dev:
    dev_files = z_dev.namelist()
    print(f"Archivos en zip de dev: {len(dev_files)}")
    
    cap4_rids = []
    in_c4 = False
    for p in doc_dev.paragraphs:
        t = p.text.strip()
        if 'CAPÍTULO IV' in t:
            in_c4 = True
        elif 'CAPÍTULO V' in t:
            in_c4 = False
            break
        if in_c4:
            for blip in p._element.xpath('.//a:blip/@r:embed'):
                cap4_rids.append(blip)

    print(f"Total blips encontrados en Cap IV: {len(cap4_rids)}")
    for rid in cap4_rids:
        rel = doc_dev.part.rels[rid]
        target = 'word/' + rel.target_ref if not rel.target_ref.startswith('word/') else rel.target_ref
        exists = target in dev_files
        print(f"  {rid} -> {target} (en zip: {exists})")

