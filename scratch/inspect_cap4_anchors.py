import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

c4_idx = None
c5_idx = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO IV' in p.text.upper() and c4_idx is None:
        c4_idx = i
    elif 'CAPÍTULO V' in p.text.upper() and c5_idx is None:
        c5_idx = i

print(f"c4_idx={c4_idx}, c5_idx={c5_idx}")
for k in range(c4_idx - 6, c5_idx + 4):
    p = doc.paragraphs[k]
    has_sp = bool(p._element.xpath('./w:pPr/w:sectPr'))
    pbb = bool(p._element.xpath('./w:pPr/w:pageBreakBefore'))
    print(f"P{k} [style={p.style.name}, sectPr={has_sp}, pbb={pbb}]: '{p.text.strip()}'")

