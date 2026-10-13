import docx
import glob

for path in sorted(glob.glob('docs/*.docx')):
    try:
        d = docx.Document(path)
        print(f"=== {path} (paras={len(d.paragraphs)}) ===")
        for i in range(min(25, len(d.paragraphs))):
            t = d.paragraphs[i].text.strip()
            if t:
                print(f"  P{i} [{d.paragraphs[i].style.name}]: {t[:75]}")
    except Exception as e:
        print(f"Error {path}: {e}")

