import docx
from copy import deepcopy

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
print("Testing deepcopy of table...")
tbl_elm = deepcopy(doc.tables[2]._element)
p = doc.paragraphs[0]._p
p.addnext(tbl_elm)
print("Deepcopy succeeded!")

