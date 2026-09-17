import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
print("Document paragraphs:", len(doc.paragraphs))

drawing_paragraphs = []
for i, p in enumerate(doc.paragraphs):
    if '<w:drawing' in p._p.xml or '<w:pict' in p._p.xml:
        drawing_paragraphs.append((i, p.text[:40]))

print(f"Total drawing paragraphs: {len(drawing_paragraphs)}")
for idx, txt in drawing_paragraphs:
    prev = doc.paragraphs[idx-1].text[:40] if idx > 0 else ""
    nxt = doc.paragraphs[idx+1].text[:60] if idx+1 < len(doc.paragraphs) else ""
    print(f"[{idx}] prev='{prev}' | nxt='{nxt}'")
