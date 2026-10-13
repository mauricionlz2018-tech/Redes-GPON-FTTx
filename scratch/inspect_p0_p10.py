import docx

doc = docx.Document('docs/Documentacion_Residencias_avance_BACKUP_08OCT_FINAL_PAG.docx')
print("--- P0 TO P10 ---")
for i in range(11):
    p = doc.paragraphs[i]
    has_sp = bool(p._element.xpath('./w:pPr/w:sectPr'))
    print(f"P{i} [sectPr={has_sp}]: '{p.text}'")

print("\n--- P155 TO P165 ---")
for i in range(155, 166):
    p = doc.paragraphs[i]
    has_sp = bool(p._element.xpath('./w:pPr/w:sectPr'))
    print(f"P{i} [sectPr={has_sp}]: '{p.text}'")

