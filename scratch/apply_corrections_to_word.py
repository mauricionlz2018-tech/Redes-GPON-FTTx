import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_run(run, name="Arial", size_pt=11.0, bold=False, italic=False, color_rgb=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb

def insert_p(ref_p, text, font_size=11.0, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=10, space_after=14, line_spacing=1.5, style="Normal", color_rgb=None):
    p = ref_p.insert_paragraph_before()
    p.style = style
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    run = p.add_run(text)
    format_run(run, name="Arial", size_pt=font_size, bold=bold, italic=italic, color_rgb=color_rgb)
    return p

def insert_figure(ref_p, img_path, width_in=6.0, caption_text="", elaboration_text="Elaboración propia"):
    p_img = ref_p.insert_paragraph_before()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_in))
    
    p_cap = ref_p.insert_paragraph_before()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(2)
    run_cap = p_cap.add_run(caption_text)
    format_run(run_cap, name="Arial", size_pt=9.5, bold=True)
    
    p_elab = ref_p.insert_paragraph_before()
    p_elab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_elab.paragraph_format.space_before = Pt(0)
    p_elab.paragraph_format.space_after = Pt(12)
    run_elab = p_elab.add_run(elaboration_text)
    format_run(run_elab, name="Arial", size_pt=9.0, italic=True)

def main():
    doc_path = 'docs/PORTADA_INSTITUCIONAL (3).docx'
    print(f"Loading {doc_path}...")
    doc = docx.Document(doc_path)
    
    # -------------------------------------------------------------
    # 1. ERRADICAR EMOJIS E ICONOS EN TODO EL DOCUMENTO
    # -------------------------------------------------------------
    print("Erradicando emojis e iconos en el documento...")
    emoji_replacements = [
        ("[✓]", "icono de verificación"),
        ("[ 🔒 ]", "icono de candado de seguridad"),
        ("[🔒]", "icono de candado de seguridad"),
        ("[⚠]", "icono de advertencia triangular"),
        ("[✕]", "icono de aspa de bloqueo"),
        ("✓", "•"),
        ("🔒", ""),
        ("⚠", ""),
        ("✕", "")
    ]
    for p in doc.paragraphs:
        for emo, rep in emoji_replacements:
            if emo in p.text:
                p.text = p.text.replace(emo, rep)
                for r in p.runs:
                    if emo in r.text:
                        r.text = r.text.replace(emo, rep)

    # -------------------------------------------------------------
    # 2. ACTUALIZAR DIAGRAMA DE CASOS DE USO LIMPIO (FIGURA 5)
    # -------------------------------------------------------------
    print("Actualizando imagen del diagrama de casos de uso (Figura 5)...")
    # Find the paragraph with Figure 5
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Figura 5. Diagrama general de casos de uso"):
            # The image paragraph is p - 1
            p_img = doc.paragraphs[i - 1]
            p_img.clear()
            run_new = p_img.add_run()
            run_new.add_picture('scratch/diagrama_casos_de_uso_clean.png', width=Inches(6.2))
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            print("Replaced Use Case Diagram with clean image.")
            break

    # -------------------------------------------------------------
    # 3. REEMPLAZAR MAQUETADO POR VISTAS INDIVIDUALES (PARTE POR PARTE)
    # -------------------------------------------------------------
    print("Reemplazando maquetado único por vistas individuales parte por parte...")
    # Find Section 3.7.1 and its figure
    idx_371 = None
    for i, p in enumerate(doc.paragraphs):
        if "3.7.1 Diseño Centrado en el Usuario" in p.text or "3.7.1 Dise" in p.text:
            idx_371 = i
            break

    if idx_371 is not None:
        # Find where Figure 8 currently is
        idx_fig8 = None
        for j in range(idx_371, idx_371 + 30):
            if doc.paragraphs[j].text.strip().startswith("Figura 8. Maquetado de interfaces"):
                idx_fig8 = j
                break
        
        if idx_fig8 is not None:
            # We want to clear the old single image and caption, and replace with 4 individual mockups
            # Image paragraph is idx_fig8 - 1, caption is idx_fig8, elaboration is idx_fig8 + 1
            p_old_img = doc.paragraphs[idx_fig8 - 1]
            p_old_cap = doc.paragraphs[idx_fig8]
            p_old_elab = doc.paragraphs[idx_fig8 + 1]
            
            p_old_img.clear()
            p_old_cap.clear()
            p_old_elab.clear()
            
            # Now let's insert the 4 views before the paragraph right after old_elab (idx_fig8 + 2)
            ref_insert = doc.paragraphs[idx_fig8 + 2]
            
            # --- VIEW 1: Visor Cartográfico GIS ---
            insert_p(ref_insert, "A continuación se desglosa la arquitectura visual de la plataforma pantalla por pantalla:")
            insert_p(ref_insert, "1. Visor Cartográfico Geoespacial (GIS): Constituye el entorno principal de monitoreo para supervisores e ingenieros. Despliega la cartografía de San José del Rincón mediante mosaicos OpenStreetMap, representando las cajas NAP con marcadores semafóricos circulares que indican su nivel de ocupación en tiempo real, junto con el trazado vectorial de los cables troncales de fibra óptica desde el ODF central.")
            insert_figure(ref_insert, 'scratch/mockup_gis_map.png', width_in=6.0,
                          caption_text="Figura 8. Maquetado de interfaz: Visor cartográfico geoespacial interactivo (GIS).",
                          elaboration_text="Elaboración propia.")
            
            # --- VIEW 2: Matriz Isomórfica 16 Puertos SC-APC ---
            insert_p(ref_insert, "2. Matriz Isomórfica de Chasis de 16 Puertos SC-APC: Reproduce fielmente la disposición física 2×8 del panel de conectores de la caja terminal NAP. Cada puerto dispone de un indicador LED dinámico con su estado operativo. Al seleccionar un conector libre, se despliega la interfaz modal de asignación rápida que ejecuta la orden con bloqueo pesimista ACID (SELECT ... FOR UPDATE) y valida la dirección MAC de la ONT.")
            insert_figure(ref_insert, 'scratch/mockup_chassis_matrix.png', width_in=6.0,
                          caption_text="Figura 9. Maquetado de interfaz: Matriz isomórfica de chasis de 16 puertos SC-APC y modal transaccional.",
                          elaboration_text="Elaboración propia.")

            # --- VIEW 3: Móvil Offline-First PWA ---
            insert_p(ref_insert, "3. Módulo Móvil Offline-First (PWA): Diseñado ergonómicamente para teléfonos inteligentes en campo. Cuando la cuadrilla pierde conectividad en zonas rurales, la interfaz activa un banner superior en color ámbar informando 'Modo sin conexión activo', almacena las asignaciones en una cola local en IndexedDB (Dexie.js) y proporciona un botón destacado para sincronizar las transacciones al recuperar cobertura.")
            insert_figure(ref_insert, 'scratch/mockup_mobile_offline.png', width_in=6.0,
                          caption_text="Figura 10. Maquetado de interfaz: Arquitectura móvil PWA Offline-First y sincronización en campo.",
                          elaboration_text="Elaboración propia.")

            # --- VIEW 4: Directorio de Clientes y Reportes ---
            insert_p(ref_insert, "4. Directorio Consolidado de Clientes y Auditoría PDF: Dispone el padrón de abonados FTTx en una tabla interactiva con búsqueda reactiva en tiempo real por número de contrato, nombre, dirección MAC o modelo de módem ONT. En la sección inferior integra las métricas globales de saturación y el botón de exportación para la generación de reportes ejecutivos mediante streaming con PDFKit.")
            insert_figure(ref_insert, 'scratch/mockup_clients_reports.png', width_in=6.0,
                          caption_text="Figura 11. Maquetado de interfaz: Directorio consolidado de clientes y consola de reportes ejecutivos.",
                          elaboration_text="Elaboración propia.")
            
            print("Successfully inserted 4 individual wireframe mockups.")

    # -------------------------------------------------------------
    # 4. ACTUALIZAR TABLA Y GRAFICA DE ADOBE COLOR
    # -------------------------------------------------------------
    print("Actualizando tabla de contraste de Adobe Color e imagen de Analizador...")
    # Find Table 13 (Adobe Color palette table)
    # The table has "Rol de Interfaz" in cell 0,0
    tbl_color = None
    for tbl in doc.tables:
        if len(tbl.rows) > 0 and len(tbl.columns) >= 5:
            if "Rol de Interfaz" in tbl.cell(0, 0).text:
                tbl_color = tbl
                break

    if tbl_color is not None:
        # We will replace the table content with the 6-column contrast detailed table
        # Remove old rows and rebuild with 6 columns
        # In python-docx, easiest way is to add the new table and replace XML or clear cells
        # Let's see the parent paragraph before which tbl_color sits
        pass # We will replace it below

    # Find Section 3.7.2 paragraph to rebuild Table 13 cleanly
    p_372 = None
    for p in doc.paragraphs:
        if "3.7.2 Definición de la Paleta Cromática" in p.text or "3.7.2 Definici" in p.text:
            p_372 = p
            break

    if p_372 is not None:
        # Find the old table 13 caption
        for idx, p in enumerate(doc.paragraphs):
            if "Tabla 13. Paleta cromática" in p.text or "Tabla 13. Paleta crom" in p.text:
                p_t13_cap = p
                # The table is right before this caption!
                # Let's locate the table in doc.tables
                for t in doc.tables:
                    if len(t.rows) > 0 and "Rol de Interfaz" in t.cell(0, 0).text:
                        # Replace tbl XML with new 6-col table
                        new_t = doc.add_table(rows=8, cols=6)
                        set_table_borders(new_t)
                        new_t.alignment = WD_TABLE_ALIGNMENT.CENTER
                        
                        hdrs = ["Rol de Interfaz", "Denominación y Código HEX", "Función Semántica en el Sistema", "vs Fondo Blanco (#FFFFFF)", "vs Fondo Oscuro (#0F172A)", "Cumplimiento WCAG 2.1"]
                        for c_idx, h in enumerate(hdrs):
                            cell = new_t.cell(0, c_idx)
                            set_cell_background(cell, "1E3A8A")
                            set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
                            p_c = cell.paragraphs[0]
                            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            r_h = p_c.add_run(h)
                            format_run(r_h, name="Arial", size_pt=8.5, bold=True, color_rgb=RGBColor(255, 255, 255))
                            
                        rows_data = [
                            ("Primario Institucional", "Azul Índigo Telecom\n#1E3A8A", "Barras de navegación fijas, cabeceras y branding corporativo.", "10.36 : 1\n(Pasa AAA)", "1.72 : 1\n(Falla texto)", "Pasa AAA para texto normal y grande sobre blanco."),
                            ("Acento Interactivo", "Cian Técnico\n#0284C7", "Botones de acción principal, enlaces activos y foco de selección.", "4.10 : 1\n(Pasa UI / AA)", "4.36 : 1\n(Pasa UI / AA)", "Pasa AA para componentes de interfaz táctiles y texto grande."),
                            ("Estado Óptimo / Libre", "Verde Esmeralda\n#10B981", "Puertos libres habilitados, cajas con saturación <80% y estado Online.", "2.54 : 1\n(Ajuste #065F46: 7.8:1)", "7.04 : 1\n(Pasa AAA)", "Pasa AAA en modo oscuro. En modo claro se usa Verde Oscuro #065F46."),
                            ("Estado Preventivo / Alerta", "Ámbar Alerta\n#F59E0B", "Puertos reservados, saturación 80-99% y sincronizaciones pendientes.", "2.15 : 1\n(Ajuste #92400E: 6.9:1)", "8.31 : 1\n(Pasa AAA)", "Pasa AAA en modo oscuro. En modo claro se usa Ámbar Oscuro #92400E."),
                            ("Estado Crítico / Dañado", "Rojo Carmesí\n#EF4444", "Puertos dañados, cajas saturadas al 100% y anomalías ópticas.", "3.76 : 1\n(Ajuste #991B1B: 7.2:1)", "4.74 : 1\n(Pasa AA)", "Pasa AA en modo oscuro. En modo claro para texto se usa Rojo #991B1B."),
                            ("Perfil Administrador", "Púrpura RBAC\n#7C3AED", "Identificación visual de perfil Admin en el conmutador de roles.", "5.70 : 1\n(Pasa AA)", "3.13 : 1\n(Pasa UI 3:1)", "Pasa AA para texto sobre blanco y componentes en modo oscuro."),
                            ("Superficie Dark Mode", "Pizarra Oscura\n#0F172A", "Fondo de modo oscuro de alto contraste para visores en campo.", "17.85 : 1\n(Pasa AAA)", "Base de referencia\n(Superficie)", "Pasa AAA con contraste máximo para visores nocturnos.")
                        ]
                        for r_idx, r_vals in enumerate(rows_data):
                            for c_idx, val in enumerate(r_vals):
                                cell = new_t.cell(r_idx + 1, c_idx)
                                bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                                set_cell_background(cell, bg)
                                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                                p_c = cell.paragraphs[0]
                                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
                                r_v = p_c.add_run(val)
                                is_b = True if c_idx in [0, 1] else False
                                format_run(r_v, name="Arial", size_pt=8.0, bold=is_b)
                                
                        # Replace the old table with new table
                        t._tbl.getparent().replace(t._tbl, new_t._tbl)
                        print("Replaced Table 13 with 6-column contrast detailed table.")
                        break
                break

    # Now replace Figure 9 (old Adobe Color image) with new Figure 12 (Adobe Color Contrast Analyzer)
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Figura 9. Evaluación de la paleta cromática"):
            # Old image is p - 1
            p_img = doc.paragraphs[i - 1]
            p_img.clear()
            r_pic = p_img.add_run()
            r_pic.add_picture('scratch/adobe_color_contrast_analyzer.png', width=Inches(6.0))
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            p.text = "Figura 12. Evaluación de contraste en Adobe Color Contrast Analyzer y cumplimiento WCAG 2.1."
            p.runs[0].text = p.text
            format_run(p.runs[0], name="Arial", size_pt=9.5, bold=True)
            print("Replaced Adobe Color graphic with Contrast Analyzer UI.")
            break

    # -------------------------------------------------------------
    # 5. RENUMBER SUBSEQUENT FIGURES IN CHAPTER 3
    # -------------------------------------------------------------
    print("Renumbering subsequent figures in Chapter 3...")
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("Figura 10. Arquitectura de componentes"):
            p.text = txt.replace("Figura 10.", "Figura 13.")
            for r in p.runs:
                if "Figura 10." in r.text:
                    r.text = r.text.replace("Figura 10.", "Figura 13.")
        elif txt.startswith("Figura 11. Flujo de decisi"):
            p.text = txt.replace("Figura 11.", "Figura 14.")
            for r in p.runs:
                if "Figura 11." in r.text:
                    r.text = r.text.replace("Figura 11.", "Figura 14.")
        elif txt.startswith("Figura 12. Flujo de generaci"):
            p.text = txt.replace("Figura 12.", "Figura 15.")
            for r in p.runs:
                if "Figura 12." in r.text:
                    r.text = r.text.replace("Figura 12.", "Figura 15.")
        elif txt.startswith("Figura 13. Arquitectura de contenerizaci"):
            p.text = txt.replace("Figura 13.", "Figura 16.")
            for r in p.runs:
                if "Figura 13." in r.text:
                    r.text = r.text.replace("Figura 13.", "Figura 16.")

    # -------------------------------------------------------------
    # 6. ANEXO A: CRONOGRAMA DE ACTIVIDADES (FORMATO OFICIAL UMB CON ASESOR LEONARDO)
    # -------------------------------------------------------------
    print("Actualizando Anexo A con Cronograma Oficial UMB y asesor Leonardo...")
    for i, p in enumerate(doc.paragraphs):
        if "Anexo A. Cronograma de actividades" in p.text:
            # Check paragraphs under Anexo A
            p_anexo = p
            # Locate subsequent paragraphs to update advisor and insert figure
            for j in range(i, min(i + 15, len(doc.paragraphs))):
                pj = doc.paragraphs[j]
                if "Asesor:" in pj.text or "Organización:" in pj.text:
                    pass
                if "Periodo:" in pj.text:
                    # After Periodo, insert the Cronograma Institutional Image
                    p_target = doc.paragraphs[j + 1] if j + 1 < len(doc.paragraphs) else pj
                    
                    # Update or ensure advisor paragraph
                    insert_p(p_target, "Asesor de Residencia Profesional: I.S.C. Leonardo Becerril Sánchez", bold=True, font_size=10.0)
                    
                    insert_figure(
                        p_target,
                        img_path='scratch/cronograma_institucional_umb.png',
                        width_in=6.5,
                        caption_text="Figura 17. Cronograma general de actividades de residencia profesional (Formato Oficial UMB).",
                        elaboration_text="Fuente: Elaboración propia según formato normativo de la Universidad Mexiquense del Bicentenario."
                    )
                    print("Inserted official UMB Cronograma figure with advisor Leonardo Becerril Sánchez.")
                    break
            break

    # Save modified document
    output_path = 'docs/PORTADA_INSTITUCIONAL (3).docx'
    print(f"Saving updated document to {output_path}...")
    doc.save(output_path)
    print("Corrections applied successfully!")

if __name__ == '__main__':
    main()

