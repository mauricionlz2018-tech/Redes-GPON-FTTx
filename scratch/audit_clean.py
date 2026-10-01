import docx
import re

for doc_path in ["docs/Documentacion_Residencias_avance.docx", "docs/Documentacion_Residencias (4).docx"]:
    print(f"\n{'='*70}\nAUDIT: {doc_path}\n{'='*70}")
    doc = docx.Document(doc_path)
    
    # 1. Check 4.1 subsections
    print("--- Subsections in 4.1 ---")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith("4.1."):
            print(f"  P{i} [{p.style.name}]: {txt}")

    # 2. Check figures count in Body
    figs = [p.text.strip() for p in doc.paragraphs if p.style.name == "Figuras"]
    print(f"\nTotal figures in body: {len(figs)}")
    print(f"  First fig: {figs[0][:50]}")
    print(f"  Fig 58: {figs[57][:50]}")
    print(f"  Fig 59: {figs[58][:50]}")
    print(f"  Fig 64: {figs[63][:50]}")
    print(f"  Fig 65: {figs[64][:50]}")
    print(f"  Last fig: {figs[-1][:50]}")

    # 3. Check TOC figures
    toc_figs = [p.text.strip() for p in doc.paragraphs if p.style.name == "toc 1" and p.text.strip().startswith("Figura ")]
    print(f"\nTotal TOC figures: {len(toc_figs)}")
    print(f"  First TOC: {toc_figs[0][:50]}")
    print(f"  TOC 58: {toc_figs[57][:50]}")
    print(f"  TOC 59: {toc_figs[58][:50]}")
    print(f"  TOC 64: {toc_figs[63][:50]}")
    print(f"  TOC 65: {toc_figs[64][:50]}")
    print(f"  Last TOC: {toc_figs[-1][:50]}")

    # 4. Check chapter headings geometry
    print("\n--- Chapter Headings Geometry ---")
    for pat in [r'^CAP[ÍI]TULO\s+I\b', r'^CAP[ÍI]TULO\s+II\b', r'^CAP[ÍI]TULO\s+III\b', r'^CAP[ÍI]TULO\s+IV\b', r'^CAP[ÍI]TULO\s+V\b']:
        for idx, p in enumerate(doc.paragraphs):
            if re.search(pat, p.text.strip(), re.IGNORECASE):
                first_spacer_pb = doc.paragraphs[idx-6].paragraph_format.page_break_before
                title_sz = [r.font.size.pt for r in p.runs if r.font.size]
                title_bold = [r.font.bold for r in p.runs]
                print(f"  {p.text}: PB={first_spacer_pb}, TitleSize={set(title_sz)}, Bold={set(title_bold)}")
                break

print("\nAUDIT COMPLETED!")
