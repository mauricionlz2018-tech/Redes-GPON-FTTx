import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')
for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    if '<w:drawing' in xml or '<w:pict' in xml:
        prev_t = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        next_t = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
        print(f"Drawing at P[{i}]: prev='{prev_t[:60]}' | next='{next_t[:60]}'")

