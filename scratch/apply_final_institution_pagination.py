import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import shutil
import subprocess

def apply_final_pagination():
    doc_path = 'docs/Documentacion_Residencias_avance.docx'
    backup_path = 'docs/Documentacion_Residencias_avance_BACKUP_08OCT_FINAL_PAG.docx'
    
    print(f"Creando respaldo en {backup_path}...")
    shutil.copyfile(doc_path, backup_path)
    print("Respaldo creado con éxito.")

    print(f"Abriendo {doc_path}...")
    doc = docx.Document(doc_path)

    def remove_page_break_before(p):
        pPr = p._element.find(qn('w:pPr'))
        if pPr is not None:
            pbb = pPr.find(qn('w:pageBreakBefore'))
            if pbb is not None:
                pPr.remove(pbb)

    # 1. Quitar sectPr de P159 si existe
    for idx in range(150, 165):
        if idx < len(doc.paragraphs):
            p = doc.paragraphs[idx]
            sects = p._element.xpath('./w:pPr/w:sectPr')
            if sects:
                p._element.pPr.remove(sects[0])
                print(f"  -> Removido sectPr en P{idx} para unificar Preliminares con Introducción y Objetivos.")

    # 2. Localizar puntos de corte exactos:
    # Fin de Objetivos: buscar el último párrafo antes de los vacíos de Capítulo I
    p_end_obj = None
    p_start_cap1 = None
    for i, p in enumerate(doc.paragraphs):
        if '3. Programar el backend del sistema' in p.text:
            p_end_obj = doc.paragraphs[i]
            p_start_cap1 = doc.paragraphs[i+1]
            break

    # Fin de Cap I: buscar párrafo antes de Cap II
    p_end_cap1 = None
    p_start_cap2 = None
    for i, p in enumerate(doc.paragraphs):
        if 'CAPÍTULO II. MARCO' in p.text.upper():
            # Buscar el párrafo previo con pageBreakBefore
            p_end_cap1 = doc.paragraphs[i-8]
            p_start_cap2 = doc.paragraphs[i-7]
            break

    # Fin de Cap II: buscar párrafo antes de Cap III
    p_end_cap2 = None
    p_start_cap3 = None
    for i, p in enumerate(doc.paragraphs):
        if 'CAPÍTULO III. DISEÑO' in p.text.upper():
            p_end_cap2 = doc.paragraphs[i-7]
            p_start_cap3 = doc.paragraphs[i-6]
            break

    # Fin de Cap III: buscar párrafo antes de Cap IV
    p_end_cap3 = None
    p_start_cap4 = None
    for i, p in enumerate(doc.paragraphs):
        if 'CAPÍTULO IV. DESARROLLO' in p.text.upper():
            p_end_cap3 = doc.paragraphs[i-7]
            p_start_cap4 = doc.paragraphs[i-6]
            break

    # Fin de Cap IV: buscar párrafo antes de Cap V
    p_end_cap4 = None
    p_start_cap5 = None
    for i, p in enumerate(doc.paragraphs):
        if 'CAPÍTULO V. PRUEBAS' in p.text.upper():
            p_end_cap4 = doc.paragraphs[i-7]
            p_start_cap5 = doc.paragraphs[i-6]
            break

    splits = [
        (p_end_obj, p_start_cap1, 'Capítulo I'),
        (p_end_cap1, p_start_cap2, 'Capítulo II'),
        (p_end_cap2, p_start_cap3, 'Capítulo III'),
        (p_end_cap3, p_start_cap4, 'Capítulo IV'),
        (p_end_cap4, p_start_cap5, 'Capítulo V'),
    ]

    for p_end, p_start, name in splits:
        if p_end is not None and p_start is not None:
            remove_page_break_before(p_start)
            pPr = p_end._element.get_or_add_pPr()
            for sp in pPr.xpath('./w:sectPr'):
                pPr.remove(sp)
            new_sp = parse_xml(f'''
            <w:sectPr {nsdecls("w")}>
                <w:type w:val="nextPage"/>
            </w:sectPr>
            ''')
            pPr.append(new_sp)
            print(f"  -> Creada sección antes de {name}.")

    # Remover pageBreakBefore de Conclusiones para evitar página en blanco
    for p in doc.paragraphs:
        if p.text.strip() == 'Conclusiones':
            remove_page_break_before(p)
            break

    # Helper para asignar pie de página
    def set_page_field_in_footer(footer):
        for p in footer.paragraphs:
            p.text = ''
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fld_xml = parse_xml(f'''
        <w:p {nsdecls("w")}>
            <w:pPr>
                <w:jc w:val="right"/>
                <w:rPr>
                    <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                    <w:sz w:val="22"/>
                </w:rPr>
            </w:pPr>
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:fldChar w:fldCharType="begin"/>
            </w:r>
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:instrText xml:space="preserve"> PAGE </w:instrText>
            </w:r>
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:fldChar w:fldCharType="separate"/>
            </w:r>
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:t>1</w:t>
            </w:r>
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:fldChar w:fldCharType="end"/>
            </w:r>
        </w:p>
        ''')
        p._element.getparent().replace(p._element, fld_xml)

    def clear_footer(footer):
        footer.is_linked_to_previous = False
        for p in list(footer.paragraphs):
            p.text = ''

    print(f"\nConfigurando propiedades de las {len(doc.sections)} secciones...")
    for i, s in enumerate(doc.sections):
        sectPr = s._sectPr
        for el in sectPr.xpath('./w:pgNumType'):
            sectPr.remove(el)
            
        if i == 0:
            # Sección 0: Carátula
            s.different_first_page_header_footer = True
            clear_footer(s.first_page_footer)
            clear_footer(s.footer)
            sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1" w:fmt="decimal"/>'))
            print("  -> Sec 0 (Carátula): Sin número visible. Cuenta págs. 1 y 2.")
            
        elif i == 1:
            # Sección 1: Preliminares (Índices, Introducción, Planteamiento, Justificación, Objetivos)
            s.different_first_page_header_footer = False
            s.footer.is_linked_to_previous = False
            set_page_field_in_footer(s.footer)
            sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="3" w:fmt="upperRoman"/>'))
            print("  -> Sec 1 (Preliminares: Índices, Introducción, Planteamiento, Justificación, Objetivos): upperRoman iniciando en III.")
            
        elif 2 <= i <= 6:
            # Secciones 2 a 6: Capítulos I, II, III, IV, V
            ch_num = i - 1
            s.different_first_page_header_footer = True
            clear_footer(s.first_page_footer)
            s.footer.is_linked_to_previous = False
            set_page_field_in_footer(s.footer)
            sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal"/>'))
            print(f"  -> Sec {i} (Capítulo {ch_num}): Portadilla sin número visible; contenido en arábigos consecutivos.")
            
        elif i == 7:
            # Sección 7: Conclusiones y Glosario
            s.different_first_page_header_footer = False
            s.footer.is_linked_to_previous = False
            set_page_field_in_footer(s.footer)
            sectPr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="upperRoman"/>'))
            print("  -> Sec 7 (Conclusiones y Glosario): upperRoman consecutivos a la paginación general.")
            
        elif i == 8:
            # Sección 8: REFERENCIAS
            s.different_first_page_header_footer = False
            clear_footer(s.footer)
            print("  -> Sec 8 (REFERENCIAS): Sin número visible.")
            
        elif i == 9:
            # Sección 9: Anexos
            s.different_first_page_header_footer = False
            clear_footer(s.footer)
            print("  -> Sec 9 (Anexos): Sin número visible.")

    doc.save(doc_path)
    print("\nGuardado exitoso en Documentacion_Residencias_avance.docx.")

if __name__ == '__main__':
    apply_final_pagination()

