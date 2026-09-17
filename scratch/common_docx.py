import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from copy import deepcopy
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.style = 'Heading 1'
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.style = 'Heading 2'
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.style = 'Heading 3'
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    return p

def add_paragraph(doc, text, bold_prefix="", indent=False):
    p = doc.add_paragraph()
    p.style = 'Normal'
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.left_indent = Inches(0.25)
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(11)
        r_pre.bold = True
        
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    return p

def add_code_block(doc, code_text, language="TypeScript"):
    # Wrapper box table or single cell for syntax highlighting appearance
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "0F172A") # Slate 900 dark background
    cell.width = Inches(6.5)

    # Padding
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="140" w:type="dxa"/><w:bottom w:w="140" w:type="dxa"/><w:left w:w="200" w:type="dxa"/><w:right w:w="200" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    # Language tag
    r_tag = p.add_run(f"// [{language}]\n")
    r_tag.font.name = 'Consolas'
    r_tag.font.size = Pt(8.5)
    r_tag.font.color.rgb = RGBColor(56, 189, 248) # Sky 400
    r_tag.bold = True

    lines = code_text.strip().split('\n')
    for idx, l in enumerate(lines):
        r = p.add_run(l + ('\n' if idx < len(lines)-1 else ''))
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
        # Check comment or keyword
        stripped = l.strip()
        if stripped.startswith('//') or stripped.startswith('*') or stripped.startswith('/*'):
            r.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
            r.italic = True
        elif stripped.startswith('import ') or stripped.startswith('export ') or stripped.startswith('const ') or stripped.startswith('interface ') or stripped.startswith('class ') or stripped.startswith('public '):
            r.font.color.rgb = RGBColor(129, 140, 248) # Indigo 400
        else:
            r.font.color.rgb = RGBColor(241, 245, 249) # Slate 100

    # Space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(6)
    p_sp.paragraph_format.line_spacing = 1.0

def add_figure(doc, img_path, caption_text, source_text, width_inches=5.8):
    p_img = doc.add_paragraph()
    p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    
    if os.path.exists(img_path):
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(width_inches))
    else:
        print(f"WARNING: Image not found: {img_path}")

    # Caption strictly below
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(2)
    p_cap.paragraph_format.keep_with_next = True
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = 'Arial'
    r_cap.font.size = Pt(9)
    r_cap.bold = True
    r_cap.font.color.rgb = RGBColor(15, 23, 42)

    # Source line strictly below
    p_src = doc.add_paragraph()
    p_src.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_src.paragraph_format.space_before = Pt(0)
    p_src.paragraph_format.space_after = Pt(14)
    r_src = p_src.add_run(source_text)
    r_src.font.name = 'Arial'
    r_src.font.size = Pt(8.5)
    r_src.italic = True
    r_src.font.color.rgb = RGBColor(71, 85, 105)

def sanitize_text(text):
    # Rule 1: No em-dashes
    text = text.replace('—', '-')
    text = text.replace('–', '-')
    # Rule 2: No semicolons
    text = text.replace(';', ',')
    # Rule 3: No emojis or brackets with checks
    for glyph in ['✓', '🔒', '⚠', '✕', '⚡', '📊', '🛡️', '🗺️', '📱', '📡', '🚀', '✅', '❌']:
        text = text.replace(glyph, '')
    # Rule 4: Ensure Leonardo Becerril Sánchez is advisor
    text = text.replace('Jorge Casio González', 'I.S.C. Leonardo Becerril Sánchez')
    text = text.replace('Jorge Casio Gonzalez', 'I.S.C. Leonardo Becerril Sánchez')
    text = text.replace('Casio', 'Leonardo Becerril Sánchez')
    return text

