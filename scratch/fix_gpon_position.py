import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import shutil

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

# 1. Remove misplaced GPON paragraph around 544
misplaced_p = None
for i, p in enumerate(doc.paragraphs):
    if 'gpon telecom s.a. de c.v. (2026). infraestructura para la vida' in p.text.lower():
        if i < 1000:
            print(f'Found misplaced GPON at P{i}: {p.text}')
            misplaced_p = p
            break

if misplaced_p is not None:
    misplaced_p._element.getparent().remove(misplaced_p._element)
    print('Misplaced paragraph removed from Chapter 2.')

# 2. Find Google Chrome Developers in references section (index > 1200)
ref_google_p = None
for i, p in enumerate(doc.paragraphs):
    if i > 1200 and 'google chrome developers. (2023)' in p.text.lower():
        print(f'Found Google Chrome Developers in references at P{i}: {p.text}')
        ref_google_p = p
        break

if ref_google_p is not None:
    new_gpon_xml = parse_xml(f'<w:p {nsdecls("w")}/>')
    ref_google_p._element.addnext(new_gpon_xml)
    p_gpon = docx.text.paragraph.Paragraph(new_gpon_xml, doc)
    p_gpon.paragraph_format.line_spacing = 1.5
    p_gpon.paragraph_format.space_after = Pt(6)
    p_gpon.paragraph_format.left_indent = Pt(36)
    p_gpon.paragraph_format.first_line_indent = Pt(-36)
    
    r1 = p_gpon.add_run('GPON Telecom S.A. de C.V. (2026). ')
    r2 = p_gpon.add_run('Infraestructura para la vida: Nosotros. ')
    r2.font.italic = True
    r3 = p_gpon.add_run('https://gpontelecom.com.mx/#div_nosotros')
    for r in [r1, r2, r3]:
        r.font.name = 'Arial'
        r.font.size = Pt(11)
    print('GPON reference correctly inserted in Referencias electronicas under G.')

doc.save('docs/Documentacion_Residencias_avance.docx')
shutil.copyfile('docs/Documentacion_Residencias_avance.docx', 'docs/Documentacion_Residencias (4).docx')
print('Saved and mirrored successfully.')

