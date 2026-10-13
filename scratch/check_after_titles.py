import docx

doc = docx.Document('docs/Documentacion_Residencias_avance_BACKUP_08OCT_FINAL_PAG.docx')
for mark in [195, 250, 759, 1198, 1205]:
    for i in range(mark + 1, min(mark + 20, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        if p.text.strip():
            has_pbb = bool(p._element.xpath('.//w:pPr/w:pageBreakBefore'))
            title_text = doc.paragraphs[mark].text.strip()
            print(f"After P{mark} ('{title_text}'): First non-empty is P{i} [pbb={has_pbb}]: '{p.text.strip()[:50]}'")
            break

