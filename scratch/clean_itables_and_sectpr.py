import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import shutil

target_avance = 'docs/Documentacion_Residencias_avance.docx'
target_mirror = 'docs/Documentacion_Residencias (4).docx'

doc = docx.Document(target_avance)

# 1. Eliminar P111 a P128 (recolección previa para no alterar índices)
to_remove = [doc.paragraphs[k]._element for k in range(111, 129)]
for el in to_remove:
    el.getparent().remove(el)

# 2. Localizar el párrafo vacío justo antes de 'Introducción' (que cierra el TOC de tablas)
p_before_intro = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'Introducción':
        p_before_intro = doc.paragraphs[i - 1]
        print(f"Introducción ahora está en P{i}; párrafo previo es P{i-1}")
        break

# 3. Asegurar que p_before_intro tenga el sectPr de Sección 1 (Índices: upperRoman start=3, sin logos, nextPage)
pPr = p_before_intro._element.get_or_add_pPr()
for sp in pPr.xpath('./w:sectPr'):
    pPr.remove(sp)

sp_sec1 = parse_xml(f'''
<w:sectPr {nsdecls("w", "r")}>
    <w:headerReference w:type="default" r:id="rId15"/>
    <w:footerReference w:type="default" r:id="rId16"/>
    <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
    <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
    <w:pgNumType w:fmt="upperRoman" w:start="3"/>
    <w:cols w:space="708"/>
    <w:type w:val="nextPage"/>
    <w:docGrid w:linePitch="360"/>
</w:sectPr>
''')
pPr.append(sp_sec1)

doc.save(target_avance)
shutil.copyfile(target_avance, target_mirror)

# Verificación inmediata
doc_v = docx.Document(target_avance)
toc_tables = [p.text.strip() for p in doc_v.paragraphs[:120] if p.text.strip().startswith('Tabla ') and '\t' in p.text]
print(f"Total entradas únicas en Índice de tablas: {len(toc_tables)}")
print(f"Total secciones en el documento: {len(doc_v.sections)}")
for idx_s, s in enumerate(doc_v.sections):
    h_draw = len(s.header._element.xpath('.//w:drawing'))
    print(f"  Sec {idx_s}: header_drawings={h_draw}")

