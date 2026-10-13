import docx
import shutil

target_avance = 'docs/Documentacion_Residencias_avance.docx'
target_mirror = 'docs/Documentacion_Residencias (4).docx'

doc = docx.Document(target_avance)

# Encontrar la primera 'Tabla 36.' y el párrafo de cierre antes de 'Introducción'
first_t36_idx = None
intro_idx = None
for i, p in enumerate(doc.paragraphs):
    if first_t36_idx is None and p.text.strip().startswith('Tabla 36.'):
        first_t36_idx = i
    if p.text.strip() == 'Introducción':
        intro_idx = i
        break

print(f"first_t36_idx={first_t36_idx}, intro_idx={intro_idx}")

# Los párrafos entre first_t36_idx + 1 y intro_idx - 2 son el segundo TOC duplicado
# Verifiquemos qué son intro_idx - 2 e intro_idx - 1:
for k in range(first_t36_idx, intro_idx + 1):
    p = doc.paragraphs[k]
    has_sp = bool(p._element.xpath('./w:pPr/w:sectPr'))
    flds = [(el.tag.split('}')[-1], el.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType')) for el in p._element.xpath('.//w:fldChar')]
    print(f"P{k} [sectPr={has_sp}, flds={flds}]: '{p.text.strip()[:50]}'")

