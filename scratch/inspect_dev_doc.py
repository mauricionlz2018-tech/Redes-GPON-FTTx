import docx

doc_dev = docx.Document('docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx')
print(f"Total párrafos en doc_dev: {len(doc_dev.paragraphs)}")
print(f"Total tablas en doc_dev: {len(doc_dev.tables)}")

for i, p in enumerate(doc_dev.paragraphs):
    if p.style.name.startswith('Heading') or p.text.strip().startswith('CAPÍTULO'):
        print(f"P{i} [{p.style.name}]: {p.text}")

