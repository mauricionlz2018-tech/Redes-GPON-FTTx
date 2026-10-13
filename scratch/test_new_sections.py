import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import shutil

src_path = 'docs/Documentacion_Residencias_avance.docx'
test_path = 'scratch/test_new_sections.docx'
shutil.copyfile(src_path, test_path)

doc = docx.Document(test_path)

# Let's inspect paragraphs around:
# 1. P159 (current section break between indices and introduccion)
# We want to REMOVE sectPr from P159.
p159 = doc.paragraphs[159]
p159_sects = p159._element.xpath('./w:pPr/w:sectPr')
if p159_sects:
    p159._element.pPr.remove(p159_sects[0])
    print("Removed sectPr from P159. Now Introduccion follows Preliminaries.")

# 2. Before Capitulo I:
# Where does Capitulo I start?
# In our previous inspection, P187 had pageBreakBefore=True, P195 was CAPITULO I.
# The paragraph that should end the Roman section is P186 (the last paragraph of Objetivos)!
# If P186 ends the Roman section, then P187 starts the next section (Capitulo I)!
p186 = doc.paragraphs[186]
# Let's check P186
print(f"P186 text: {repr(p186.text[:40])}")

doc.save(test_path)
print("Saved basic test.")

