import docx
from copy import deepcopy

source_doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')
tables_source = [deepcopy(t) for t in source_doc.tables]
print(f"Captured {len(tables_source)} source tables.")

# Test clearing body in target
target_doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_PRE_RESTRUCTURE.docx')
body = target_doc._body._element

# Retain sectPr
sectPr = body.xpath('w:sectPr')
# Remove all children
for child in list(body):
    if child.tag.endswith('sectPr'):
        continue
    body.remove(child)

print(f"Body elements remaining: {len(body)}")
# Add a test paragraph
p = target_doc.add_paragraph("Test after clear")
target_doc.save('scratch/test_cleared.docx')
print("Saved scratch/test_cleared.docx successfully!")

