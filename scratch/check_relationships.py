import docx

doc_dev = docx.Document('docs/Documentacion_Desarrollo_Codificacion_y_Pruebas.docx')
doc_main = docx.Document('docs/Documentacion_Residencias_avance.docx')

dev_rels = doc_dev.part.rels
main_rels = doc_main.part.rels

# Find all rIds used in Cap IV of doc_dev
cap4_rids = set()
in_c4 = False
for p in doc_dev.paragraphs:
    t = p.text.strip()
    if 'CAPÍTULO IV' in t:
        in_c4 = True
    elif 'CAPÍTULO V' in t:
        in_c4 = False
        break
    if in_c4:
        for blip in p._element.xpath('.//a:blip/@r:embed'):
            cap4_rids.add(blip)

print(f"Total rIds used in Cap IV drawings: {len(cap4_rids)}")
missing_in_main = [rid for rid in cap4_rids if rid not in main_rels]
print(f"rIds missing in doc_main: {len(missing_in_main)}: {missing_in_main}")

matching_targets = []
for rid in cap4_rids:
    if rid in main_rels:
        t_dev = dev_rels[rid].target_ref
        t_main = main_rels[rid].target_ref
        matching_targets.append(t_dev == t_main)
print(f"For existing rIds, do targets match? {all(matching_targets) if matching_targets else 'N/A'}")

