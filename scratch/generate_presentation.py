import os
import math
from PIL import Image, ImageDraw, ImageFont
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# -------------------------------------------------------------
# 1. GENERATE HIGH-RES BACKGROUND TEMPLATES
# -------------------------------------------------------------
def generate_background_images():
    w, h = 1920, 1080
    
    # Palette
    C_LIGHT = (205, 216, 160, 255) # Light sage / olive wave
    C_DARK = (123, 139, 59, 255)   # Main institutional olive green wave
    
    # A. Content Background
    bg_content = Image.new('RGBA', (w, h), (255, 255, 255, 255))
    draw_c = ImageDraw.Draw(bg_content)
    
    # Upper wave
    pts_light = [(0, 800)]
    for x in range(0, w + 1, 10):
        y = 800 + 90 * math.sin((x / w) * 2.2) + 20 * math.cos((x / w) * 1.5) - 20
        pts_light.append((x, y))
    pts_light.extend([(w, h), (0, h)])
    draw_c.polygon(pts_light, fill=C_LIGHT)
    
    # Lower wave
    pts_dark = [(0, 840)]
    for x in range(0, w + 1, 10):
        y = 840 + 95 * math.sin((x / w) * 2.1) + 20 * math.cos((x / w) * 1.5) - 20
        pts_dark.append((x, y))
    pts_dark.extend([(w, h), (0, h)])
    draw_c.polygon(pts_dark, fill=C_DARK)
    
    # UMB Typographic Logo
    try:
        font_umb = ImageFont.truetype('arialbd.ttf', 70)
        font_sub = ImageFont.truetype('arialbd.ttf', 17)
    except:
        font_umb = ImageFont.load_default()
        font_sub = ImageFont.load_default()
        
    draw_c.text((90, 910), 'UMB', font=font_umb, fill=(255, 255, 255, 255))
    draw_c.text((90, 995), 'UNIVERSIDAD MEXIQUENSE DEL BICENTENARIO', font=font_sub, fill=(255, 255, 255, 255))
    
    bg_content.save('scratch/slide_bg_content.png')
    
    # B. Cover Background
    bg_cover = Image.new('RGBA', (w, h), (255, 255, 255, 255))
    draw_cov = ImageDraw.Draw(bg_cover)
    
    # Upper wave for cover (slightly higher)
    pts_cov_light = [(0, 780)]
    for x in range(0, w + 1, 10):
        y = 780 + 100 * math.sin((x / w) * 2.2) + 20 * math.cos((x / w) * 1.5) - 20
        pts_cov_light.append((x, y))
    pts_cov_light.extend([(w, h), (0, h)])
    draw_cov.polygon(pts_cov_light, fill=C_LIGHT)
    
    # Lower wave for cover
    pts_cov_dark = [(0, 820)]
    for x in range(0, w + 1, 10):
        y = 820 + 105 * math.sin((x / w) * 2.1) + 20 * math.cos((x / w) * 1.5) - 20
        pts_cov_dark.append((x, y))
    pts_cov_dark.extend([(w, h), (0, h)])
    draw_cov.polygon(pts_cov_dark, fill=C_DARK)
    
    draw_cov.text((90, 900), 'UMB', font=font_umb, fill=(255, 255, 255, 255))
    draw_cov.text((90, 985), 'UNIVERSIDAD MEXIQUENSE DEL BICENTENARIO', font=font_sub, fill=(255, 255, 255, 255))
    
    # Subtle top accent bar in olive green
    draw_cov.rectangle([(0, 0), (w, 14)], fill=C_DARK)
    
    bg_cover.save('scratch/slide_bg_cover.png')
    print("Generated slide_bg_content.png and slide_bg_cover.png successfully.")

# -------------------------------------------------------------
# 2. POWERPOINT BUILDER
# -------------------------------------------------------------
def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout
    
    # Color Constants
    CLR_OLIVE = RGBColor(123, 139, 59)     # #7B8B3B
    CLR_DARK_OLIVE = RGBColor(90, 104, 40) # #5A6828
    CLR_TEXT_DARK = RGBColor(31, 41, 55)   # #1F2937 Charcoal
    CLR_TEXT_MUTED = RGBColor(75, 85, 99)  # #4B5563
    CLR_BG_CARD = RGBColor(248, 250, 245)  # Light olive-tinted card
    CLR_BORDER_CARD = RGBColor(218, 226, 185) # Soft sage border
    CLR_WHITE = RGBColor(255, 255, 255)
    
    def add_top_right_title(slide, title_text):
        # Calculate width dynamically based on text length
        width_in = max(4.0, len(title_text) * 0.16 + 1.2)
        left_in = 13.333 - width_in - 0.8
        
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_in), Inches(0.45), Inches(width_in), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = CLR_WHITE
        box.line.color.rgb = CLR_OLIVE
        box.line.width = Pt(1.5)
        
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title_text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = CLR_DARK_OLIVE
        return box

    # =========================================================
    # SLIDE 1: PORTADA
    # =========================================================
    slide1 = prs.slides.add_slide(blank_layout)
    slide1.shapes.add_picture('scratch/slide_bg_cover.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    
    # Institutional Crest / Logo (Top-Left)
    if os.path.exists('scratch/logo_portada.png'):
        slide1.shapes.add_picture('scratch/logo_portada.png', Inches(0.8), Inches(0.45), height=Inches(1.2))
        
    # Company Logo (Top-Right)
    if os.path.exists('backend/assets/logo-gpon.png'):
        slide1.shapes.add_picture('backend/assets/logo-gpon.png', Inches(10.2), Inches(0.45), height=Inches(1.1))
        
    # Institutional Header (Center)
    head_box = slide1.shapes.add_textbox(Inches(2.2), Inches(0.4), Inches(7.8), Inches(1.2))
    tf_head = head_box.text_frame
    tf_head.word_wrap = True
    
    p1 = tf_head.paragraphs[0]
    p1.text = "UNIVERSIDAD MEXIQUENSE DEL BICENTENARIO"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = "Arial"
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = CLR_DARK_OLIVE
    
    p2 = tf_head.add_paragraph()
    p2.text = "Unidad de Estudios Superiores San José del Rincón"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = "Arial"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = CLR_TEXT_DARK
    
    p3 = tf_head.add_paragraph()
    p3.text = "Ingeniería en Sistemas Computacionales"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = "Arial"
    p3.font.size = Pt(12)
    p3.font.color.rgb = CLR_TEXT_MUTED
    
    # Project Title Card (Center)
    title_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.85), Inches(10.933), Inches(2.2))
    title_box.fill.solid()
    title_box.fill.fore_color.rgb = CLR_BG_CARD
    title_box.line.color.rgb = CLR_BORDER_CARD
    title_box.line.width = Pt(1.5)
    
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    pt0 = tf_title.paragraphs[0]
    pt0.text = "PROYECTO DE RESIDENCIA PROFESIONAL"
    pt0.alignment = PP_ALIGN.CENTER
    pt0.font.name = "Arial"
    pt0.font.size = Pt(12)
    pt0.font.bold = True
    pt0.font.color.rgb = CLR_OLIVE
    
    pt1 = tf_title.add_paragraph()
    pt1.text = "Sistema de Inventario y Mapeo Lógico de Redes GPON / FTTx"
    pt1.alignment = PP_ALIGN.CENTER
    pt1.font.name = "Arial"
    pt1.font.size = Pt(25)
    pt1.font.bold = True
    pt1.font.color.rgb = CLR_TEXT_DARK
    
    pt2 = tf_title.add_paragraph()
    pt2.text = "Empresa: GPON TELECOM S.A. de C.V.  |  San José del Rincón, Estado de México"
    pt2.alignment = PP_ALIGN.CENTER
    pt2.font.name = "Arial"
    pt2.font.size = Pt(12.5)
    pt2.font.color.rgb = CLR_TEXT_MUTED
    
    # Presenter & Advisor Info (Two Cards Bottom-Middle)
    # Left Card: Presenta
    card_pres = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.25), Inches(4.8), Inches(1.6))
    card_pres.fill.solid()
    card_pres.fill.fore_color.rgb = CLR_WHITE
    card_pres.line.color.rgb = CLR_BORDER_CARD
    card_pres.line.width = Pt(1)
    tf_pres = card_pres.text_frame
    tf_pres.word_wrap = True
    
    pp0 = tf_pres.paragraphs[0]
    pp0.text = "PRESENTA:"
    pp0.font.name = "Arial"
    pp0.font.size = Pt(11)
    pp0.font.bold = True
    pp0.font.color.rgb = CLR_DARK_OLIVE
    
    pp1 = tf_pres.add_paragraph()
    pp1.text = "Mauricio Nolazco Lonjino"
    pp1.font.name = "Arial"
    pp1.font.size = Pt(15)
    pp1.font.bold = True
    pp1.font.color.rgb = CLR_TEXT_DARK
    
    pp2 = tf_pres.add_paragraph()
    pp2.text = "Matrícula: Residente Profesional UMB"
    pp2.font.name = "Arial"
    pp2.font.size = Pt(11)
    pp2.font.color.rgb = CLR_TEXT_MUTED
    
    # Right Card: Asesor
    card_adv = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(4.25), Inches(4.8), Inches(1.6))
    card_adv.fill.solid()
    card_adv.fill.fore_color.rgb = CLR_WHITE
    card_adv.line.color.rgb = CLR_BORDER_CARD
    card_adv.line.width = Pt(1)
    tf_adv = card_adv.text_frame
    tf_adv.word_wrap = True
    
    pa0 = tf_adv.paragraphs[0]
    pa0.text = "ASESOR DE RESIDENCIA PROFESIONAL:"
    pa0.font.name = "Arial"
    pa0.font.size = Pt(11)
    pa0.font.bold = True
    pa0.font.color.rgb = CLR_DARK_OLIVE
    
    pa1 = tf_adv.add_paragraph()
    pa1.text = "I.S.C. Leonardo Becerril Sánchez"
    pa1.font.name = "Arial"
    pa1.font.size = Pt(15)
    pa1.font.bold = True
    pa1.font.color.rgb = CLR_TEXT_DARK
    
    pa2 = tf_adv.add_paragraph()
    pa2.text = "Asesor Institucional - UES San José del Rincón"
    pa2.font.name = "Arial"
    pa2.font.size = Pt(11)
    pa2.font.color.rgb = CLR_TEXT_MUTED

    # =========================================================
    # SLIDE 2: INTRODUCCIÓN
    # =========================================================
    slide2 = prs.slides.add_slide(blank_layout)
    slide2.shapes.add_picture('scratch/slide_bg_content.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_top_right_title(slide2, "Introducción")
    
    # Main narrative card (Left)
    card_main2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), Inches(7.5), Inches(4.4))
    card_main2.fill.solid()
    card_main2.fill.fore_color.rgb = CLR_WHITE
    card_main2.line.color.rgb = CLR_BORDER_CARD
    card_main2.line.width = Pt(1.5)
    tf_m2 = card_main2.text_frame
    tf_m2.word_wrap = True
    
    p = tf_m2.paragraphs[0]
    p.text = "Contexto y Necesidad del Proyecto"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CLR_DARK_OLIVE
    
    p = tf_m2.add_paragraph()
    p.text = (
        "Las redes ópticas pasivas con capacidad de gigabit (GPON) y arquitecturas de fibra hasta el hogar (FTTH) "
        "representan la infraestructura de telecomunicaciones más veloz y confiable en la actualidad. "
        "Sin embargo, su despliegue físico en municipios rurales como San José del Rincón enfrenta desafíos críticos de dispersión geográfica y orografía montañosa."
    )
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = CLR_TEXT_DARK
    p.space_after = Pt(10)
    
    p = tf_m2.add_paragraph()
    p.text = (
        "La empresa GPON TELECOM S.A. de C.V. requería una solución tecnológica integral que facilite y optimice la gestión "
        "de sus activos de red: paneles ODF, hilos de fibra, mufas de empalme, splitters PLC y cajas terminales de acceso a la red (NAP)."
    )
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = CLR_TEXT_DARK
    p.space_after = Pt(10)
    
    p = tf_m2.add_paragraph()
    p.text = (
        "Se propone el desarrollo de un Sistema Web Progresivo (PWA) de Inventario y Mapeo Lógico Georreferenciado, "
        "dotado de mecanismos de control transaccional ACID contra colisiones y capacidad de operación sin conexión (Offline-First)."
    )
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.color.rgb = CLR_TEXT_DARK

    # Right Column Highlights (3 Cards)
    highlights = [
        ("Mapeo Geoespacial Activo", "Visualización satelital precisa de cajas NAP, postes y rutas de tendido sobre OpenStreetMap.", "scratch/real_gis_map.png"),
        ("Trazabilidad Integral", "Mapeo lógico extremo a extremo desde el puerto PON en la OLT hasta la roseta del abonado.", "scratch/real_chassis_matrix.png"),
        ("Operatividad Fuera de Línea", "Funcionamiento continuo de cuadrillas en zonas rurales sin cobertura celular mediante IndexedDB.", "scratch/real_modal_nap.png")
    ]
    
    for idx, (h_title, h_desc, h_img) in enumerate(highlights):
        top_y = 1.55 + idx * 1.5
        c_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(top_y), Inches(4.0), Inches(1.35))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = CLR_BG_CARD
        c_box.line.color.rgb = CLR_BORDER_CARD
        c_box.line.width = Pt(1)
        tf_c = c_box.text_frame
        tf_c.word_wrap = True
        
        ph0 = tf_c.paragraphs[0]
        ph0.text = f"• {h_title}"
        ph0.font.name = "Arial"
        ph0.font.size = Pt(13)
        ph0.font.bold = True
        ph0.font.color.rgb = CLR_DARK_OLIVE
        
        ph1 = tf_c.add_paragraph()
        ph1.text = h_desc
        ph1.font.name = "Arial"
        ph1.font.size = Pt(11)
        ph1.font.color.rgb = CLR_TEXT_DARK

    # =========================================================
    # SLIDE 3: PLANTEAMIENTO DEL PROBLEMA
    # =========================================================
    slide3 = prs.slides.add_slide(blank_layout)
    slide3.shapes.add_picture('scratch/slide_bg_content.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_top_right_title(slide3, "Planteamiento del Problema")
    
    # 4 Problem Cards in 2x2 Grid
    problems = [
        ("1. Registro Manual y Descentralizado",
         "La administración física se realizaba mediante libretas de papel y hojas de cálculo locales desarticuladas, provocando discrepancias inmediatas entre oficina y cuadrillas de campo."),
        ("2. Cruces y Sobreventa de Puertos",
         "Múltiples técnicos asignaban abonados simultáneamente sobre un mismo conector de splitter, generando desconexiones accidentales a clientes activos y quejas operativas."),
        ("3. Aislamiento en Zonas sin Señal",
         "En postes rurales sin cobertura móvil, el personal no podía verificar la disponibilidad del splitter, viéndose obligado a realizar conexiones a ciegas con alto índice de error."),
        ("4. Carencia de Georreferenciación y Presupuesto",
         "Ausencia de coordenadas GPS para ubicar cajas NAP en campo, además de desconocimiento del presupuesto de atenuación óptica (dBm), originando enlaces degradados.")
    ]
    
    positions = [
        (0.8, 1.55), (6.9, 1.55),
        (0.8, 3.65), (6.9, 3.65)
    ]
    
    for (x, y), (p_title, p_desc) in zip(positions, problems):
        p_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.6), Inches(1.95))
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = CLR_WHITE
        p_card.line.color.rgb = CLR_BORDER_CARD
        p_card.line.width = Pt(1.5)
        tf_p = p_card.text_frame
        tf_p.word_wrap = True
        
        pp0 = tf_p.paragraphs[0]
        pp0.text = p_title
        pp0.font.name = "Arial"
        pp0.font.size = Pt(14)
        pp0.font.bold = True
        pp0.font.color.rgb = RGBColor(165, 42, 42) # Crimson accent for problem
        
        pp1 = tf_p.add_paragraph()
        pp1.text = p_desc
        pp1.font.name = "Arial"
        pp1.font.size = Pt(12)
        pp1.font.color.rgb = CLR_TEXT_DARK
        pp1.space_before = Pt(6)

    # =========================================================
    # SLIDE 4: JUSTIFICACIÓN
    # =========================================================
    slide4 = prs.slides.add_slide(blank_layout)
    slide4.shapes.add_picture('scratch/slide_bg_content.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_top_right_title(slide4, "Justificación")
    
    # 3 Column Cards: Operativa, Técnica, Económica
    justifications = [
        ("Dimensión Operativa",
         "Optimización del Servicio de Campo",
         [
             "Reducción de hasta un 70% en tiempos de atención e instalación de nuevos abonados.",
             "Eliminación de traslados infructuosos de cuadrillas por falta de puertos libres en la caja NAP.",
             "Auditoría visual inmediata del estado y saturación de la infraestructura óptica en San José del Rincón."
         ]),
        ("Dimensión Técnica",
         "Integridad y Continuidad Transaccional",
         [
             "Blindaje transaccional ACID en PostgreSQL con bloqueo pesimista (SELECT ... FOR UPDATE) para cero colisiones.",
             "Arquitectura Offline-First con Dexie.js / IndexedDB para funcionamiento ininterrumpido sin señal celular.",
             "Cumplimiento de estándares internacionales ITU-T G.984 (GPON) y directrices de accesibilidad WCAG 2.1."
         ]),
        ("Dimensión Económica",
         "Eficiencia y Rentabilidad Empresarial",
         [
             "Ahorro directo en combustible y costos de desplazamiento de vehículos operativos.",
             "Disminución de penalizaciones y quejas de usuarios por cortes involuntarios de conexión.",
             "Plataforma escalable construida 100% sobre tecnologías Open Source sin costos de licencias privativas."
         ])
    ]
    
    for idx, (j_cat, j_sub, j_items) in enumerate(justifications):
        jx = 0.8 + idx * 4.0
        j_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(jx), Inches(1.55), Inches(3.7), Inches(4.3))
        j_card.fill.solid()
        j_card.fill.fore_color.rgb = CLR_WHITE
        j_card.line.color.rgb = CLR_BORDER_CARD
        j_card.line.width = Pt(1.5)
        tf_j = j_card.text_frame
        tf_j.word_wrap = True
        
        pj0 = tf_j.paragraphs[0]
        pj0.text = j_cat
        pj0.font.name = "Arial"
        pj0.font.size = Pt(15)
        pj0.font.bold = True
        pj0.font.color.rgb = CLR_DARK_OLIVE
        
        pj1 = tf_j.add_paragraph()
        pj1.text = j_sub
        pj1.font.name = "Arial"
        pj1.font.size = Pt(11)
        pj1.font.italic = True
        pj1.font.color.rgb = CLR_TEXT_MUTED
        pj1.space_after = Pt(12)
        
        for item in j_items:
            p_it = tf_j.add_paragraph()
            p_it.text = f"- {item}"
            p_it.font.name = "Arial"
            p_it.font.size = Pt(11.5)
            p_it.font.color.rgb = CLR_TEXT_DARK
            p_it.space_after = Pt(8)

    # =========================================================
    # SLIDE 5: OBJETIVOS
    # =========================================================
    slide5 = prs.slides.add_slide(blank_layout)
    slide5.shapes.add_picture('scratch/slide_bg_content.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_top_right_title(slide5, "Objetivos del Proyecto")
    
    # Top Card: Objetivo General
    card_gen = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), Inches(11.733), Inches(1.65))
    card_gen.fill.solid()
    card_gen.fill.fore_color.rgb = CLR_BG_CARD
    card_gen.line.color.rgb = CLR_BORDER_CARD
    card_gen.line.width = Pt(1.5)
    tf_gen = card_gen.text_frame
    tf_gen.word_wrap = True
    
    pg0 = tf_gen.paragraphs[0]
    pg0.text = "OBJETIVO GENERAL"
    pg0.font.name = "Arial"
    pg0.font.size = Pt(13)
    pg0.font.bold = True
    pg0.font.color.rgb = CLR_DARK_OLIVE
    
    pg1 = tf_gen.add_paragraph()
    pg1.text = (
        "Desarrollar una aplicación web para el inventario, mapeo lógico y georreferenciación de infraestructura de red GPON / FTTx "
        "para la empresa GPON TELECOM S.A. de C.V., optimizando la trazabilidad de puertos ópticos y la asignación de abonados mediante "
        "tecnologías modernas y capacidades de sincronización fuera de línea."
    )
    pg1.font.name = "Arial"
    pg1.font.size = Pt(13)
    pg1.font.color.rgb = CLR_TEXT_DARK
    pg1.space_before = Pt(4)

    # Bottom 4 Cards: Objetivos Específicos
    spec_objs = [
        ("1. Analizar y Especificar",
         "Levantar la información técnica de planta externa en San José del Rincón y modelar los casos de uso bajo estándar UML e IEEE 830."),
        ("2. Diseñar la Base de Datos",
         "Modelar el esquema relacional en PostgreSQL bajo Tercera Forma Normal (3FN), definiendo diccionarios y reglas de integridad."),
        ("3. Programar el Backend Modular",
         "Construir la API RESTful en Node.js y TypeScript, incorporando control transaccional pesimista contra sobreventa de puertos y JWT/RBAC."),
        ("4. Implementar Interfaz y Offline",
         "Desarrollar el visor GIS con Leaflet, matriz de 16 puertos en React, sincronización offline con Dexie.js y streaming de reportes en PDF.")
    ]
    
    for idx, (s_title, s_desc) in enumerate(spec_objs):
        sx = 0.8 + idx * 2.98
        s_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(sx), Inches(3.4), Inches(2.78), Inches(2.45))
        s_card.fill.solid()
        s_card.fill.fore_color.rgb = CLR_WHITE
        s_card.line.color.rgb = CLR_BORDER_CARD
        s_card.line.width = Pt(1.5)
        tf_s = s_card.text_frame
        tf_s.word_wrap = True
        
        ps0 = tf_s.paragraphs[0]
        ps0.text = s_title
        ps0.font.name = "Arial"
        ps0.font.size = Pt(12)
        ps0.font.bold = True
        ps0.font.color.rgb = CLR_DARK_OLIVE
        
        ps1 = tf_s.add_paragraph()
        ps1.text = s_desc
        ps1.font.name = "Arial"
        ps1.font.size = Pt(11)
        ps1.font.color.rgb = CLR_TEXT_DARK
        ps1.space_before = Pt(6)

    # =========================================================
    # SLIDE 6: HERRAMIENTAS Y STACK TECNOLÓGICO
    # =========================================================
    slide6 = prs.slides.add_slide(blank_layout)
    slide6.shapes.add_picture('scratch/slide_bg_content.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_top_right_title(slide6, "Herramientas y Stack Tecnológico")
    
    # 4 Stack Layer Cards (2x2 Grid)
    tech_layers = [
        ("Backend & Servicios RESTful",
         "Capa de Lógica de Negocio y Seguridad",
         [
             "Node.js & Express.js: Entorno de ejecución asíncrono no bloqueante.",
             "TypeScript: Tipado estático estricto y prevención de errores.",
             "Zod: Validación declarativa y sanitización de esquemas de datos.",
             "JSON Web Tokens (JWT) & Bcrypt: Autenticación criptográfica sin estado y control RBAC."
         ]),
        ("Base de Datos & Concurrencia",
         "Capa de Persistencia y Blindaje ACID",
         [
             "PostgreSQL 16: Motor relacional con integridad referencial robusta.",
             "Neon Serverless: Base de datos en la nube con almacenamiento desagregado.",
             "Sequelize ORM: Abstracción objeto-relacional con migraciones tipadas.",
             "Bloqueo Pesimista (SELECT ... FOR UPDATE): Eliminación total de sobreventa de puertos."
         ]),
        ("Frontend & Cartografía GIS",
         "Capa de Experiencia de Usuario Reactiva",
         [
             "React 18 & Vite: Arquitectura modular por componentes y compilación ultrarrápida.",
             "Tailwind CSS: Diseño utility-first adaptativo con paleta institucional UMB.",
             "Leaflet & React-Leaflet: Renderizado de cartografía web geoespacial interactiva.",
             "Adobe Color Contrast Analyzer: Verificación formal de accesibilidad WCAG 2.1 AA/AAA."
         ]),
        ("Movilidad, Reportes & Despliegue",
         "Capa de Infraestructura y Operación de Campo",
         [
             "PWA & Service Workers: Capacidad de instalación multiplataforma y caché web.",
             "Dexie.js & IndexedDB: Almacenamiento local en navegador y sincronización diferida.",
             "PDFKit: Generación dinámica en memoria de reportes técnicos ejecutivos en PDF.",
             "Docker & Docker Compose: Contenerización y despliegue modular reproducible."
         ])
    ]
    
    tech_positions = [
        (0.8, 1.55), (6.9, 1.55),
        (0.8, 3.65), (6.9, 3.65)
    ]
    
    for (tx, ty), (t_title, t_sub, t_items) in zip(tech_positions, tech_layers):
        t_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(tx), Inches(ty), Inches(5.6), Inches(1.95))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = CLR_WHITE
        t_card.line.color.rgb = CLR_BORDER_CARD
        t_card.line.width = Pt(1.5)
        tf_t = t_card.text_frame
        tf_t.word_wrap = True
        
        pt0 = tf_t.paragraphs[0]
        pt0.text = t_title
        pt0.font.name = "Arial"
        pt0.font.size = Pt(13)
        pt0.font.bold = True
        pt0.font.color.rgb = CLR_DARK_OLIVE
        
        pt1 = tf_t.add_paragraph()
        pt1.text = t_sub
        pt1.font.name = "Arial"
        pt1.font.size = Pt(10.5)
        pt1.font.italic = True
        pt1.font.color.rgb = CLR_TEXT_MUTED
        pt1.space_after = Pt(4)
        
        for item in t_items:
            pit = tf_t.add_paragraph()
            pit.text = f"• {item}"
            pit.font.name = "Arial"
            pit.font.size = Pt(10)
            pit.font.color.rgb = CLR_TEXT_DARK

    # =========================================================
    # SLIDE 7: AVANCE DEL PROYECTO (EVIDENCIAS DE DESARROLLO)
    # =========================================================
    slide7 = prs.slides.add_slide(blank_layout)
    slide7.shapes.add_picture('scratch/slide_bg_content.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_top_right_title(slide7, "Avance del Proyecto")
    
    # 4 Module Cards with Screen captures
    modules = [
        ("Visor Cartográfico GIS", "scratch/real_gis_map.png", "Georreferenciación interactiva de cajas NAP con semaforización de ocupación en tiempo real."),
        ("Matriz de 16 Puertos", "scratch/real_chassis_matrix.png", "Chasis visual de distribución para asignación atómica de abonados y prevención de sobreventa."),
        ("Modal de Registro GPS", "scratch/real_modal_nap.png", "Formulario validado con Zod para alta y calibración de coordenadas en sitio de nuevas cajas."),
        ("Directorio de Clientes", "scratch/real_clients_table.png", "Padrón general con cálculo de presupuesto óptico (dBm) y exportación a PDF vía streaming.")
    ]
    
    for idx, (m_title, m_img_path, m_desc) in enumerate(modules):
        mx = 0.8 + idx * 2.98
        m_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(mx), Inches(1.55), Inches(2.78), Inches(4.3))
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = CLR_WHITE
        m_card.line.color.rgb = CLR_BORDER_CARD
        m_card.line.width = Pt(1.5)
        
        tf_m = m_card.text_frame
        tf_m.word_wrap = True
        
        pm0 = tf_m.paragraphs[0]
        pm0.text = m_title
        pm0.font.name = "Arial"
        pm0.font.size = Pt(12)
        pm0.font.bold = True
        pm0.font.color.rgb = CLR_DARK_OLIVE
        pm0.alignment = PP_ALIGN.CENTER
        
        # Add screenshot inside card
        if os.path.exists(m_img_path):
            slide7.shapes.add_picture(m_img_path, Inches(mx + 0.15), Inches(2.05), width=Inches(2.48))
            
        # Description at bottom of card
        desc_box = slide7.shapes.add_textbox(Inches(mx + 0.1), Inches(4.35), Inches(2.58), Inches(1.4))
        tf_d = desc_box.text_frame
        tf_d.word_wrap = True
        pd = tf_d.paragraphs[0]
        pd.text = m_desc
        pd.font.name = "Arial"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = CLR_TEXT_DARK

    # =========================================================
    # SLIDE 8: CONCLUSIÓN Y PRÓXIMOS PASOS
    # =========================================================
    slide8 = prs.slides.add_slide(blank_layout)
    slide8.shapes.add_picture('scratch/slide_bg_content.png', 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_top_right_title(slide8, "Conclusiones")
    
    # Left Card: Conclusiones del Avance
    card_c1 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), Inches(5.6), Inches(4.3))
    card_c1.fill.solid()
    card_c1.fill.fore_color.rgb = CLR_WHITE
    card_c1.line.color.rgb = CLR_BORDER_CARD
    card_c1.line.width = Pt(1.5)
    tf_c1 = card_c1.text_frame
    tf_c1.word_wrap = True
    
    pc1_0 = tf_c1.paragraphs[0]
    pc1_0.text = "Logros Alcanzados en el Avance"
    pc1_0.font.name = "Arial"
    pc1_0.font.size = Pt(15)
    pc1_0.font.bold = True
    pc1_0.font.color.rgb = CLR_DARK_OLIVE
    pc1_0.space_after = Pt(10)
    
    concl_points = [
        "Resolución de la problemática central: Se erradicó el registro manual en papel, unificando el catastro de fibra óptica en una plataforma centralizada.",
        "Blindaje anti-colisión comprobado: La implementación del bloqueo pesimista en PostgreSQL garantiza cero sobreventa de puertos en condiciones de alta concurrencia.",
        "Operatividad en zonas remotas: La arquitectura Offline-First con Dexie.js permite a las cuadrillas registrar instalaciones en postes sin conexión celular.",
        "Cumplimiento de estándares de calidad: Se validaron los contrastes cromáticos bajo normas internacionales WCAG 2.1 y la topología según ITU-T G.984."
    ]
    for pt in concl_points:
        p_c = tf_c1.add_paragraph()
        p_c.text = f"- {pt}"
        p_c.font.name = "Arial"
        p_c.font.size = Pt(11.5)
        p_c.font.color.rgb = CLR_TEXT_DARK
        p_c.space_after = Pt(8)

    # Right Card: Próximas Actividades
    card_c2 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.55), Inches(5.6), Inches(4.3))
    card_c2.fill.solid()
    card_c2.fill.fore_color.rgb = CLR_BG_CARD
    card_c2.line.color.rgb = CLR_BORDER_CARD
    card_c2.line.width = Pt(1.5)
    tf_c2 = card_c2.text_frame
    tf_c2.word_wrap = True
    
    pc2_0 = tf_c2.paragraphs[0]
    pc2_0.text = "Próximas Actividades y Cierre"
    pc2_0.font.name = "Arial"
    pc2_0.font.size = Pt(15)
    pc2_0.font.bold = True
    pc2_0.font.color.rgb = CLR_DARK_OLIVE
    pc2_0.space_after = Pt(10)
    
    next_points = [
        "Pruebas de campo exhaustivas: Validación directa de cuadrillas técnicas en la ruta troncal de San José del Rincón.",
        "Calibración de telemetría: Integración de mediciones de potencia óptica en vivo (ONT/OLT) para cálculo automático de atenuación.",
        "Despliegue cloud serverless: Puesta en marcha productiva sobre Neon PostgreSQL y plataforma de alojamiento en la nube.",
        "Documentación formal: Integración de métricas finales de rendimiento para la memoria de residencia profesional."
    ]
    for pt in next_points:
        p_n = tf_c2.add_paragraph()
        p_n.text = f"- {pt}"
        p_n.font.name = "Arial"
        p_n.font.size = Pt(11.5)
        p_n.font.color.rgb = CLR_TEXT_DARK
        p_n.space_after = Pt(8)

    # Save presentation
    output_path = "docs/PRESENTACION_AVANCE_RESIDENCIA_GPON.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    generate_background_images()
    build_presentation()
