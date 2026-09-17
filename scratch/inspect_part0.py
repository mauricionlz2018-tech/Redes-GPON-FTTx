import docx

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
for i in range(0, 130):
    t = doc.paragraphs[i].text.strip()
    if t:
        print(f"P[{i}] ({doc.paragraphs[i].style.name}): {t[:75]}")

