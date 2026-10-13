import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

print("--- INSPECCIÓN DE TABLA 13 (Matriz de trazabilidad) y TABLA 2 ---")
for idx in [2, 13, 15]:
    t = doc.tables[idx]
    print(f"\nTabla idx {idx}: style={t.style.name}, rows={len(t.rows)}, cols={len(t.columns)}")
    tblPr = t._tbl.tblPr
    print("tblPr XML:", tblPr.xml if tblPr is not None else "None")
    c0 = t.rows[0].cells[0]
    tcPr0 = c0._tc.tcPr
    print("Header cell tcPr XML:", tcPr0.xml if tcPr0 is not None else "None")
    if c0.paragraphs and c0.paragraphs[0].runs:
        r = c0.paragraphs[0].runs[0]
        print(f"Header run: font={r.font.name}, size={r.font.size}, bold={r.font.bold}, color={r.font.color.rgb if r.font.color else None}")
    c1 = t.rows[1].cells[0]
    tcPr1 = c1._tc.tcPr
    print("Body cell tcPr XML:", tcPr1.xml if tcPr1 is not None else "None")
    if c1.paragraphs and c1.paragraphs[0].runs:
        r = c1.paragraphs[0].runs[0]
        print(f"Body run: font={r.font.name}, size={r.font.size}, bold={r.font.bold}, color={r.font.color.rgb if r.font.color else None}")

print("\n--- PÁRRAFOS CON ESTILO 'Tablas' EN EL DOCUMENTO ---")
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'Tablas' or (p.text.strip().startswith('Tabla ') and '\t' not in p.text):
        print(f"P{i} [{p.style.name}]: {p.text}")
        if i+1 < len(doc.paragraphs):
            print(f"  Next P{i+1} [{doc.paragraphs[i+1].style.name}]: {doc.paragraphs[i+1].text[:60]}")

