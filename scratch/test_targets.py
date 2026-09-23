import docx
import re
import os

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

# Check target 1:
p1 = None
for p in doc.paragraphs:
    if '4.- API Geoespacial GPS / Leaflet' in p.text:
        p1 = p
        break

# Check target 2:
p2 = None
for i, p in enumerate(doc.paragraphs):
    if 'Figura 29. Cadena de distribuci' in p.text:
        p2 = doc.paragraphs[i+1] # The Elaboración propia
        break

# Check target 3:
p3 = None
for i, p in enumerate(doc.paragraphs):
    if 'Figura 31. Diagrama de secuencia transaccional' in p.text:
        p3 = doc.paragraphs[i+1] # The Elaboración propia
        break

# Check target 4:
p4 = None
for p in doc.paragraphs:
    if '4.- Wireframe 4: Padr' in p.text:
        p4 = p
        break

print(f"Target 1 found: {p1 is not None} -> {p1.text[:50] if p1 else ''}")
print(f"Target 2 found: {p2 is not None} -> {p2.text[:50] if p2 else ''}")
print(f"Target 3 found: {p3 is not None} -> {p3.text[:50] if p3 else ''}")
print(f"Target 4 found: {p4 is not None} -> {p4.text[:50] if p4 else ''}")

# Check all 7 images exist
imgs = [
    'scratch/diag_robustez_asignacion.png',
    'scratch/diag_topologia_gpon.png',
    'scratch/diag_estados_puerto.png',
    'scratch/diag_secuencia_offline.png',
    'scratch/diag_navegacion_sistema.png',
    'scratch/diag_arquitectura_informacion.png',
    'scratch/diag_paquetes_componentes.png'
]
for img in imgs:
    print(f"Image {img} exists: {os.path.exists(img)} ({os.path.getsize(img)} bytes)")

