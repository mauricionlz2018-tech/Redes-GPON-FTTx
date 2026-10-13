import docx

doc = docx.Document('docs/Documentacion_Residencias_avance_BACKUP_08OCT_FINAL_PAG.docx')

def check_range(title, start, end):
    print(f"\n=== {title} (P{start} to P{end}) ===")
    for i in range(max(0, start), min(len(doc.paragraphs), end + 1)):
        p = doc.paragraphs[i]
        has_sp = bool(p._element.xpath('./w:pPr/w:sectPr'))
        t = p.text.strip()[:60]
        if t or has_sp:
            print(f"P{i} [sectPr={has_sp}]: '{t}'")

check_range("Indices to Intro", 155, 163)
check_range("Objetivos to Cap I", 183, 196)
check_range("Cap I to Cap II", 240, 252)
check_range("Cap II to Cap III", 750, 762)
check_range("Cap III to Cap IV to Cap V", 1188, 1205)
check_range("Cap V to Conclusiones", 1205, 1215)
check_range("Conclusiones to Referencias", 1234, 1242)
check_range("Referencias to Anexos", 1370, 1378)

