import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document('docs/PORTADA_INSTITUCIONAL (3)_FINAL.docx')
target_p = doc.paragraphs[433]
print(f"Target paragraph: '{target_p.text[:60]}'")

# Test inserting a paragraph after target_p
new_p = doc.add_paragraph()
new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = new_p.add_run("TEST INSERT AFTER 433")
r.font.name = "Arial"
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(0, 0, 0)
target_p._p.addnext(new_p._p)

doc.save('scratch/test_insert_after.docx')
print("Successfully saved test_insert_after.docx")

# Verify
doc2 = docx.Document('scratch/test_insert_after.docx')
for i in range(430, 437):
    print(f"P[{i}]: '{doc2.paragraphs[i].text[:60]}'")

