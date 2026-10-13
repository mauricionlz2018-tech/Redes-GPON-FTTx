import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')
p886 = doc.paragraphs[886]
p887 = doc.paragraphs[887]
t13 = doc.tables[13]

print("--- P886 XML ---")
print(p886._element.xml)
print("--- P887 XML ---")
print(p887._element.xml)
print("--- T13 Row 0 Cell 0 XML ---")
print(t13.rows[0].cells[0]._tc.xml)
print("--- T13 Row 1 Cell 0 XML ---")
print(t13.rows[1].cells[0]._tc.xml)

