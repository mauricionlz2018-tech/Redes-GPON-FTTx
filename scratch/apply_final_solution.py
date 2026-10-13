import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import shutil

src_path = 'docs/Documentacion_Residencias_avance_BACKUP_08OCT_FINAL_PAG.docx'
target_avance = 'docs/Documentacion_Residencias_avance.docx'
target_mirror = 'docs/Documentacion_Residencias (4).docx'

doc = docx.Document(src_path)

def remove_page_break_before(p):
    pPr = p._element.find(qn('w:pPr'))
    if pPr is not None:
        pbb = pPr.find(qn('w:pageBreakBefore'))
        if pbb is not None:
            pPr.remove(pbb)

print(f"Total párrafos: {len(doc.paragraphs)}")

# 1. SECCIÓN 1: ÍNDICES (Párrafos 5 a 159)
# Número romano mayúscula iniciando en III (Carátula y oficio = I y II)
# Header: rId13 (vacío, sin logos)
# Footer: rId14 (número de página)
p159 = doc.paragraphs[159]
p159_sp = p159._element.xpath('./w:pPr/w:sectPr')[0]
for el in list(p159_sp):
    p159_sp.remove(el)

p159_sp.append(parse_xml(f'<w:headerReference {nsdecls("w", "r")} w:type="default" r:id="rId13"/>'))
p159_sp.append(parse_xml(f'<w:footerReference {nsdecls("w", "r")} w:type="default" r:id="rId14"/>'))
p159_sp.append(parse_xml(f'<w:pgSz {nsdecls("w")} w:w="12242" w:h="15842" w:code="1"/>'))
p159_sp.append(parse_xml(f'<w:pgMar {nsdecls("w")} w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>'))
p159_sp.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="upperRoman" w:start="3"/>'))
p159_sp.append(parse_xml(f'<w:cols {nsdecls("w")} w:space="708"/>'))
p159_sp.append(parse_xml(f'<w:type {nsdecls("w")} w:val="nextPage"/>'))
p159_sp.append(parse_xml(f'<w:docGrid {nsdecls("w")} w:linePitch="360"/>'))
print("1. Índices: Romano start=3, header limpio sin logos.")

# 2. SECCIÓN 2: INTRODUCCIÓN, PLANTEAMIENTO, JUSTIFICACIÓN, OBJETIVOS (Párrafos 160 a 186)
# Romanos mayúsculas consecutivos de la sección anterior. Header vacío, footer con página.
idx_obj = -1
for i, p in enumerate(doc.paragraphs):
    if '3. Programar el backend del sistema' in p.text:
        p_end_obj = p
        idx_obj = i
        break

pPr_obj = p_end_obj._element.get_or_add_pPr()
for sp in pPr_obj.xpath('./w:sectPr'):
    pPr_obj.remove(sp)

new_sp_obj = parse_xml(f'''
<w:sectPr {nsdecls("w", "r")}>
    <w:headerReference w:type="default" r:id="rId13"/>
    <w:footerReference w:type="default" r:id="rId14"/>
    <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
    <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
    <w:pgNumType w:fmt="upperRoman"/>
    <w:cols w:space="708"/>
    <w:type w:val="nextPage"/>
    <w:docGrid w:linePitch="360"/>
</w:sectPr>
''')
pPr_obj.append(new_sp_obj)
print("2. Introducción hasta Objetivos: Romano consecutivo, página nueva garantizada, sin logos.")

# Quitar salto manual redundante al inicio de Cap I
remove_page_break_before(doc.paragraphs[idx_obj + 1])

# Función para sectPr de Capítulos (arábigos, portadilla sin número visible pero contada)
def make_chapter_sectPr():
    return parse_xml(f'''
    <w:sectPr {nsdecls("w", "r")}>
        <w:headerReference w:type="default" r:id="rId13"/>
        <w:headerReference w:type="first" r:id="rId13"/>
        <w:footerReference w:type="default" r:id="rId92"/>
        <w:footerReference w:type="first" r:id="rId10"/>
        <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
        <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
        <w:pgNumType w:fmt="decimal"/>
        <w:cols w:space="708"/>
        <w:type w:val="nextPage"/>
        <w:titlePg/>
        <w:docGrid w:linePitch="360"/>
    </w:sectPr>
    ''')

# 3. CAPÍTULO I (termina antes de Cap II)
p_end_cap1 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO II. MARCO' in p.text.upper():
        p_end_cap1 = doc.paragraphs[i-8]
        remove_page_break_before(doc.paragraphs[i-7])
        break

pPr_c1 = p_end_cap1._element.get_or_add_pPr()
for sp in pPr_c1.xpath('./w:sectPr'): pPr_c1.remove(sp)
pPr_c1.append(make_chapter_sectPr())
print("3. Capítulo I: Arábigos consecutivos, portadilla no numerada pero contada.")

# 4. CAPÍTULO II (termina antes de Cap III)
p_end_cap2 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO III. DISEÑO' in p.text.upper():
        p_end_cap2 = doc.paragraphs[i-7]
        remove_page_break_before(doc.paragraphs[i-6])
        break

pPr_c2 = p_end_cap2._element.get_or_add_pPr()
for sp in pPr_c2.xpath('./w:sectPr'): pPr_c2.remove(sp)
pPr_c2.append(make_chapter_sectPr())
print("4. Capítulo II: Arábigos consecutivos, portadilla no numerada pero contada.")

# 5. CAPÍTULO III (termina antes de Cap IV)
p_end_cap3 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO IV. DESARROLLO' in p.text.upper():
        p_end_cap3 = doc.paragraphs[i-7]
        remove_page_break_before(doc.paragraphs[i-6])
        break

pPr_c3 = p_end_cap3._element.get_or_add_pPr()
for sp in pPr_c3.xpath('./w:sectPr'): pPr_c3.remove(sp)
pPr_c3.append(make_chapter_sectPr())
print("5. Capítulo III: Arábigos consecutivos, portadilla no numerada pero contada.")

# 6. CAPÍTULO IV (termina antes de Cap V)
p_end_cap4 = None
for i, p in enumerate(doc.paragraphs):
    if 'CAPÍTULO V. PRUEBAS' in p.text.upper():
        p_end_cap4 = doc.paragraphs[i-7]
        remove_page_break_before(doc.paragraphs[i-6])
        break

pPr_c4 = p_end_cap4._element.get_or_add_pPr()
for sp in pPr_c4.xpath('./w:sectPr'): pPr_c4.remove(sp)
pPr_c4.append(make_chapter_sectPr())
print("6. Capítulo IV: Arábigos consecutivos, portadilla no numerada pero contada.")

# 7. CAPÍTULO V (termina antes de Conclusiones)
p_end_cap5 = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'Conclusiones':
        p_end_cap5 = doc.paragraphs[i-1]
        remove_page_break_before(p)
        break

pPr_c5 = p_end_cap5._element.get_or_add_pPr()
for sp in pPr_c5.xpath('./w:sectPr'): pPr_c5.remove(sp)
pPr_c5.append(make_chapter_sectPr())
print("7. Capítulo V: Arábigos consecutivos, portadilla no numerada pero contada.")

# 8. CONCLUSIONES Y GLOSARIO (termina antes de REFERENCIAS)
# Romanos mayúsculas consecutivos a la paginación
p_end_concl = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'REFERENCIAS':
        p_end_concl = doc.paragraphs[i-1]
        break

pPr_concl = p_end_concl._element.get_or_add_pPr()
for sp in pPr_concl.xpath('./w:sectPr'): pPr_concl.remove(sp)
new_sp_concl = parse_xml(f'''
<w:sectPr {nsdecls("w", "r")}>
    <w:headerReference w:type="default" r:id="rId13"/>
    <w:footerReference w:type="default" r:id="rId14"/>
    <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
    <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
    <w:pgNumType w:fmt="upperRoman"/>
    <w:cols w:space="708"/>
    <w:type w:val="nextPage"/>
    <w:docGrid w:linePitch="360"/>
</w:sectPr>
''')
pPr_concl.append(new_sp_concl)
print("8. Conclusiones y Glosario: Romanos consecutivos, sin logos.")

# 9. REFERENCIAS (termina antes de Anexos)
# Sin numeración de página (footer rId10 vacío), sin logos
p_end_ref = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'Anexos':
        p_end_ref = doc.paragraphs[i-1]
        break

pPr_ref = p_end_ref._element.get_or_add_pPr()
for sp in pPr_ref.xpath('./w:sectPr'): pPr_ref.remove(sp)
new_sp_ref = parse_xml(f'''
<w:sectPr {nsdecls("w", "r")}>
    <w:headerReference w:type="default" r:id="rId13"/>
    <w:footerReference w:type="default" r:id="rId10"/>
    <w:pgSz w:w="12242" w:h="15842" w:code="1"/>
    <w:pgMar w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>
    <w:cols w:space="708"/>
    <w:type w:val="nextPage"/>
    <w:docGrid w:linePitch="360"/>
</w:sectPr>
''')
pPr_ref.append(new_sp_ref)
print("9. Referencias: Sin numeración visible, sin logos.")

# 10. ANEXOS (última sección en body)
# Sin numeración de página (footer rId10 vacío), sin logos
body_sp = doc._body._element.xpath('./w:sectPr')[0]
for el in list(body_sp): body_sp.remove(el)
body_sp.append(parse_xml(f'<w:headerReference {nsdecls("w", "r")} w:type="default" r:id="rId13"/>'))
body_sp.append(parse_xml(f'<w:footerReference {nsdecls("w", "r")} w:type="default" r:id="rId10"/>'))
body_sp.append(parse_xml(f'<w:pgSz {nsdecls("w")} w:w="12242" w:h="15842" w:code="1"/>'))
body_sp.append(parse_xml(f'<w:pgMar {nsdecls("w")} w:top="1417" w:right="1701" w:bottom="1417" w:left="1701" w:header="708" w:footer="708" w:gutter="0"/>'))
body_sp.append(parse_xml(f'<w:cols {nsdecls("w")} w:space="708"/>'))
body_sp.append(parse_xml(f'<w:docGrid {nsdecls("w")} w:linePitch="360"/>'))
print("10. Anexos: Sin numeración visible, sin logos.")

# Guardar en Documentacion_Residencias_avance.docx y espejo
doc.save(target_avance)
shutil.copyfile(target_avance, target_mirror)
print("\n¡Documentos guardados exitosamente!")
