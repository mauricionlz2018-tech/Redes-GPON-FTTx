import docx

doc = docx.Document('docs/Documentacion_Residencias_avance.docx')

with open('scratch/sec_3_2_dump.txt', 'w', encoding='utf-8') as f:
    f.write("=== PÁRRAFOS DE SECCIÓN 3.2 (P735 a P915) ===\n")
    for i in range(735, min(915, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        f.write(f"P{i} [{p.style.name}]: {p.text}\n")

    f.write("\n=== TABLAS EXISTENTES EN EL DOCUMENTO ===\n")
    f.write(f"Total tablas: {len(doc.tables)}\n")
    for idx, t in enumerate(doc.tables):
        rows = len(t.rows)
        cols = len(t.columns)
        headers = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells] if rows > 0 else []
        f.write(f"Tabla idx {idx}: {rows}x{cols} | style='{t.style.name}' | headers={headers}\n")

print("Volcado completado.")

