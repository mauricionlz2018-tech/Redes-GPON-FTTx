import pptx
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(os.path.join('docs', 'PRESENTACION_AVANCE_RESIDENCIA_GPON_FINAL.pptx'))

for i, s in enumerate(prs.slides):
    print(f"\n================ SLIDE {i+1} ================")
    for shp in s.shapes:
        txt = shp.text.strip() if shp.has_text_frame else ''
        if txt:
            print(f"  Shape '{shp.name}':")
            for p in shp.text_frame.paragraphs:
                ptxt = p.text.strip()
                if ptxt:
                    print(f"    - {ptxt}")
        elif shp.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  Picture '{shp.name}' pos=({shp.left/914400:.2f}, {shp.top/914400:.2f}), size=({shp.width/914400:.2f} x {shp.height/914400:.2f})")
